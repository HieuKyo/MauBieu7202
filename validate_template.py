#!/usr/bin/env python3
"""
Script để validate Jinja2 syntax trong Word template
Giúp phát hiện lỗi trước khi upload lên hệ thống
"""
import sys
from docx import Document
from jinja2 import Environment, BaseLoader, TemplateSyntaxError

def validate_word_template(template_path):
    """
    Validate tất cả Jinja2 syntax trong Word template

    Args:
        template_path: Đường dẫn đến file .docx

    Returns:
        (is_valid, errors): Tuple (bool, list of error messages)
    """
    errors = []
    jinja_env = Environment(loader=BaseLoader())

    try:
        doc = Document(template_path)
    except Exception as e:
        return False, [f"Không thể mở file Word: {e}"]

    # Check paragraphs
    for i, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text
        if '{%' in text or '{{' in text:
            try:
                template = jinja_env.from_string(text)
                # Try to render with empty context to check syntax
                template.render({})
            except TemplateSyntaxError as e:
                errors.append(f"Paragraph {i+1}: {e}")
            except Exception as e:
                # Some runtime errors are OK (e.g., undefined variables)
                pass

    # Check tables
    for table_idx, table in enumerate(doc.tables):
        for row_idx, row in enumerate(table.rows):
            for cell_idx, cell in enumerate(row.cells):
                for para_idx, paragraph in enumerate(cell.paragraphs):
                    text = paragraph.text
                    if '{%' in text or '{{' in text:
                        try:
                            template = jinja_env.from_string(text)
                            template.render({})
                        except TemplateSyntaxError as e:
                            errors.append(
                                f"Table {table_idx+1}, Row {row_idx+1}, "
                                f"Cell {cell_idx+1}, Para {para_idx+1}: {e}"
                            )
                        except Exception:
                            pass

    # Check headers/footers
    for section_idx, section in enumerate(doc.sections):
        for para_idx, paragraph in enumerate(section.header.paragraphs):
            text = paragraph.text
            if '{%' in text or '{{' in text:
                try:
                    template = jinja_env.from_string(text)
                    template.render({})
                except TemplateSyntaxError as e:
                    errors.append(f"Header Section {section_idx+1}, Para {para_idx+1}: {e}")
                except Exception:
                    pass

        for para_idx, paragraph in enumerate(section.footer.paragraphs):
            text = paragraph.text
            if '{%' in text or '{{' in text:
                try:
                    template = jinja_env.from_string(text)
                    template.render({})
                except TemplateSyntaxError as e:
                    errors.append(f"Footer Section {section_idx+1}, Para {para_idx+1}: {e}")
                except Exception:
                    pass

    is_valid = len(errors) == 0
    return is_valid, errors


def main():
    if len(sys.argv) < 2:
        print("Cách dùng: python validate_template.py <path_to_template.docx>")
        sys.exit(1)

    template_path = sys.argv[1]

    print(f"Đang kiểm tra template: {template_path}\n")

    is_valid, errors = validate_word_template(template_path)

    if is_valid:
        print("✅ Template hợp lệ! Không có lỗi Jinja2 syntax.")
    else:
        print("❌ Template có lỗi!\n")
        print(f"Tìm thấy {len(errors)} lỗi:\n")
        for i, error in enumerate(errors, 1):
            print(f"{i}. {error}")
        print("\n💡 Hướng dẫn sửa lỗi:")
        print("- Kiểm tra tất cả {% if %} có {% endif %}")
        print("- Kiểm tra tất cả {% for %} có {% endfor %}")
        print("- Kiểm tra tất cả {% set %} đúng syntax")
        print("- Kiểm tra tất cả {% macro %} có {% endmacro %}")
        print("- Kiểm tra {{ variable }} không có ký tự đặc biệt")
        sys.exit(1)


if __name__ == '__main__':
    main()
