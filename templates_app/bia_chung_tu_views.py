"""
Tạo bìa chứng từ (Kế toán Ngân quỹ) — in bìa lưu trữ Nhật ký chứng từ (mẫu M01/LHS)
kèm mã QR, thay cho file "Mau in bia chung tu.xlsb" + mail merge Word.

- Trưởng/Phó phòng Kế toán & Ngân quỹ cấp số hồ sơ theo ngày cho các GDV cùng đơn vị,
  sửa/xoá, nhập số tờ, in PDF; import file Excel để in bù bìa cũ.
- Số hồ sơ "<Mã CN>-<năm>.<NNNN>" đánh riêng theo từng đơn vị (UserProfile.branch),
  mỗi năm đếm lại từ 0001.
- Nội dung QR giữ đúng như file xlsb: <Mã CN>-<UserIPCAS>-<dd/mm/yyyy>-<tập>/<tổng>-<Số hồ sơ>.
"""
import io
import os
import re
from datetime import date, datetime

import openpyxl
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.paginator import Paginator
from django.db import transaction
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.http import require_http_methods, require_POST
from reportlab.graphics import renderPDF
from reportlab.graphics.barcode.qr import QrCodeWidget
from reportlab.graphics.shapes import Drawing
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

from .models import BiaChungTu, BiaThuTuIn, BranchConfig, UserProfile

# Mẫu Word dùng Times New Roman; không có thì rơi về Helvetica (không dấu).
FONT_REGULAR = 'Helvetica'
FONT_BOLD = 'Helvetica-Bold'
if os.path.exists(r'C:\Windows\Fonts\times.ttf'):
    try:
        pdfmetrics.registerFont(TTFont('BiaTimes', r'C:\Windows\Fonts\times.ttf'))
        pdfmetrics.registerFont(TTFont('BiaTimesBold', r'C:\Windows\Fonts\timesbd.ttf'))
        FONT_REGULAR, FONT_BOLD = 'BiaTimes', 'BiaTimesBold'
    except Exception:
        pass

MANAGER_POSITIONS = ('TRUONG_PHONG', 'PHO_PHONG')
TITLE_RE = re.compile(r'ngày\s+(\d{1,2}/\d{1,2}/\d{4})\s*[-–]\s*Tập\s*(\d+)\s*/\s*(\d+)', re.IGNORECASE)


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _manager_branch(user):
    """Đơn vị của user nếu là Trưởng/Phó phòng KT&NQ (hoặc superuser), ngược lại None."""
    profile = getattr(user, 'profile', None)
    if not profile or not profile.branch:
        return None
    if user.is_superuser or (profile.department == 'KE_TOAN' and profile.position in MANAGER_POSITIONS):
        return profile.branch
    return None


def _deny(request):
    messages.error(request, 'Chỉ Trưởng/Phó phòng Kế toán & Ngân quỹ (đã khai báo đơn vị) mới dùng được chức năng Tạo bìa chứng từ.')
    return redirect('dashboard')


def _back(request):
    """Quay lại trang danh sách (giữ bộ lọc) — chỉ nhận đường dẫn nội bộ."""
    nxt = request.POST.get('next', '')
    return redirect(nxt if nxt.startswith('/') and not nxt.startswith('//') else 'bia_chung_tu_list')


def _ma_cn(branch):
    cfg = BranchConfig.objects.filter(branch_code=branch).first()
    return (cfg.ma_chi_nhanh if cfg else '') or ''


def _ipcas(user):
    profile = getattr(user, 'profile', None)
    return (profile.ipcas_user or '').strip() if profile else ''


def _gdv_qs(branch):
    """GDV phòng KT&NQ cùng đơn vị (đang hoạt động)."""
    return UserProfile.objects.filter(
        branch=branch, department='KE_TOAN', job_function='GIAO_DICH_VIEN', user__is_active=True,
    ).order_by('full_name')


def _gdv_choices(branch):
    """[(user_id, họ tên)] cho các ô chọn GDV."""
    return list(_gdv_qs(branch).values_list('user_id', 'full_name'))


def _gdv_ipcas_map(branch):
    """{IPCAS (in hoa): user_id} của GDV cùng đơn vị — dùng cho cột User khi import."""
    return {(ip or '').strip().upper(): uid for uid, ip in _gdv_qs(branch).values_list('user_id', 'ipcas_user') if ip}


def _ho_ten(bia):
    """Họ tên in trên nhãn QR: GDV sở hữu tập, bìa import (không có GDV) thì là người import."""
    profile = getattr(bia.gdv or bia.nguoi_lap, 'profile', None)
    return profile.full_name if profile else ''


def _next_so(branch, ma_cn, year):
    """Số thứ tự tiếp theo trong năm của đơn vị (lớn nhất + 1)."""
    prefix = f'{ma_cn}-{year}.'
    nums = [
        int(s[len(prefix):])
        for s in BiaChungTu.objects.filter(branch=branch, so_ho_so__startswith=prefix).values_list('so_ho_so', flat=True)
        if s[len(prefix):].isdigit()
    ]
    return max(nums, default=0) + 1


def _format_so(ma_cn, year, n):
    return f'{ma_cn}-{year}.{n:04d}'


def _lech_tap_days(branch, days):
    """Các ngày có Tập x/N không khớp (thiếu/thừa tập, N khác nhau) — chỉ để cảnh báo."""
    by_day = {}
    for d, tap, tong in BiaChungTu.objects.filter(branch=branch, ngay_chung_tu__in=days).values_list(
            'ngay_chung_tu', 'tap_so', 'tong_tap'):
        by_day.setdefault(d, []).append((tap, tong))
    lech = set()
    for d, items in by_day.items():
        n = len(items)
        if sorted(t for t, _ in items) != list(range(1, n + 1)) or any(tong != n for _, tong in items):
            lech.add(d)
    return lech


def _parse_date(val):
    if isinstance(val, datetime):
        return val.date()
    if isinstance(val, date):
        return val
    s = str(val or '').strip()
    for fmt in ('%d/%m/%Y', '%Y-%m-%d'):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def _qr_text(bia, ma_cn):
    return f'{ma_cn}-{bia.user_ipcas}-{bia.ngay_chung_tu:%d/%m/%Y}-{bia.tap_so}/{bia.tong_tap}-{bia.so_ho_so}'


# ---------------------------------------------------------------------------
# Danh sách / CRUD
# ---------------------------------------------------------------------------

@login_required
def bia_list(request):
    branch = _manager_branch(request.user)
    if not branch:
        return _deny(request)

    qs = BiaChungTu.objects.filter(branch=branch).select_related('gdv__profile')
    filters = {k: request.GET.get(k, '').strip() for k in ('tu_ngay', 'den_ngay', 'gdv', 'trang_thai', 'q')}
    tu_ngay, den_ngay = _parse_date(filters['tu_ngay']), _parse_date(filters['den_ngay'])
    if tu_ngay:
        qs = qs.filter(ngay_chung_tu__gte=tu_ngay)
    if den_ngay:
        qs = qs.filter(ngay_chung_tu__lte=den_ngay)
    if filters['gdv'].isdigit():
        qs = qs.filter(gdv_id=int(filters['gdv']))
    if filters['trang_thai'] == 'da_in':
        qs = qs.filter(ngay_in__isnull=False)
    elif filters['trang_thai'] == 'chua_in':
        qs = qs.filter(ngay_in__isnull=True)
    if filters['q']:
        qs = qs.filter(so_ho_so__icontains=filters['q'])

    page_obj = Paginator(qs, 100).get_page(request.GET.get('page'))
    lech = _lech_tap_days(branch, {b.ngay_chung_tu for b in page_obj})

    query = request.GET.copy()
    query.pop('page', None)
    return render(request, 'templates_app/bia_chung_tu/list.html', {
        'page_obj': page_obj,
        'lech': lech,
        'filters': filters,
        'gdv_choices': _gdv_choices(branch),
        'query_string': query.urlencode(),
        'ipcas': _ipcas(request.user),
    })


@login_required
@require_http_methods(['GET', 'POST'])
def bia_cap_so(request):
    """Cấp số hồ sơ cho các GDV trong ngày: mỗi GDV n tập → n số liên tiếp, Tập 1/N..N/N.
    QR mang UserIPCAS của GDV sở hữu tập."""
    branch = _manager_branch(request.user)
    if not branch:
        return _deny(request)
    ma_cn = _ma_cn(branch)
    gdv_rows = list(_gdv_qs(branch).values_list('user_id', 'full_name', 'ipcas_user'))
    ngay = _parse_date(request.POST.get('ngay') or request.GET.get('ngay')) or date.today()

    if request.method == 'POST':
        # Lưu "Thứ tự in" đang nhập để lần sau tự điền sẵn (ô trống = bỏ thứ tự đã nhớ)
        for uid, _, _ in gdv_rows:
            thu_tu = request.POST.get(f'thu_tu_{uid}', '').strip()
            if thu_tu.isdigit() and int(thu_tu) > 0:
                BiaThuTuIn.objects.update_or_create(user_id=uid, defaults={'thu_tu': int(thu_tu)})
            else:
                BiaThuTuIn.objects.filter(user_id=uid).delete()

    # GDV có "Thứ tự in" đứng trước (tăng dần), còn lại giữ mặc định (họ tên A→Z)
    thu_tu_map = dict(BiaThuTuIn.objects.filter(user_id__in=[r[0] for r in gdv_rows]).values_list('user_id', 'thu_tu'))
    gdv_rows = sorted(
        [(uid, name, ipcas, thu_tu_map.get(uid)) for uid, name, ipcas in gdv_rows],
        key=lambda r: (r[3] is None, r[3] or 0),
    )

    if request.method == 'POST':
        chon, thieu_ipcas = [], []
        for uid, name, ipcas, _ in gdv_rows:
            if request.POST.get(f'chk_{uid}'):
                so_tap = request.POST.get(f'so_tap_{uid}', '1')
                chon.append((uid, (ipcas or '').strip(), int(so_tap) if so_tap.isdigit() and int(so_tap) > 0 else 1))
                if not (ipcas or '').strip():
                    thieu_ipcas.append(name)
        so_bat_dau = request.POST.get('so_bat_dau', '').strip()

        if not ma_cn:
            messages.error(request, f'Chưa cấu hình Mã chi nhánh cho đơn vị {branch} (Cấu hình Chi nhánh).')
        elif thieu_ipcas:
            messages.error(request, f'GDV chưa khai báo Tài khoản IPCAS (QR cần UserIPCAS): {", ".join(thieu_ipcas)}.')
        elif not chon:
            messages.error(request, 'Chưa chọn GDV nào.')
        elif not so_bat_dau.isdigit() or int(so_bat_dau) < 1:
            messages.error(request, 'Chưa nhập Số bắt đầu (hoặc số không hợp lệ).')
        else:
            with transaction.atomic():
                cung_ngay = BiaChungTu.objects.filter(branch=branch, ngay_chung_tu=ngay)
                da_co = cung_ngay.count()
                tong = da_co + sum(n for _, _, n in chon)
                start = int(so_bat_dau)
                so_moi = [_format_so(ma_cn, ngay.year, start + i) for i in range(tong - da_co)]
                trung = list(BiaChungTu.objects.filter(branch=branch, so_ho_so__in=so_moi)
                             .values_list('so_ho_so', flat=True))
                if trung:
                    messages.error(request, f'Số hồ sơ đã tồn tại: {", ".join(trung)}. Hãy chọn số bắt đầu khác.')
                else:
                    cung_ngay.update(tong_tap=tong)
                    tap, idx = da_co, 0
                    for uid, ipcas, n in chon:
                        for _ in range(n):
                            tap += 1
                            BiaChungTu.objects.create(
                                branch=branch, so_ho_so=so_moi[idx], ngay_chung_tu=ngay,
                                tap_so=tap, tong_tap=tong, gdv_id=uid,
                                user_ipcas=ipcas, nguoi_lap=request.user,
                            )
                            idx += 1
                    messages.success(request, f'Đã cấp {len(so_moi)} số hồ sơ: {so_moi[0]} → {so_moi[-1]}.')
                    return redirect(f"{reverse('bia_chung_tu_list')}?tu_ngay={ngay:%d/%m/%Y}&den_ngay={ngay:%d/%m/%Y}")

    return render(request, 'templates_app/bia_chung_tu/cap_so.html', {
        'ngay': ngay,
        'gdv_rows': gdv_rows,
        'so_bat_dau': request.POST.get('so_bat_dau', ''),
        'so_goi_y': _next_so(branch, ma_cn, ngay.year) if ma_cn else '',
        'ma_cn': ma_cn,
        'ipcas': _ipcas(request.user),
        'da_co': BiaChungTu.objects.filter(branch=branch, ngay_chung_tu=ngay).count(),
    })


@login_required
@require_http_methods(['GET', 'POST'])
def bia_edit(request, pk):
    branch = _manager_branch(request.user)
    if not branch:
        return _deny(request)
    bia = get_object_or_404(BiaChungTu, pk=pk, branch=branch)
    gdv_choices = _gdv_choices(branch)
    if bia.gdv_id and bia.gdv_id not in {uid for uid, _ in gdv_choices}:
        gdv_choices.append((bia.gdv_id, getattr(getattr(bia.gdv, 'profile', None), 'full_name', '') or bia.gdv.username))

    if request.method == 'POST':
        p = request.POST
        so_ho_so = p.get('so_ho_so', '').strip()
        ngay = _parse_date(p.get('ngay_chung_tu'))
        ints = {k: p.get(k, '') for k in ('tap_so', 'tong_tap', 'so_to')}
        errors = []
        if not so_ho_so:
            errors.append('Số hồ sơ không được trống.')
        elif BiaChungTu.objects.filter(branch=branch, so_ho_so=so_ho_so).exclude(pk=bia.pk).exists():
            errors.append(f'Số hồ sơ {so_ho_so} đã tồn tại.')
        if not ngay:
            errors.append('Ngày chứng từ không hợp lệ.')
        if not all(v.isdigit() and int(v) > 0 for v in ints.values()):
            errors.append('Tập số, Tổng số tập, Số tờ phải là số nguyên dương.')
        if not p.get('user_ipcas', '').strip():
            errors.append('UserIPCAS không được trống.')
        so_to_tap = p.get('so_to_tap', '').strip()
        if so_to_tap and not (so_to_tap.isdigit() and int(so_to_tap) > 0):
            errors.append('Số tờ của tập phải là số nguyên dương (hoặc để trống).')
        if errors:
            for e in errors:
                messages.error(request, e)
        else:
            bia.so_ho_so = so_ho_so
            bia.ngay_chung_tu = ngay
            bia.tap_so, bia.tong_tap, bia.so_to = (int(ints[k]) for k in ('tap_so', 'tong_tap', 'so_to'))
            bia.gdv_id = int(p['gdv']) if p.get('gdv', '').isdigit() else None
            bia.don_vi = p.get('don_vi', '').strip() or 'Quyển'
            bia.thoi_han_bao_quan = p.get('thoi_han_bao_quan', '').strip() or '20 năm'
            bia.user_ipcas = p['user_ipcas'].strip()
            bia.so_to_tap = int(so_to_tap) if so_to_tap else None
            bia.save()
            messages.success(request, f'Đã cập nhật bìa {bia.so_ho_so}.')
            return redirect('bia_chung_tu_list')

    return render(request, 'templates_app/bia_chung_tu/edit.html', {
        'bia': bia, 'gdv_choices': gdv_choices, 'ipcas': _ipcas(request.user),
    })


@login_required
@require_POST
def bia_delete(request, pk):
    branch = _manager_branch(request.user)
    if not branch:
        return _deny(request)
    bia = get_object_or_404(BiaChungTu, pk=pk, branch=branch)
    so = bia.so_ho_so
    bia.delete()
    messages.success(request, f'Đã xoá bìa {so}.')
    return _back(request)


@login_required
@require_POST
def bia_save_so_to(request):
    """Lưu nhanh Số tờ nhập trực tiếp trên danh sách (input name="so_to_<id>")."""
    branch = _manager_branch(request.user)
    if not branch:
        return _deny(request)
    messages.success(request, f'Đã lưu số tờ ({_apply_so_to(request, branch)} bìa thay đổi).')
    return _back(request)


def _apply_so_to(request, branch):
    """so_to_<id>: Số tờ trên bìa; so_to_tap_<id>: Số tờ của tập ở nhãn QR (trống = "….")."""
    changed = 0
    for key, val in request.POST.items():
        val = val.strip()
        if key.startswith('so_to_tap_') and key[10:].isdigit():
            moi = int(val) if val.isdigit() and int(val) > 0 else None
            changed += BiaChungTu.objects.filter(branch=branch, pk=int(key[10:])).exclude(so_to_tap=moi).update(so_to_tap=moi)
        elif key.startswith('so_to_') and key[6:].isdigit() and val.isdigit() and int(val) > 0:
            changed += BiaChungTu.objects.filter(branch=branch, pk=int(key[6:])).exclude(so_to=int(val)).update(so_to=int(val))
    return changed


# ---------------------------------------------------------------------------
# Import file Excel (in bù bìa cũ)
# ---------------------------------------------------------------------------

def _parse_import_file(f):
    """Đọc sheet đầu: Số hồ sơ | Tiêu đề hồ sơ | Thời hạn bảo quản | Ngày mở | Ngày công việc kết thúc | Số tờ | Ghi chú
    | User | Số tờ của tập chứng từ (2 cột cuối không bắt buộc; mỗi dòng = 1 tập).
    Bỏ qua dòng không có tiêu đề. Trả về list dict (giá trị chuỗi để lưu session)."""
    ws = openpyxl.load_workbook(f, data_only=True, read_only=True).worksheets[0]
    rows = ws.iter_rows(values_only=True)
    header = [str(c or '').strip().lower() for c in next(rows, [])]

    def col(name):
        return header.index(name) if name in header else None

    c_so, c_tieu_de = col('số hồ sơ'), col('tiêu đề hồ sơ')
    if c_so is None or c_tieu_de is None:
        raise ValueError('Không tìm thấy cột "Số hồ sơ" và "Tiêu đề hồ sơ" ở dòng đầu tiên.')
    c_han, c_mo, c_to, c_ghi_chu = col('thời hạn bảo quản'), col('ngày mở'), col('số tờ'), col('ghi chú')
    c_user, c_to_tap = col('user'), col('số tờ của tập chứng từ')

    def get(r, c):
        return r[c] if c is not None and c < len(r) else None

    result = []
    for i, r in enumerate(rows, start=2):
        so = str(get(r, c_so) or '').strip()
        tieu_de = str(get(r, c_tieu_de) or '').strip()
        if not tieu_de:
            continue
        m = TITLE_RE.search(tieu_de)
        ngay = _parse_date(m.group(1)) if m else _parse_date(get(r, c_mo))
        so_to = get(r, c_to)
        so_to_tap = get(r, c_to_tap)  # "…." = chưa có, hậu kiểm điền tay
        item = {
            'dong': i, 'so_ho_so': so, 'tieu_de': tieu_de,
            'ngay': ngay.isoformat() if ngay else '',
            'tap_so': int(m.group(2)) if m else 0, 'tong_tap': int(m.group(3)) if m else 0,
            'so_to': int(so_to) if isinstance(so_to, (int, float)) and so_to > 0 else 1,
            'don_vi': str(get(r, c_ghi_chu) or '').strip() or 'Quyển',
            'thoi_han': str(get(r, c_han) or '').strip() or '20 năm',
            'user': str(get(r, c_user) or '').strip().upper(),
            'so_to_tap': int(so_to_tap) if isinstance(so_to_tap, (int, float)) and so_to_tap > 0 else None,
            'hanh_dong': 'Tạo mới',
            'loi': '',
        }
        if not so:
            item['loi'] = 'Thiếu số hồ sơ'
        elif not m or not ngay:
            item['loi'] = 'Tiêu đề không đúng dạng "Nhật ký chứng từ ngày dd/mm/yyyy - Tập x/y"'
        result.append(item)
    return result


@login_required
@require_http_methods(['GET', 'POST'])
def bia_import(request):
    branch = _manager_branch(request.user)
    if not branch:
        return _deny(request)

    if request.method == 'POST' and request.FILES.get('file'):
        try:
            rows = _parse_import_file(request.FILES['file'])
        except Exception as e:
            messages.error(request, f'Lỗi đọc file: {e}')
            return redirect('bia_chung_tu_import')
        counts = {}
        for r in rows:
            counts[r['so_ho_so']] = counts.get(r['so_ho_so'], 0) + 1
        da_co = set(BiaChungTu.objects.filter(branch=branch, so_ho_so__in=list(counts))
                    .values_list('so_ho_so', flat=True))
        gdv_ipcas = _gdv_ipcas_map(branch)
        for r in rows:
            if r['loi']:
                continue
            if counts[r['so_ho_so']] > 1:
                r['loi'] = 'Số hồ sơ bị trùng trong file'
            elif r['user'] and r['user'] not in gdv_ipcas:
                r['loi'] = f'Không tìm thấy GDV có IPCAS {r["user"]}'
            elif r['so_ho_so'] in da_co:
                # Bìa đã import trước: chỉ cập nhật User + Số tờ của tập (nếu file có), giữ nguyên phần đã sửa tay
                if r['user'] or r['so_to_tap']:
                    r['hanh_dong'] = 'Cập nhật'
                else:
                    r['loi'] = 'Đã có trong hệ thống (bỏ qua)'
        request.session['bia_import_rows'] = rows
        return render(request, 'templates_app/bia_chung_tu/import.html', {
            'rows': rows,
            'so_hop_le': sum(1 for r in rows if not r['loi']),
            'ipcas': _ipcas(request.user),
        })

    return render(request, 'templates_app/bia_chung_tu/import.html', {'ipcas': _ipcas(request.user)})


@login_required
@require_POST
def bia_import_confirm(request):
    branch = _manager_branch(request.user)
    if not branch:
        return _deny(request)
    ipcas = _ipcas(request.user)
    rows = [r for r in request.session.get('bia_import_rows', []) if not r['loi']]
    da_co = {b.so_ho_so: b for b in BiaChungTu.objects.filter(branch=branch, so_ho_so__in=[r['so_ho_so'] for r in rows])}
    # Dòng không có cột User dùng UserIPCAS của người import
    if not ipcas and any(not r['user'] and r['so_ho_so'] not in da_co for r in rows):
        messages.error(request, 'Có dòng không ghi User mà bạn chưa khai báo Tài khoản IPCAS trong hồ sơ nhân viên — QR cần UserIPCAS.')
        return redirect('bia_chung_tu_import')
    request.session.pop('bia_import_rows', None)
    gdv_ipcas = _gdv_ipcas_map(branch)

    moi, cap_nhat = [], []
    for r in rows:
        bia = da_co.get(r['so_ho_so'])
        if bia is None:
            moi.append(BiaChungTu(
                branch=branch, so_ho_so=r['so_ho_so'], ngay_chung_tu=date.fromisoformat(r['ngay']),
                tap_so=r['tap_so'], tong_tap=r['tong_tap'], so_to=r['so_to'], don_vi=r['don_vi'],
                thoi_han_bao_quan=r['thoi_han'], so_to_tap=r['so_to_tap'],
                user_ipcas=r['user'] or ipcas, gdv_id=gdv_ipcas.get(r['user']), nguoi_lap=request.user,
            ))
            continue
        if r['user']:
            bia.user_ipcas, bia.gdv_id = r['user'], gdv_ipcas[r['user']]
        if r['so_to_tap']:
            bia.so_to_tap = r['so_to_tap']
        cap_nhat.append(bia)
    with transaction.atomic():
        BiaChungTu.objects.bulk_create(moi)
        BiaChungTu.objects.bulk_update(cap_nhat, ['user_ipcas', 'gdv', 'so_to_tap'])
    messages.success(request, f'Đã import {len(moi)} bìa mới, cập nhật {len(cap_nhat)} bìa đã có.')
    return redirect('bia_chung_tu_list')


# ---------------------------------------------------------------------------
# In PDF
# ---------------------------------------------------------------------------

def _draw_nhan_qr(c, bia, ma_cn, ten_cn, x0, y_top):
    """Khối nhãn QR như sheet "Mau_in Kho A4" của file xlsb (rộng 130mm, cao 65mm).
    y_top: toạ độ cạnh trên (điểm PDF)."""
    w = 130 * mm
    c1 = 42 * mm                     # cột QR
    c2 = 18 * mm                     # cột nhãn "Tập"
    rows = [7, 15, 8, 12, 7, 8, 8]   # mm: đơn vị | ngày CT | họ tên | chuỗi QR | tập | hậu kiểm | đóng CT
    ys = [y_top]
    for h in rows:
        ys.append(ys[-1] - h * mm)
    xr, xa, xb = x0 + w, x0 + c1, x0 + c1 + c2

    c.setLineWidth(0.6)
    c.rect(x0, ys[-1], w, y_top - ys[-1])
    for i in (1, 5, 6):
        c.line(x0, ys[i], xr, ys[i])
    for i in (2, 3, 4):
        c.line(xa, ys[i], xr, ys[i])
    c.line(xa, ys[1], xa, ys[5])
    c.line(xb, ys[3], xb, ys[5])

    def mid(i, size):
        return (ys[i] + ys[i + 1]) / 2 - size * 0.35

    def center(text, xl, xr_, i, size, dy=0):
        c.setFont(FONT_BOLD, size)
        c.drawCentredString((xl + xr_) / 2, mid(i, size) + dy, text)

    c.setFont(FONT_BOLD, 11)
    c.drawString(x0 + 2 * mm, mid(0, 11), f'{ma_cn}-Agribank chi nhánh {ten_cn}')

    ngay = f'{bia.ngay_chung_tu:%d/%m/%Y}'
    # Số tờ ở nhãn = số tờ của tập chứng từ; chưa có thì để "…." cho hậu kiểm điền tay
    center(f'Chứng từ ngày {ngay} ({bia.so_to_tap or "…."} tờ)', xa, xr, 1, 12, 3 * mm)
    center(f'Số tập chứng từ trong ngày: {bia.tong_tap}', xa, xr, 1, 12, -3 * mm)
    center(_ho_ten(bia), xa, xr, 2, 12)

    qr_text = _qr_text(bia, ma_cn)
    lines = ['']  # chuỗi QR không có dấu cách → ngắt dòng sau dấu "-"
    for part in re.findall(r'[^-]+-?', qr_text):
        if lines[-1] and pdfmetrics.stringWidth(lines[-1] + part, FONT_BOLD, 11) > xr - xb - 4 * mm:
            lines.append(part)
        else:
            lines[-1] += part
    for k, ln in enumerate(lines):
        center(ln, xb, xr, 3, 11, (len(lines) - 1) * 2.2 * mm - k * 4.4 * mm)
    center('Tập', xa, xb, 4, 12)
    center(f'{bia.tap_so}/{bia.tong_tap}', xb, xr, 4, 12)

    c.setFont(FONT_REGULAR, 12)
    for i, label in ((5, 'Người hậu kiểm chứng từ: '), (6, 'Người đóng chứng từ: ')):
        text = label
        while c.stringWidth(text + '.', FONT_REGULAR, 12) < w - 6 * mm:
            text += '.'
        c.drawString(x0 + 2 * mm, mid(i, 12), text)

    size = 40 * mm
    qr = QrCodeWidget(qr_text, barLevel='M', barBorder=1)
    bx1, by1, bx2, by2 = qr.getBounds()
    d = Drawing(size, size, transform=[size / (bx2 - bx1), 0, 0, size / (by2 - by1), 0, 0])
    d.add(qr)
    renderPDF.draw(d, c, x0 + (c1 - size) / 2, (ys[1] + ys[5]) / 2 - size / 2)


def _draw_bia(c, bia, ten_cn_hoa):
    W, _ = A4
    cx = W / 2
    x0, w = 22 * mm, W - 44 * mm
    top, bottom = 272 * mm, 22 * mm

    # Ô "Mẫu: M01/LHS" góc trên phải, ngoài khung
    c.setFont(FONT_REGULAR, 11)
    c.setLineWidth(0.6)
    c.rect(x0 + w - 36 * mm, top + 3 * mm, 36 * mm, 7 * mm)
    c.drawCentredString(x0 + w - 18 * mm, top + 5.2 * mm, 'Mẫu: M01/LHS')

    # Khung viền đôi
    c.setLineWidth(1)
    c.rect(x0, bottom, w, top - bottom)
    c.rect(x0 + 1.4 * mm, bottom + 1.4 * mm, w - 2.8 * mm, top - bottom - 2.8 * mm)

    def line(text, y, size, bold=False):
        c.setFont(FONT_BOLD if bold else FONT_REGULAR, size)
        c.drawCentredString(cx, top - y * mm, text)

    line('NGÂN HÀNG NÔNG NGHIỆP', 11, 12)
    line('VÀ PHÁT TRIỂN NÔNG THÔN VIỆT NAM', 18, 12)
    line(f'CHI NHÁNH {ten_cn_hoa}', 25.5, 11, bold=True)
    line('PHÒNG KẾ TOÁN VÀ NGÂN QUỸ', 32, 11, bold=True)
    c.setLineWidth(0.6)
    c.line(cx - 19.5 * mm, top - 34 * mm, cx + 19.5 * mm, top - 34 * mm)

    ngay = f'{bia.ngay_chung_tu:%d/%m/%Y}'
    line('Ký hiệu thông tin:', 47, 14)
    line('HỒ SƠ', 70, 35)
    line(f'Ngày mở: {ngay}', 85, 20)
    line(f'Nhật ký chứng từ ngày {ngay} – Tập {bia.tap_so}/{bia.tong_tap}', 100, 14)
    line(f'Ngày bắt đầu: {ngay}    Ngày kết thúc: {ngay}', 110, 14)
    line(f'Gồm: {bia.so_to:02d} tờ (01 {bia.don_vi.lower()})', 120, 14)

    # Khối thông tin cuối bìa: trái / phải
    c.setFont(FONT_REGULAR, 14)
    left, right = x0 + 12 * mm, x0 + w * 0.72
    for y, l_text, r_text in (
        (198, 'Phông số :…….……', 'Thời hạn bảo quản'),
        (207, 'Sổ quản lý hồ sơ số :…….', bia.thoi_han_bao_quan),
        (216, f'Hồ sơ số: {bia.so_ho_so}', 'Ngày công việc kết thúc'),
        (225, '', ngay),
    ):
        c.drawString(left, top - y * mm, l_text)
        c.drawCentredString(right, top - y * mm, r_text)


@login_required
@require_http_methods(['GET', 'POST'])
def bia_print(request):
    """In PDF: POST ids=<id>... (in nhiều, tích chọn trên danh sách) hoặc GET ?id=<id> (in 1 bìa)."""
    branch = _manager_branch(request.user)
    if not branch:
        return _deny(request)
    if request.method == 'POST':
        _apply_so_to(request, branch)  # in luôn số tờ vừa nhập trên danh sách, không cần bấm Lưu trước
    ids = request.POST.getlist('ids') if request.method == 'POST' else request.GET.getlist('id')
    bias = list(BiaChungTu.objects.filter(branch=branch, pk__in=[i for i in ids if i.isdigit()])
                .select_related('gdv__profile', 'nguoi_lap__profile').order_by('so_ho_so'))
    if not bias:
        messages.error(request, 'Chưa chọn bìa nào để in.')
        return redirect('bia_chung_tu_list')

    ma_cn = _ma_cn(branch)
    cfg = BranchConfig.objects.filter(branch_code=branch).first()
    ten_cn, ten_cn_hoa = (cfg.ten_chi_nhanh, cfg.ten_chi_nhanh_hoa) if cfg else ('', '')

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    c.setTitle('Bìa chứng từ')
    for bia in bias:
        _draw_bia(c, bia, ten_cn_hoa)
        c.showPage()
        # Trang sau: nhãn QR riêng, căn giữa phía trên trang
        _draw_nhan_qr(c, bia, ma_cn, ten_cn, A4[0] / 2 - 65 * mm, A4[1] - 20 * mm)
        c.showPage()
    c.save()
    BiaChungTu.objects.filter(pk__in=[b.pk for b in bias]).update(ngay_in=timezone.now())

    resp = HttpResponse(buf.getvalue(), content_type='application/pdf')
    resp['Content-Disposition'] = f'inline; filename="bia_chung_tu_{date.today():%Y%m%d}.pdf"'
    return resp
