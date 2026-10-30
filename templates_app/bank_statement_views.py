"""
Views phân tích sao kê ngân hàng: upload, xem kết quả, nhập tên tài khoản, xuất Excel.
"""
# Standard library imports
import os

# Django imports
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.db import models
from django.shortcuts import render, get_object_or_404, redirect
from django.urls import reverse


# ====================
# Bank Statement Analyzer Views
# ====================

@login_required
def bank_statement_upload(request):
    """
    Trang upload file sao kê ngân hàng
    """
    from .models import BankStatement

    if request.method == 'POST' and request.FILES.get('statement_file'):
        try:
            from .bank_statement_parser import BankStatementParser
            from .models import Transaction
            import os
            import uuid
            from django.conf import settings

            uploaded_file = request.FILES['statement_file']
            original_name = uploaded_file.name
            ext = os.path.splitext(original_name)[1].lower()

            # Tạo thư mục upload nếu chưa có
            upload_dir = os.path.join(settings.MEDIA_ROOT, 'bank_statements')
            os.makedirs(upload_dir, exist_ok=True)

            # Dùng UUID để tránh trùng tên file giữa các user
            unique_name = f"{uuid.uuid4().hex}{ext}"
            file_path = os.path.join(upload_dir, unique_name)

            with open(file_path, 'wb+') as destination:
                for chunk in uploaded_file.chunks():
                    destination.write(chunk)

            # File đối soát CSP / MSPH02 / FX tùy chọn
            from .bank_statement_enricher import AUX_FILES, enrich
            aux_paths = {}
            for kind in AUX_FILES:
                aux_file = request.FILES.get(f'{kind}_file')
                if aux_file:
                    aux_ext = os.path.splitext(aux_file.name)[1].lower()
                    aux_paths[kind] = os.path.join(upload_dir, f"{kind}_{uuid.uuid4().hex}{aux_ext}")
                    with open(aux_paths[kind], 'wb+') as f:
                        for chunk in aux_file.chunks():
                            f.write(chunk)

            try:
                # Parse file
                parser = BankStatementParser(file_path)

                is_valid, error_msg = parser.validate_file()
                if not is_valid:
                    messages.error(request, f'File không hợp lệ: {error_msg}')
                    return redirect('bank_statement_upload')

                enrichment, unmatched_rows = enrich(parser.df, aux_paths)
                from .models import AccountName
                name_book = {(a.account_key, a.bank_key): a.name for a in AccountName.objects.all()}
                transactions_data = parser.process(enrichment=enrichment, name_book=name_book)
                summary = parser.get_summary()

                # Tạo BankStatement
                statement = BankStatement.objects.create(
                    file_name=original_name,
                    total_transactions=summary['total_transactions'],
                    total_debit=summary['total_debit'],
                    total_credit=summary['total_credit'],
                    final_balance=summary['final_balance'],
                    processed=True,
                    unmatched_rows=unmatched_rows,
                    uploaded_by=request.user
                )

                # Lưu transactions bằng bulk_create (1 query thay vì N queries)
                Transaction.objects.bulk_create([
                    Transaction(
                        statement=statement,
                        stt=t['stt'],
                        transaction_date=t['ngay_giao_dich'],
                        debit_amount=t['so_tien_ghi_no'],
                        credit_amount=t['so_tien_ghi_co'],
                        balance=t['so_du_sau_gd'],
                        bank_name=t['ngan_hang'],
                        account_number=t['so_tai_khoan'],
                        beneficiary_name=t['ten_nguoi'],
                        description=t['noi_dung'],
                        transaction_type=t['ghi_chu'],
                        source=t['nguon'],
                        raw_trcdnm=t['raw_trcdnm'],
                        raw_tomgntno=t['raw_tomgntno'],
                    )
                    for t in transactions_data
                ])

                msg = f'Đã phân tích thành công {summary["total_transactions"]} giao dịch!'
                if aux_paths:
                    msg += f' Bổ sung thông tin đối tác cho {len(enrichment)} giao dịch từ file đối soát'
                    msg += f', {len(unmatched_rows)} dòng không khớp.' if unmatched_rows else '.'
                messages.success(request, msg)
                return redirect('bank_statement_result', statement_id=statement.id)

            finally:
                # Luôn xóa file tạm sau khi xử lý xong (dù thành công hay lỗi)
                if os.path.exists(file_path):
                    os.remove(file_path)
                for aux_path in aux_paths.values():
                    if os.path.exists(aux_path):
                        os.remove(aux_path)

        except Exception as e:
            messages.error(request, f'Lỗi khi xử lý file: {str(e)}')
            return redirect('bank_statement_upload')

    # Lấy danh sách các statement đã upload
    statements = BankStatement.objects.filter(uploaded_by=request.user).order_by('-uploaded_at')[:10]

    context = {
        'statements': statements,
    }
    return render(request, 'templates_app/bank_statement_upload.html', context)


# Cụm từ ngân hàng không phải tên người — dùng chung cho result và export
_NON_NAME_PHRASES = [
    'CHUYEN KHOAN', 'CHUYEN TIEN', 'THANH TOAN', 'GUI TIEN',
    'NAP TIEN', 'RUT TIEN', 'TRA GOP', 'TRA NO', 'TIEN DIEN',
    'TIEN NUOC', 'HOA DON', 'PHI DICH VU', 'PHI SMS',
    'NGAN HANG', 'TAI KHOAN', 'SO TK',
    'CHUYENKHOAN', 'THANHTOAN', 'DICH VU',
]


def _is_proper_name(name):
    """
    Trả True nếu tên trông như tên người thật:
    - > 70% ký tự chữ cái là CHỮ HOA
    - Không chứa chữ số (loại mã giao dịch)
    - Có ít nhất 2 từ (loại tên đơn như LE, TRAN, MOMO)
    - Không chứa cụm từ ngân hàng
    """
    if not name:
        return False
    name = name.strip()
    alpha_chars = [c for c in name if c.isalpha()]
    if len(alpha_chars) < 2:
        return False
    if any(c.isdigit() for c in name):
        return False
    upper_ratio = sum(1 for c in alpha_chars if c.isupper()) / len(alpha_chars)
    if upper_ratio <= 0.7:
        return False
    if len(name.split()) < 2:
        return False
    name_upper = name.upper()
    if any(phrase in name_upper for phrase in _NON_NAME_PHRASES):
        return False
    return True


def _frequent_accounts(statement, amount_field):
    """
    Gom giao dịch theo (số TK, ngân hàng) cho chiều tiền ra (debit_amount) hoặc vào (credit_amount).
    Tên ưu tiên lấy từ dòng có nguồn đã xác thực (file đối soát, danh bạ — không lọc);
    còn lại tên parse từ rem phải qua _is_proper_name.
    """
    accounts = list(
        statement.transactions.filter(**{f'{amount_field}__gt': 0})
        .exclude(account_number='')
        .values('account_number', 'bank_name')
        .annotate(
            count=models.Count('id'),
            total_amount=models.Sum(amount_field),
            name=models.Max('beneficiary_name'),
            aux_name=models.Max('beneficiary_name',
                                filter=~models.Q(source='') & ~models.Q(source__contains='chưa xác thực')),
        )
        .order_by('-count')
    )
    for acc in accounts:
        name = acc.get('name', '') or ''
        if acc['aux_name']:
            name = acc['aux_name']
        elif not _is_proper_name(name):
            name = ''
        acc['name'] = name
    return accounts


@login_required
def bank_statement_result(request, statement_id):
    """
    Hiển thị kết quả phân tích sao kê
    """
    from .models import BankStatement
    from django.core.paginator import Paginator

    statement = get_object_or_404(BankStatement, id=statement_id, uploaded_by=request.user)

    # Lấy tham số filter
    transaction_type_filter = request.GET.get('type', '')

    # Query transactions
    transactions = statement.transactions.all()

    # Apply filter
    if transaction_type_filter:
        transactions = transactions.filter(transaction_type=transaction_type_filter)

    # Pagination
    paginator = Paginator(transactions, 50)  # 50 giao dịch mỗi trang
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)

    # Thống kê theo loại giao dịch
    transaction_types = statement.transactions.exclude(transaction_type='').values('transaction_type').annotate(
        count=models.Count('id'),
        total_debit=models.Sum('debit_amount'),
        total_credit=models.Sum('credit_amount')
    ).order_by('transaction_type')

    # Tài khoản nhận tiền nhiều lần (chuyển đi) / chuyển tiền đến nhiều lần (nhận về)
    frequent_recipients = _frequent_accounts(statement, 'debit_amount')
    frequent_senders = _frequent_accounts(statement, 'credit_amount')

    # Danh sách các loại giao dịch để filter — phải khớp chính xác với classify_transaction()
    filter_options = [
        "Mở tài khoản",
        "Chuyển khoản nội bộ Agribank",
        "Nhận chuyển khoản nội bộ Agribank",
        "Nhận chuyển khoản từ ATM",
        "Chuyển khoản đi khác ngân hàng",
        "Nhận chuyển khoản liên ngân hàng",
        "Thanh toán qua MCC",
        "Nhận thanh toán MCC",
        "Rút tiền ATM",
        "Rút tiền ATM cùng hệ thống",
        "Phí rút tiền ATM",
        "Phí rút tiền ATM cùng hệ thống",
        "Phí chuyển khoản ATM",
        "Hủy rút tiền ATM cùng hệ thống",
        "Hoàn phí rút tiền ATM cùng hệ thống",
        "Rút tiền ATM khác hệ thống",
        "Rút tiền mặt",
        "Rút tiền mặt cùng hệ thống",
        "Phí rút tiền mặt cùng hệ thống",
        "Phí rút tiền ATM khác hệ thống",
        "Nộp tiền tại quầy",
        "Nộp tiền qua ATM",
        "Nộp tiền tại Agribank",
        "Nộp tiền mặt",
        "Thanh toán thẻ",
        "Thanh toán POS",
        "Nạp tiền điện thoại",
        "Thanh toán tiền điện",
        "Thanh toán hóa đơn",
        "Thanh toán dịch vụ",
        "Thanh toán dịch vụ (VNPT)",
        "Phí dịch vụ",
        "Nộp thuế / NSNN",
        "Trả lãi tiền gửi",
        "Trả lãi tiền gửi hàng tháng",
        "Trả lãi tiền gửi hằng tháng",
        "Giải ngân",
        "Giao dịch OSB",
    ]

    context = {
        'statement': statement,
        'page_obj': page_obj,
        'transaction_types': transaction_types,
        'filter_options': filter_options,
        'selected_type': transaction_type_filter,
        'frequent_recipients': frequent_recipients,
        'frequent_senders': frequent_senders,
    }

    return render(request, 'templates_app/bank_statement_result.html', context)


@login_required
def bank_statement_import_names(request, statement_id):
    """
    Nhận lại file Excel đã xuất (người dùng điền tên ở sheet 'TK giao dịch nhiều lần'):
    lưu tên mới/đã sửa vào danh bạ dùng chung và cập nhật giao dịch của sao kê này.
    """
    from .models import AccountName, BankStatement, Transaction
    from .bank_statement_parser import account_key, name_book_source
    from openpyxl import load_workbook

    statement = get_object_or_404(BankStatement, id=statement_id, uploaded_by=request.user)
    result_url = reverse('bank_statement_result', args=[statement.id])
    names_file = request.FILES.get('names_file')
    if request.method != 'POST' or not names_file:
        messages.error(request, 'Vui lòng chọn file Excel đã điền tên.')
        return redirect(result_url)

    try:
        wb = load_workbook(names_file, read_only=True, data_only=True)
        rows = list(wb['TK giao dịch nhiều lần'].iter_rows(values_only=True))
        wb.close()
    except Exception:
        messages.error(request, "File không đúng: cần file Excel xuất từ chương trình "
                                "(có sheet 'TK giao dịch nhiều lần').")
        return redirect(result_url)

    # Tên đang hiển thị hiện tại — chỉ lưu những tên người dùng mới điền hoặc đã sửa
    current = {}
    for amount_field in ('debit_amount', 'credit_amount'):
        for acc in _frequent_accounts(statement, amount_field):
            current[account_key(acc['account_number'], acc['bank_name'])] = acc['name']

    # Dòng dữ liệu: cột A = số thứ tự, B = Ngân hàng, C = Số TK, D = Tên (bỏ qua dòng tiêu đề)
    entries = {}
    for row in rows:
        if len(row) < 4 or not isinstance(row[0], (int, float)):
            continue
        bank, acct, name = (str(v or '').strip() for v in row[1:4])
        if not acct or not name:
            continue
        key = account_key(acct, bank)
        if current.get(key, '') != name:
            entries[key] = (acct, bank, name)

    if not entries:
        messages.warning(request, 'Không tìm thấy tên mới nào trong file.')
        return redirect(result_url)

    for (acct_key, bank_key), (acct, bank, name) in entries.items():
        AccountName.objects.update_or_create(
            account_key=acct_key, bank_key=bank_key,
            defaults={'account_number': acct, 'bank_name': bank, 'name': name, 'updated_by': request.user},
        )

    updated = []
    for trans in statement.transactions.exclude(account_number=''):
        entry = entries.get(account_key(trans.account_number, trans.bank_name))
        if entry and trans.beneficiary_name != entry[2]:
            trans.beneficiary_name = entry[2]
            trans.source = name_book_source(trans.source)
            updated.append(trans)
    Transaction.objects.bulk_update(updated, ['beneficiary_name', 'source'])

    messages.success(request, f'Đã lưu {len(entries)} tên vào danh bạ, cập nhật {len(updated)} giao dịch.')
    return redirect(result_url)


@login_required
def bank_statement_huong_dan(request):
    """Tải hướng dẫn dạng Word (bản .docx sinh từ docs/HUONG_DAN_PHAN_TICH_SAO_KE.md)."""
    from django.conf import settings
    from django.http import FileResponse
    path = os.path.join(settings.BASE_DIR, 'docs', 'HUONG_DAN_PHAN_TICH_SAO_KE.docx')
    return FileResponse(open(path, 'rb'), as_attachment=True, filename='Huong_dan_Phan_tich_sao_ke.docx')


@login_required
def bank_statement_export(request, statement_id):
    """
    Export báo cáo sao kê ra file Excel
    """
    from .models import BankStatement
    from io import BytesIO
    import xlsxwriter
    from django.http import HttpResponse

    statement = get_object_or_404(BankStatement, id=statement_id, uploaded_by=request.user)

    # Tạo file Excel trong memory
    output = BytesIO()

    workbook = xlsxwriter.Workbook(output)

    # Format cho số tiền
    money_format = workbook.add_format({
        'num_format': '#,##0',
        'align': 'right'
    })

    # Format cho header
    header_format = workbook.add_format({
        'bold': True,
        'bg_color': '#4472C4',
        'font_color': 'white',
        'align': 'center',
        'valign': 'vcenter',
        'border': 1
    })

    # Format cho ngày
    date_format = workbook.add_format({
        'num_format': 'dd/mm/yyyy',
        'align': 'center'
    })

    # Sheet 1: Chi tiết giao dịch
    worksheet1 = workbook.add_worksheet('Chi tiết')

    # Header
    headers = [
        'STT', 'Ngày GD', 'Số tiền ghi nợ', 'Số tiền ghi có',
        'Số dư sau GD', 'Ngân hàng', 'Số TK', 'Tên người',
        'Nội dung', 'Loại giao dịch', 'Nguồn'
    ]

    for col_num, header in enumerate(headers):
        worksheet1.write(0, col_num, header, header_format)

    # Set column widths
    worksheet1.set_column('A:A', 8)   # STT
    worksheet1.set_column('B:B', 12)  # Ngày
    worksheet1.set_column('C:E', 15)  # Các cột tiền
    worksheet1.set_column('F:F', 15)  # Ngân hàng
    worksheet1.set_column('G:G', 18)  # Số TK
    worksheet1.set_column('H:H', 25)  # Tên người
    worksheet1.set_column('I:I', 40)  # Nội dung
    worksheet1.set_column('J:J', 30)  # Loại GD
    worksheet1.set_column('K:K', 25)  # Nguồn

    # Bảng TK giao dịch nhiều lần (sheet 3) — tính trước để sheet 1 tham chiếu bằng công thức
    frequent_recipients_export = _frequent_accounts(statement, 'debit_amount')
    frequent_senders_export = _frequent_accounts(statement, 'credit_amount')
    start_row = len(frequent_recipients_export) + 4  # dòng tiêu đề bảng 2 (0-based)

    # Vùng dữ liệu (Excel, 1-based) của 2 bảng ở sheet 3: cột B=Ngân hàng, C=Số TK, D=Tên
    sheet3 = "'TK giao dịch nhiều lần'"
    rec_rows = (3, 2 + len(frequent_recipients_export))
    send_rows = (start_row + 3, start_row + 2 + len(frequent_senders_export))

    def name_lookup_formula(excel_row, rows):
        # Tên điền tay ở sheet 3 → tự hiện ở cột Tên người (khớp Ngân hàng + Số TK)
        first, last = rows
        bank = f'{sheet3}!$B${first}:$B${last}'
        acct = f'{sheet3}!$C${first}:$C${last}'
        name = f'{sheet3}!$D${first}:$D${last}'
        return (f'=IFERROR(LOOKUP(2,1/(({bank}=F{excel_row})*({acct}=G{excel_row})*({name}<>"")),{name}),"")')

    # Data
    transactions = statement.transactions.all()
    for row_num, trans in enumerate(transactions, start=1):
        worksheet1.write(row_num, 0, trans.stt)
        worksheet1.write(row_num, 1, trans.transaction_date, date_format)
        worksheet1.write(row_num, 2, float(trans.debit_amount), money_format)
        worksheet1.write(row_num, 3, float(trans.credit_amount), money_format)
        worksheet1.write(row_num, 4, float(trans.balance), money_format)
        worksheet1.write(row_num, 5, trans.bank_name)
        worksheet1.write(row_num, 6, trans.account_number)
        rows = rec_rows if trans.debit_amount > 0 else send_rows
        if not trans.beneficiary_name and trans.account_number and rows[1] >= rows[0]:
            worksheet1.write_formula(row_num, 7, name_lookup_formula(row_num + 1, rows), None, '')
        else:
            worksheet1.write(row_num, 7, trans.beneficiary_name)
        worksheet1.write(row_num, 8, trans.description)
        worksheet1.write(row_num, 9, trans.transaction_type)
        worksheet1.write(row_num, 10, trans.source or 'Nội dung GD')

    # Sheet 2: Thống kê
    worksheet2 = workbook.add_worksheet('Thống kê')

    # Thống kê tổng quan
    worksheet2.write(0, 0, 'THỐNG KÊ TỔNG QUAN', header_format)
    worksheet2.write(1, 0, 'Tên file:')
    worksheet2.write(1, 1, statement.file_name)
    worksheet2.write(2, 0, 'Ngày upload:')
    worksheet2.write(2, 1, statement.uploaded_at.strftime('%d/%m/%Y %H:%M'))
    worksheet2.write(3, 0, 'Tổng số giao dịch:')
    worksheet2.write(3, 1, statement.total_transactions)
    worksheet2.write(4, 0, 'Tổng tiền ghi nợ:')
    worksheet2.write(4, 1, float(statement.total_debit), money_format)
    worksheet2.write(5, 0, 'Tổng tiền ghi có:')
    worksheet2.write(5, 1, float(statement.total_credit), money_format)
    worksheet2.write(6, 0, 'Số dư cuối kỳ:')
    worksheet2.write(6, 1, float(statement.final_balance), money_format)

    # Thống kê theo loại giao dịch
    worksheet2.write(8, 0, 'THỐNG KÊ THEO LOẠI GIAO DỊCH', header_format)
    worksheet2.write(9, 0, 'Loại giao dịch', header_format)
    worksheet2.write(9, 1, 'Số lượng', header_format)
    worksheet2.write(9, 2, 'Tổng ghi nợ', header_format)
    worksheet2.write(9, 3, 'Tổng ghi có', header_format)

    transaction_types = statement.transactions.exclude(transaction_type='').values('transaction_type').annotate(
        count=models.Count('id'),
        total_debit=models.Sum('debit_amount'),
        total_credit=models.Sum('credit_amount')
    ).order_by('transaction_type')

    for row_num, item in enumerate(transaction_types, start=10):
        worksheet2.write(row_num, 0, item['transaction_type'])
        worksheet2.write(row_num, 1, item['count'])
        worksheet2.write(row_num, 2, float(item['total_debit'] or 0), money_format)
        worksheet2.write(row_num, 3, float(item['total_credit'] or 0), money_format)

    # Thống kê theo nguồn thông tin đối tác (gộp các biến thể "CSP chiều đi (tên từ ...)")
    by_source = {}
    for item in statement.transactions.values('source').annotate(
        count=models.Count('id'),
        total_debit=models.Sum('debit_amount'),
        total_credit=models.Sum('credit_amount'),
    ):
        label = (item['source'] or 'Nội dung GD').split(' (')[0]
        agg = by_source.setdefault(label, [0, 0, 0])
        agg[0] += item['count']
        agg[1] += float(item['total_debit'] or 0)
        agg[2] += float(item['total_credit'] or 0)

    src_row = 10 + len(transaction_types) + 2
    worksheet2.write(src_row, 0, 'THỐNG KÊ THEO NGUỒN', header_format)
    for col, h in enumerate(['Nguồn', 'Số lượng', 'Tổng ghi nợ', 'Tổng ghi có']):
        worksheet2.write(src_row + 1, col, h, header_format)
    for row_num, (label, (count, debit, credit)) in enumerate(sorted(by_source.items()), start=src_row + 2):
        worksheet2.write(row_num, 0, label)
        worksheet2.write(row_num, 1, count)
        worksheet2.write(row_num, 2, debit, money_format)
        worksheet2.write(row_num, 3, credit, money_format)
    note_row = src_row + 2 + len(by_source)
    worksheet2.write(note_row, 0, 'Dòng file đối soát không khớp DPTB18:')
    worksheet2.write(note_row, 1, len(statement.unmatched_rows))

    worksheet2.set_column('A:A', 40)
    worksheet2.set_column('B:D', 15)

    # Sheet 3: Tài khoản giao dịch nhiều lần
    worksheet3 = workbook.add_worksheet('TK giao dịch nhiều lần')

    header_red = workbook.add_format({
        'bold': True, 'bg_color': '#C00000', 'font_color': 'white',
        'align': 'center', 'valign': 'vcenter', 'border': 1
    })
    # Ô tên còn trống: tô vàng để người dùng điền tay (sheet Chi tiết tự cập nhật)
    fill_name_format = workbook.add_format({'bg_color': '#FFF2CC'})
    header_green = workbook.add_format({
        'bold': True, 'bg_color': '#375623', 'font_color': 'white',
        'align': 'center', 'valign': 'vcenter', 'border': 1
    })

    # --- Bảng 1: Tài khoản nhận tiền nhiều lần (tiền ra) ---
    worksheet3.write(0, 0, 'TÀI KHOẢN NHẬN TIỀN NHIỀU LẦN (tất cả tài khoản)', header_red)
    worksheet3.merge_range(0, 0, 0, 5, 'TÀI KHOẢN NHẬN TIỀN NHIỀU LẦN (tất cả tài khoản)', header_red)

    rec_headers = ['#', 'Ngân hàng', 'Số tài khoản', 'Tên', 'Số lần', 'Tổng tiền ra']
    for col, h in enumerate(rec_headers):
        worksheet3.write(1, col, h, header_format)


    for i, acc in enumerate(frequent_recipients_export, start=1):
        row = i + 1
        acct = acc.get('account_number', '')
        display_name = acc['name']
        worksheet3.write(row, 0, i)
        worksheet3.write(row, 1, acc['bank_name'] or '')
        worksheet3.write(row, 2, acct)
        worksheet3.write(row, 3, display_name, None if display_name else fill_name_format)
        worksheet3.write(row, 4, acc['count'])
        worksheet3.write(row, 5, float(acc['total_amount'] or 0), money_format)

    # --- Bảng 2: Tài khoản chuyển tiền đến nhiều lần (tiền vào) ---

    worksheet3.merge_range(start_row, 0, start_row, 5, 'TÀI KHOẢN CHUYỂN TIỀN ĐẾN NHIỀU LẦN (tất cả tài khoản)', header_green)

    send_headers = ['#', 'Ngân hàng', 'Số tài khoản', 'Tên', 'Số lần', 'Tổng tiền vào']
    for col, h in enumerate(send_headers):
        worksheet3.write(start_row + 1, col, h, header_format)

    for i, acc in enumerate(frequent_senders_export, start=1):
        row = start_row + 1 + i
        acct = acc.get('account_number', '')
        display_name = acc['name']
        worksheet3.write(row, 0, i)
        worksheet3.write(row, 1, acc['bank_name'] or '')
        worksheet3.write(row, 2, acct)
        worksheet3.write(row, 3, display_name, None if display_name else fill_name_format)
        worksheet3.write(row, 4, acc['count'])
        worksheet3.write(row, 5, float(acc['total_amount'] or 0), money_format)

    worksheet3.set_column('A:A', 6)
    worksheet3.set_column('B:B', 18)
    worksheet3.set_column('C:C', 20)
    worksheet3.set_column('D:D', 28)
    worksheet3.set_column('E:E', 10)
    worksheet3.set_column('F:F', 18)

    # Sheet 4: Dòng file đối soát không khớp DPTB18
    if statement.unmatched_rows:
        worksheet4 = workbook.add_worksheet('Không khớp')
        unmatched_headers = ['#', 'File', 'Thời gian', 'Số tiền', 'Trace', 'Số TK đối tác', 'Tên đối tác', 'Nội dung']
        for col, h in enumerate(unmatched_headers):
            worksheet4.write(0, col, h, header_format)
        for i, item in enumerate(statement.unmatched_rows, start=1):
            worksheet4.write(i, 0, i)
            worksheet4.write(i, 1, item.get('file', ''))
            worksheet4.write(i, 2, item.get('time', ''))
            worksheet4.write(i, 3, float(item.get('amount') or 0), money_format)
            worksheet4.write(i, 4, item.get('trace', ''))
            worksheet4.write(i, 5, item.get('account', ''))
            worksheet4.write(i, 6, item.get('name', ''))
            worksheet4.write(i, 7, item.get('content', ''))
        worksheet4.set_column('A:A', 6)
        worksheet4.set_column('B:C', 20)
        worksheet4.set_column('D:E', 14)
        worksheet4.set_column('F:G', 22)
        worksheet4.set_column('H:H', 60)

    workbook.close()

    # Chuẩn bị response
    output.seek(0)
    response = HttpResponse(
        output.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="BaoCao_SaoKe_{statement.id}.xlsx"'

    return response
