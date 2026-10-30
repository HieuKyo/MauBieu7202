"""
Báo cáo Đối chiếu huy động vốn (File Chương trình vs File Hệ thống IPCAS) và các hàm chuẩn hóa dữ liệu HĐV dùng chung.
"""
from django.shortcuts import render, redirect
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.views.decorators.http import require_http_methods
from django.contrib import messages
import math
import pandas as pd
import io
import re
import traceback
from openpyxl.utils import get_column_letter
from datetime import datetime
from .report_views import _find_col


# ---------------------------------------------------------------------------
# Báo cáo Đối chiếu huy động vốn (File Chương trình vs File Hệ thống IPCAS)
# ---------------------------------------------------------------------------

def _hdv_clean_str(val):
    """Xóa dấu nháy đơn, khoảng trắng thừa, đuôi '.0' (khi Excel lưu mã dạng số
    khiến pandas đọc thành số thực trước khi ép chuỗi, vd '123456.0'). Trả '' nếu rỗng/NaN."""
    if pd.isna(val):
        return ''
    s = str(val).replace("'", '').strip()
    return s[:-2] if s.endswith('.0') else s


def _hdv_remove_leading_zeros(val):
    return _hdv_clean_str(val).lstrip('0')


def _hdv_to_float(val):
    try:
        f = float(val)
        return 0.0 if math.isnan(f) else f
    except (TypeError, ValueError):
        pass
    # Một số file IPCAS lưu số dạng chuỗi có đơn vị/ký tự kèm theo (vd '01 Tháng', '1,0')
    m = re.search(r'-?\d+(?:[.,]\d+)?', _hdv_clean_str(val))
    if not m:
        return 0.0
    return float(m.group(0).replace(',', '.'))


def _hdv_format_date_dmy(val):
    """Chuyển các định dạng ngày thường gặp ('2026-06-23 00:00:00', '20261225')
    sang dd/mm/yyyy. Giữ nguyên chuỗi gốc nếu không nhận diện được."""
    s = _hdv_clean_str(val)
    if not s:
        return s
    date_part = s.split(' ')[0]
    for fmt in ('%Y-%m-%d', '%Y%m%d', '%d/%m/%Y'):
        try:
            return datetime.strptime(date_part, fmt).strftime('%d/%m/%Y')
        except ValueError:
            continue
    return s


def _hdv_join_unique(values):
    """Nối các giá trị khác rỗng, không trùng lặp, theo thứ tự xuất hiện."""
    seen = []
    for v in values:
        if v and v not in seen:
            seen.append(v)
    return '; '.join(seen)


# Theo văn bản 329/NHNo.CT-KH&QLRR ngày 05/6/2026 (mục II.1):
# Số tiền thưởng = tỷ lệ thưởng x Σ số tiền gửi mới/tăng thêm phát sinh trong tháng,
# làm tròn đến 10.000đ, tối đa 5.000.000đ/cán bộ/tháng.
# Chưa có dữ liệu phân loại lãi suất niêm yết (0,1%) / chính sách (0,05%) trong file nguồn
# nên tạm áp dụng cố định mức niêm yết 0,1%.
HDV_TY_LE_THUONG = 0.001
HDV_TRAN_THUONG_THANG = 5_000_000


def _hdv_round_10k(amount):
    return round(amount / 10_000) * 10_000


@login_required
def huy_dong_von_report_view(request):
    """Giao diện báo cáo Đối chiếu huy động vốn"""
    return render(request, 'templates_app/reports/huy_dong_von.html')


@login_required
@require_http_methods(["POST"])
def process_huy_dong_von_report(request):
    """Đối chiếu file Chương trình (đăng ký HĐV) với file Hệ thống IPCAS."""
    try:
        file_ct = request.FILES.get('file_chuong_trinh')
        file_ht = request.FILES.get('file_he_thong')
        if not all([file_ct, file_ht]):
            messages.error(request, "Vui lòng tải lên đủ 2 file.")
            return redirect('huy_dong_von_report')

        try:
            df_ct = pd.read_excel(file_ct, dtype=str)
            df_ht = pd.read_excel(file_ht, dtype=str)
        except Exception as e:
            messages.error(request, f"Không đọc được file Excel: {e}")
            return redirect('huy_dong_von_report')

        required_ht = ['ID_NUMBER', 'EMPLOYEE_NUMBER', 'EMPLOYEE_NAME', 'CCY',
                       'CURRENT_BALANCE', 'OPENING_DATE', 'MATURITY_DATE', 'MONTH_TERM', 'RATE',
                       'SO_TAI_KHOAN', 'ACCOUNT_STATUS', 'MA_KH', 'MA_CN']
        missing_ht = [c for c in required_ht if _find_col(df_ht.columns, [c]) is None]
        if missing_ht:
            messages.error(request, f"File Hệ thống IPCAS thiếu cột: {', '.join(missing_ht)}")
            return redirect('huy_dong_von_report')

        col_id     = _find_col(df_ht.columns, ['ID_NUMBER'])
        col_emp    = _find_col(df_ht.columns, ['EMPLOYEE_NUMBER'])
        col_ename  = _find_col(df_ht.columns, ['EMPLOYEE_NAME'])
        col_ccy    = _find_col(df_ht.columns, ['CCY'])
        col_bal    = _find_col(df_ht.columns, ['CURRENT_BALANCE'])
        col_open   = _find_col(df_ht.columns, ['OPENING_DATE'])
        col_mat    = _find_col(df_ht.columns, ['MATURITY_DATE'])
        col_term   = _find_col(df_ht.columns, ['MONTH_TERM'])
        col_rate   = _find_col(df_ht.columns, ['RATE'])
        col_sotk   = _find_col(df_ht.columns, ['SO_TAI_KHOAN'])
        col_status = _find_col(df_ht.columns, ['ACCOUNT_STATUS'])
        col_makh   = _find_col(df_ht.columns, ['MA_KH'])
        col_ma_cn  = _find_col(df_ht.columns, ['MA_CN'])
        col_ngt    = _find_col(df_ht.columns, ['NGUOI_GIOI_THIEU'])
        col_ten_ngt = _find_col(df_ht.columns, ['TEN_NGUOI_GIOI_THIEU'])

        required_ct = {
            'CCCD/GPĐKKD/GCNĐT/Mã số DN/MST': None,
            'Mã cán bộ': None,
            'Số tiền': None,
        }
        for name in required_ct:
            required_ct[name] = _find_col(df_ct.columns, [name])
        missing_ct = [name for name, col in required_ct.items() if col is None]
        if missing_ct:
            messages.error(request, f"File Chương trình thiếu cột: {', '.join(missing_ct)}")
            return redirect('huy_dong_von_report')

        col_cccd   = required_ct['CCCD/GPĐKKD/GCNĐT/Mã số DN/MST']
        col_emp_ct = required_ct['Mã cán bộ']
        col_amount = required_ct['Số tiền']

        col_ngay_dk = _find_col(df_ct.columns, ['Ngày ĐK huy động'])
        if col_ngay_dk:
            df_ct[col_ngay_dk] = df_ct[col_ngay_dk].apply(_hdv_format_date_dmy)

        # Gộp số dư IPCAS hợp lệ (VND, kỳ hạn >= 1 tháng theo MONTH_TERM) theo CCCD_Mã cán bộ
        # + cccd_employees: CCCD -> tập mã cán bộ đã từng được gán cho KH này (kể cả sổ đã tất toán)
        # + cccd_balance:   CCCD -> tổng số dư các sổ đang Normal (dùng khi KH tất toán sổ gốc rồi gửi lại
        #                   sổ khác chưa kịp gán mã — vẫn tính chỉ tiêu cho cán bộ đã được xác nhận ở trên)
        sys_agg = {}
        cccd_employees = {}
        cccd_balance = {}
        vnd_count = 0
        term_ok_count = 0
        for _, row in df_ht.iterrows():
            if _hdv_clean_str(row[col_ccy]).upper() != 'VND':
                continue
            vnd_count += 1
            if _hdv_to_float(row[col_term]) < 1:
                continue
            term_ok_count += 1

            cccd = _hdv_remove_leading_zeros(row[col_id])
            emp  = _hdv_clean_str(row[col_emp])
            detail = {
                'employee_name': _hdv_clean_str(row[col_ename]),
                'rate': _hdv_clean_str(row[col_rate]),
                'opening_date': _hdv_format_date_dmy(row[col_open]),
                'maturity_date': _hdv_format_date_dmy(row[col_mat]),
                'month_term': _hdv_clean_str(row[col_term]),
                'so_tai_khoan': _hdv_clean_str(row[col_sotk]),
                'ma_kh': (
                    f"{_hdv_clean_str(row[col_ma_cn])}-{_hdv_clean_str(row[col_makh])}"
                    if _hdv_clean_str(row[col_makh]) else ''
                ),
                'ma_can_bo_gt': _hdv_clean_str(row[col_ngt]) if col_ngt else '',
                'ten_can_bo_gt': _hdv_clean_str(row[col_ten_ngt]) if col_ten_ngt else '',
            }
            balance = _hdv_to_float(row[col_bal])

            key = f"{cccd}_{emp}"
            entry = sys_agg.setdefault(key, {
                'balance': 0.0, 'employee_names': [], 'rates': [], 'opening_dates': [],
                'maturity_dates': [], 'month_terms': [], 'so_tai_khoans': [], 'ma_khs': [],
                'ma_can_bo_gts': [], 'ten_can_bo_gts': [],
            })
            entry['balance'] += balance
            entry['employee_names'].append(detail['employee_name'])
            entry['rates'].append(detail['rate'])
            entry['opening_dates'].append(detail['opening_date'])
            entry['maturity_dates'].append(detail['maturity_date'])
            entry['month_terms'].append(detail['month_term'])
            entry['so_tai_khoans'].append(detail['so_tai_khoan'])
            entry['ma_khs'].append(detail['ma_kh'])
            entry['ma_can_bo_gts'].append(detail['ma_can_bo_gt'])
            entry['ten_can_bo_gts'].append(detail['ten_can_bo_gt'])

            if emp:
                cccd_employees.setdefault(cccd, set()).add(emp)

            if _hdv_clean_str(row[col_status]).lower() == 'normal':
                b_entry = cccd_balance.setdefault(cccd, {
                    'balance': 0.0, 'employee_names': [], 'rates': [], 'opening_dates': [],
                    'maturity_dates': [], 'month_terms': [], 'so_tai_khoans': [], 'ma_khs': [],
                    'ma_can_bo_gts': [], 'ten_can_bo_gts': [],
                })
                b_entry['balance'] += balance
                b_entry['employee_names'].append(detail['employee_name'])
                b_entry['rates'].append(detail['rate'])
                b_entry['opening_dates'].append(detail['opening_date'])
                b_entry['maturity_dates'].append(detail['maturity_date'])
                b_entry['month_terms'].append(detail['month_term'])
                b_entry['so_tai_khoans'].append(detail['so_tai_khoan'])
                b_entry['ma_khs'].append(detail['ma_kh'])
                b_entry['ma_can_bo_gts'].append(detail['ma_can_bo_gt'])
                b_entry['ten_can_bo_gts'].append(detail['ten_can_bo_gt'])

        if not sys_agg:
            messages.warning(
                request,
                f"Cảnh báo: Không có khoản HĐV nào trong file IPCAS thỏa điều kiện đối chiếu "
                f"(tổng {len(df_ht)} dòng; VND: {vnd_count} dòng; kỳ hạn ≥ 1 tháng: {term_ok_count} dòng). "
                f"Toàn bộ kết quả sẽ hiển thị 'Không khớp IPCAS' — kiểm tra lại định dạng cột CCY/MONTH_TERM."
            )
        else:
            matched_keys = {
                f"{_hdv_remove_leading_zeros(r[col_cccd])}_{_hdv_clean_str(r[col_emp_ct])}"
                for _, r in df_ct.iterrows()
            } & sys_agg.keys()
            if not matched_keys:
                messages.warning(
                    request,
                    f"Cảnh báo: {len(sys_agg)} khoản IPCAS hợp lệ nhưng không khoản nào khớp với "
                    f"CCCD + Mã cán bộ trong file Chương trình — kiểm tra lại 2 cột này (định dạng, số 0 đầu, mã cán bộ)."
                )

        # Cán bộ đăng ký nhiều lần cho cùng 1 KH (vd sửa/cập nhật số tiền đăng ký):
        # chỉ lần đăng ký mới nhất (theo "Ngày ĐK huy động") được tính, các lần đăng ký
        # cũ hơn cho cùng cặp CCCD + Mã cán bộ bị coi là đã thay thế.
        reg_latest_date = {}
        if col_ngay_dk:
            for _, row in df_ct.iterrows():
                k = f"{_hdv_remove_leading_zeros(row[col_cccd])}_{_hdv_clean_str(row[col_emp_ct])}"
                try:
                    d = datetime.strptime(_hdv_clean_str(row[col_ngay_dk]), '%d/%m/%Y')
                except ValueError:
                    continue
                if k not in reg_latest_date or d > reg_latest_date[k]:
                    reg_latest_date[k] = d

        # Đối chiếu từng dòng file Chương trình
        so_du_ipcas, chenh_lech, trang_thai, co_so_thuong = [], [], [], []
        employee_name_out, rate_out, maturity_out, month_term_out, so_tk_out = [], [], [], [], []
        opening_out = []
        ma_kh_out = []
        ma_can_bo_gt_out, ten_can_bo_gt_out = [], []
        ghi_chu_out = []
        can_bo_agg = {}  # mã cán bộ (đã làm sạch) -> tổng cơ sở tính thưởng
        for _, row in df_ct.iterrows():
            ma_cb   = _hdv_clean_str(row[col_emp_ct])
            cccd    = _hdv_remove_leading_zeros(row[col_cccd])
            key     = f"{cccd}_{ma_cb}"
            amount  = _hdv_to_float(row[col_amount])
            ghi_chu = ''

            # Cán bộ đăng ký nhiều lần cho cùng 1 KH: lần đăng ký cũ hơn bị coi là không hợp lệ,
            # không hiển thị dữ liệu đối chiếu IPCAS (chỉ giữ ghi chú + trạng thái).
            superseded_by = None
            if col_ngay_dk:
                try:
                    reg_date = datetime.strptime(_hdv_clean_str(row[col_ngay_dk]), '%d/%m/%Y')
                except ValueError:
                    reg_date = None
                latest = reg_latest_date.get(key)
                if reg_date is not None and latest is not None and reg_date < latest:
                    superseded_by = latest

            if superseded_by is not None:
                entry = None
                matched = 0.0
                status = f"Không hợp lệ do đã gửi vào ngày {superseded_by.strftime('%d/%m/%Y')}"
                co_so = 0.0
            else:
                entry   = sys_agg.get(key)
                matched = entry['balance'] if entry else 0.0

                # KH tất toán sổ đã gán mã cán bộ rồi gửi lại sổ khác chưa kịp gán mã:
                # nếu cán bộ này đã từng được xác nhận gắn với CCCD này (kể cả ở sổ đã tất toán),
                # vẫn tính chỉ tiêu theo tổng số dư các sổ đang hoạt động (Normal) của KH đó.
                if matched <= 0 and ma_cb and ma_cb in cccd_employees.get(cccd, set()):
                    broadened = cccd_balance.get(cccd)
                    if broadened and broadened['balance'] > 0:
                        entry = broadened
                        matched = broadened['balance']
                        ghi_chu = 'Tính theo sổ khác của KH (đã tất toán sổ gốc, sổ mới chưa gán mã)'

                if matched >= amount and amount > 0:
                    status = "Khớp lệ (Đạt/Vượt chỉ tiêu)"
                elif 0 < matched < amount:
                    status = "Khớp lệ (Thực gửi ít hơn đăng ký)"
                else:
                    status = "Không khớp IPCAS (Chưa gán mã/Chưa gửi)"

                # Cơ sở tính thưởng = min(số tiền đăng ký, số dư thực tế IPCAS)
                co_so = min(amount, matched) if amount > 0 else 0.0
                co_so = max(co_so, 0.0)

            so_du_ipcas.append(matched if superseded_by is None else '')
            chenh_lech.append((matched - amount) if superseded_by is None else '')
            trang_thai.append(status)
            co_so_thuong.append(co_so)
            employee_name_out.append(_hdv_join_unique(entry['employee_names']) if entry else '')
            rate_out.append(_hdv_join_unique(entry['rates']) if entry else '')
            opening_out.append(_hdv_join_unique(entry['opening_dates']) if entry else '')
            maturity_out.append(_hdv_join_unique(entry['maturity_dates']) if entry else '')
            month_term_out.append(_hdv_join_unique(entry['month_terms']) if entry else '')
            so_tk_out.append(_hdv_join_unique(entry['so_tai_khoans']) if entry else '')
            ma_kh_out.append(_hdv_join_unique(entry['ma_khs']) if entry else '')
            ma_can_bo_gt_out.append(_hdv_join_unique(entry['ma_can_bo_gts']) if entry else '')
            ten_can_bo_gt_out.append(_hdv_join_unique(entry['ten_can_bo_gts']) if entry else '')
            ghi_chu_out.append(ghi_chu)

            if ma_cb:
                can_bo_agg[ma_cb] = can_bo_agg.get(ma_cb, 0.0) + co_so

        df_ct['Mã khách hàng (MA_KH)'] = ma_kh_out
        df_ct['Tên cán bộ IPCAS (EMPLOYEE_NAME)'] = employee_name_out
        df_ct['Mã cán bộ giới thiệu (NGUOI_GIOI_THIEU)'] = ma_can_bo_gt_out
        df_ct['Tên cán bộ giới thiệu (TEN_NGUOI_GIOI_THIEU)'] = ten_can_bo_gt_out
        df_ct['Lãi suất (RATE)'] = rate_out
        df_ct['Ngày tới hạn (MATURITY_DATE)'] = maturity_out
        df_ct['Ngày gửi tiền thực tế (OPENING_DATE)'] = opening_out
        df_ct['Kỳ hạn (MONTH_TERM)'] = month_term_out
        df_ct['Số tài khoản (SO_TAI_KHOAN)'] = so_tk_out
        df_ct['Số dư thực tế IPCAS'] = so_du_ipcas
        df_ct['Chênh lệch'] = chenh_lech
        df_ct['Ghi chú đối chiếu'] = ghi_chu_out
        df_ct['Trạng thái đối chiếu'] = trang_thai
        df_ct['Số tiền tính thưởng'] = co_so_thuong

        # Tổng hợp thưởng theo cán bộ (mục II.1 văn bản 329/NHNo.CT-KH&QLRR)
        thuong_rows = []
        for ma_cb, tong_co_so in sorted(can_bo_agg.items(), key=lambda x: -x[1]):
            thuong_truoc_tran = _hdv_round_10k(tong_co_so * HDV_TY_LE_THUONG)
            thuong = min(thuong_truoc_tran, HDV_TRAN_THUONG_THANG)
            thuong_rows.append({
                'Mã cán bộ': ma_cb,
                'Tổng số tiền tính thưởng': tong_co_so,
                'Tỷ lệ thưởng': HDV_TY_LE_THUONG,
                'Số tiền thưởng (trước trần)': thuong_truoc_tran,
                'Số tiền thưởng': thuong,
                'Đã áp trần 5tr/tháng': thuong_truoc_tran > HDV_TRAN_THUONG_THANG,
            })
        df_thuong = pd.DataFrame(thuong_rows, columns=[
            'Mã cán bộ', 'Tổng số tiền tính thưởng', 'Tỷ lệ thưởng',
            'Số tiền thưởng (trước trần)', 'Số tiền thưởng', 'Đã áp trần 5tr/tháng',
        ])

        output = io.BytesIO()
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            df_ct.to_excel(writer, index=False, sheet_name='Ket Qua Doi Chieu')
            df_thuong.to_excel(writer, index=False, sheet_name='Tinh Thuong HDV')

            ws = writer.sheets['Ket Qua Doi Chieu']
            for col_name in [col_amount, 'Số dư thực tế IPCAS']:
                col_idx = df_ct.columns.get_loc(col_name) + 1
                col_letter = get_column_letter(col_idx)
                for cell in ws[col_letter][1:]:
                    cell.number_format = '#,##0'
        output.seek(0)

        response = HttpResponse(
            output.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="Ket_Qua_Doi_Chieu_HDV.xlsx"'
        return response

    except Exception as e:
        traceback.print_exc()
        messages.error(request, f"Đã xảy ra lỗi: {e}")
        return redirect('huy_dong_von_report')
