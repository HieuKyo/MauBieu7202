"""
KPI Report Generator
====================
Tạo file Excel báo cáo KPI tháng với các chỉ tiêu theo từng ngày.
Không cần file mẫu - tự sinh Excel từ dữ liệu.
"""

import calendar
import random
from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


# ── Màu sắc ──────────────────────────────────────────────────────
COLOR_HEADER_BG  = "1F6B35"   # Xanh Agribank (header)
COLOR_TITLE_BG   = "145228"   # Xanh đậm (tiêu đề)
COLOR_TOTAL_BG   = "FFF2CC"   # Vàng nhạt (cột Tổng)
COLOR_DERIVED_BG = "EBF5EB"   # Xanh nhạt (chỉ tiêu tính toán)
COLOR_FIXED_BG   = "EBF3FB"   # Xanh dương nhạt (chỉ tiêu cố định)
COLOR_GRATHIEU   = "FEF9E7"   # Vàng rất nhạt (chỉ tiêu riêng GRATHIEU)
COLOR_WHITE      = "FFFFFF"
COLOR_ROW_ALT    = "F5F5F5"   # Xám nhạt xen kẽ
COLOR_WEEKEND    = "EFEFEF"   # Cột cuối tuần

THIN  = Side(style="thin")
THICK = Side(style="medium")
THIN_BORDER  = Border(left=THIN,  right=THIN,  top=THIN,  bottom=THIN)
THICK_BORDER = Border(left=THICK, right=THICK, top=THICK, bottom=THICK)

# GDV có các chỉ tiêu đặc biệt
SPECIAL_GDV = "GRATHIEU"


def _cell(ws, row, col, value=None, bold=False, font_size=10,
          bg_color=None, font_color="000000",
          h_align="center", v_align="center",
          border=None, wrap=False, number_format=None, italic=False):
    """Ghi giá trị và áp style cho 1 ô."""
    c = ws.cell(row=row, column=col, value=value)
    c.font = Font(name="Times New Roman", bold=bold, size=font_size,
                  color=font_color, italic=italic)
    c.alignment = Alignment(horizontal=h_align, vertical=v_align,
                            wrap_text=wrap)
    if bg_color:
        c.fill = PatternFill("solid", fgColor=bg_color)
    if border:
        c.border = border
    if number_format:
        c.number_format = number_format
    return c


def _add_note(ws, row, col, text):
    """Thêm chú thích (?) vào ô — hiện khi hover trong Excel."""
    c = ws.cell(row=row, column=col)
    c.comment = Comment(text, "Hệ thống KPI")


def _get_working_days(year, month):
    """
    Trả về set các ngày làm việc (Thứ 2 – Thứ 6) trong tháng.
    weekday(): 0=Thứ 2, 1=Thứ 3, 2=Thứ 4, 3=Thứ 5, 4=Thứ 6, 5=Thứ 7, 6=CN
    """
    total_days = calendar.monthrange(year, month)[1]
    return {d for d in range(1, total_days + 1)
            if date(year, month, d).weekday() < 5}


def generate_kpi_report(month, year, user_id,
                        cif_personal_by_day,
                        cif_corporate_by_day,
                        card_by_day,
                        sms_by_day,
                        emobile_by_day,
                        billpayment_by_day=None):
    """
    Tạo Excel báo cáo KPI tháng.

    Tham số (dict {ngày: số_lượng}):
        cif_personal_by_day  : KH cá nhân mở TK theo ngày
        cif_corporate_by_day : KH tổ chức mở TK theo ngày
        card_by_day          : Thẻ phát hành theo ngày
        sms_by_day           : SMS đăng ký theo ngày
        emobile_by_day       : E-Mobile đăng ký theo ngày
        billpayment_by_day   : Hạch toán Billpayment TM/CK theo ngày

    Trả về: BytesIO chứa file .xlsx
    """
    if billpayment_by_day is None:
        billpayment_by_day = {}
    days         = calendar.monthrange(year, month)[1]
    working_days = _get_working_days(year, month)
    is_grathieu  = (user_id == SPECIAL_GDV)

    wb = Workbook()
    ws = wb.active
    ws.title = f"KPI T{month:02d}.{year}"

    # ── Độ rộng cột ──────────────────────────────────────────────
    ws.column_dimensions["A"].width = 5    # STT
    ws.column_dimensions["B"].width = 38   # Chỉ tiêu
    for col_idx in range(3, 3 + days):
        ws.column_dimensions[get_column_letter(col_idx)].width = 4.5
    total_col = 3 + days                   # Cột Tổng
    ws.column_dimensions[get_column_letter(total_col)].width = 8

    # ── Dòng 1: Tiêu đề ─────────────────────────────────────────
    ws.merge_cells(start_row=1, start_column=1,
                   end_row=1, end_column=total_col)
    _cell(ws, 1, 1,
          value=f"BÁO CÁO KPI THÁNG {month:02d}/{year}  –  GDV: {user_id}",
          bold=True, font_size=14,
          bg_color=COLOR_TITLE_BG, font_color=COLOR_WHITE,
          border=THICK_BORDER)
    ws.row_dimensions[1].height = 34

    # ── Dòng 2: Header ngày ──────────────────────────────────────
    _cell(ws, 2, 1, "STT",      bold=True, font_size=10,
          bg_color=COLOR_HEADER_BG, font_color=COLOR_WHITE, border=THIN_BORDER)
    _cell(ws, 2, 2, "Chỉ tiêu", bold=True, font_size=10,
          bg_color=COLOR_HEADER_BG, font_color=COLOR_WHITE, border=THIN_BORDER)

    for d in range(1, days + 1):
        is_weekend = d not in working_days
        bg = COLOR_WEEKEND if is_weekend else COLOR_HEADER_BG
        fc = "999999"    if is_weekend else COLOR_WHITE
        _cell(ws, 2, 2 + d, str(d), bold=True, font_size=9,
              bg_color=bg, font_color=fc, border=THIN_BORDER)

    _cell(ws, 2, total_col, "Tổng", bold=True, font_size=10,
          bg_color=COLOR_HEADER_BG, font_color=COLOR_WHITE, border=THIN_BORDER)
    ws.row_dimensions[2].height = 20

    # ── Tính toán dữ liệu theo ngày ──────────────────────────────
    def get(d_dict, day):
        return d_dict.get(day, 0) or 0

    # Kích hoạt thẻ: random [4,10] mỗi ngày làm việc (chỉ GRATHIEU)
    activation_by_day = {
        d: random.randint(4, 10) for d in working_days
    } if is_grathieu else {}

    # Sắp xếp Báo Nợ/Báo có/In sổ phụ: tổng tháng, gán vào ngày 1 (chỉ GRATHIEU)
    baono_total = (
        random.randint(350, 370) * 2 + random.randint(78, 85)
        if is_grathieu else 0
    )

    rows_data = []
    for d in range(1, days + 1):
        wd = date(year, month, d).weekday()   # 0=T2 … 4=T6, 5=T7, 6=CN
        is_work = d in working_days

        cif_p = get(cif_personal_by_day,  d)
        cif_c = get(cif_corporate_by_day, d)
        card  = get(card_by_day,          d)
        sms   = get(sms_by_day,           d)
        emob  = get(emobile_by_day,       d)
        bill  = get(billpayment_by_day,   d)

        sign    = (cif_p + cif_c) * 2
        # Lưu trữ HS = TK mở + Thẻ PH + SMS + E-Mobile (E-Banking)
        archive = cif_p + cif_c + card + sms + emob

        # ── Chỉ tiêu chỉ có cho GRATHIEU ──
        activ = activation_by_day.get(d, 0)
        pos   = (1   if is_work else 0) if is_grathieu else 0
        evn   = (3   if is_work else 0) if is_grathieu else 0

        # Quản lý ATM: mọi ngày kể cả T7/CN = 1.5 (chỉ GRATHIEU)
        atm_manage = 1.5 if is_grathieu else 0

        # Tiếp quỹ ATM: T2/T6 = 4.5 | T5 = 3 | T3/T4 = 0 (chỉ GRATHIEU)
        if is_grathieu and is_work:
            if wd in (0, 4):   # Thứ 2, Thứ 6
                atm_refill = 4.5
            elif wd == 3:      # Thứ 5
                atm_refill = 3.0
            else:
                atm_refill = 0
        else:
            atm_refill = 0

        # ── Chỉ tiêu cho tất cả GDV ──
        sort_doc = 1 if is_work else 0
        # CB KT Quản lý tin học: random 0–2 mỗi ngày làm việc (chỉ GRATHIEU)
        cbkt = (random.randint(0, 2) if is_work else 0) if is_grathieu else 0

        if d == 1:
            report = 4
        elif d == 15:
            report = 1
        elif d == days:
            report = 2
        else:
            report = 0

        # Sắp xếp Báo Nợ/Báo có/In sổ phụ: toàn bộ tổng tháng gán vào ngày 1
        baono = baono_total if d == 1 else 0

        rows_data.append({
            "cif_p":      cif_p,
            "cif_c":      cif_c,
            "sign":       sign,
            "archive":    archive,
            "card":       card,
            "activ":      activ,
            "pos":        pos,
            "evn":        evn,
            "sms":        sms,
            "emob":       emob,
            "bill":       bill,
            "baono":      baono,
            "sort_doc":   sort_doc,
            "report":     report,
            "atm_manage": atm_manage,
            "atm_refill": atm_refill,
            "cbkt":       cbkt,
        })

    # ── Định nghĩa chỉ tiêu ──────────────────────────────────────
    # (key, tên, màu nền, chú thích ?, định dạng số, override_total)
    # override_total = None  → tổng = tổng cộng các ô ngày (bình thường)
    # override_total = value → tổng dùng giá trị cố định, ô ngày để trống
    # Dòng GRATHIEU dùng COLOR_GRATHIEU để phân biệt
    metrics = [
        (
            "cif_p",
            "Đăng ký KH Cá Nhân",
            COLOR_WHITE,
            "Nguồn: File CIF\n"
            "• Lọc cột tellernm = Mã GDV\n"
            "• Lọc cột locdpnm ∈ {\n"
            "  'Tiền gửi thanh toán cá nhân',\n"
            "  'TG KKH CB lương Ngân sách',\n"
            "  'TG KKH Cá nhân (Số đẹp)',\n"
            "  'TG thanh toán cá nhân eKYC'}\n"
            "• Đếm số lượng mở trong ngày (cột opndt)",
            "0",
            None,
        ),
        (
            "cif_c",
            "Đăng ký KH Tổ Chức",
            COLOR_WHITE,
            "Nguồn: File CIF\n"
            "• Lọc cột tellernm = Mã GDV\n"
            "• Lọc cột locdpnm ∈ {\n"
            "  'TG KKH TCKT',\n"
            "  'Tg KKH TCKT (Số đẹp)'}\n"
            "• Đếm số lượng mở trong ngày (cột opndt)",
            "0", None,
        ),
        (
            "sign",
            "Quét chữ ký KH",
            COLOR_DERIVED_BG,
            "Tính toán tự động:\n"
            "= (KH Cá Nhân + KH Tổ Chức) × 2\n"
            "Mỗi khách hàng mở TK cần quét 2 chữ ký",
            "0", None,
        ),
        (
            "archive",
            "Lưu Trữ HS Mở TK, PH Thẻ, SMS, E-Banking",
            COLOR_DERIVED_BG,
            "Tính toán tự động:\n"
            "= KH Cá Nhân + KH Tổ Chức + Phát hành Thẻ + Đăng ký SMS + E-Mobile\n"
            "Mỗi nghiệp vụ phát sinh = 1 bộ hồ sơ cần lưu trữ",
            "0", None,
        ),
        (
            "card",
            "Phát hành Thẻ",
            COLOR_WHITE,
            "Nguồn: File Thẻ\n"
            "• Cột CUSER dùng định dạng riêng, hệ thống tự quy đổi:\n"
            "  GRANTHAO→7202CTHAOTN  GRALTHUC→7202cthuclt\n"
            "  GRATHIEU→7202chieutt  GRANSINH→7202CSINHNT\n"
            "  GRATTHAO→7202CTHAOTLT GRASHANH→7202canhsh\n"
            "  GRACACHI→7202CCHICA   GRATNNHI→7202cnhitn\n"
            "• Đếm tất cả thẻ theo ngày CDATE\n"
            "  (CSP_New = phát hành mới, CSP_Reissue = phát hành lại)",
            "0", None,
        ),
        (
            "activ",
            "Kích hoạt Thẻ  [GRATHIEU]",
            COLOR_GRATHIEU,
            "Chỉ áp dụng cho GDV GRATHIEU.\n"
            "Số liệu ước tính:\n"
            "• Ngày làm việc (T2–T6): lấy ngẫu nhiên từ 4 đến 10\n"
            "• Thứ 7, Chủ nhật: = 0\n"
            "(Số thực tế lấy từ màn hình kích hoạt thẻ hệ thống)",
            "0", None,
        ),
        (
            "pos",
            "Quản lý POS  [GRATHIEU]",
            COLOR_GRATHIEU,
            "Chỉ áp dụng cho GDV GRATHIEU.\n"
            "Số liệu cố định:\n"
            "• Ngày làm việc (T2–T6): = 1\n"
            "• Thứ 7, Chủ nhật: = 0",
            "0", None,
        ),
        (
            "evn",
            "Quản lý vốn Điện lực EVN  [GRATHIEU]",
            COLOR_GRATHIEU,
            "Chỉ áp dụng cho GDV GRATHIEU.\n"
            "Số liệu cố định:\n"
            "• Ngày làm việc (T2–T6): = 3\n"
            "• Thứ 7, Chủ nhật: = 0",
            "0", None,
        ),
        (
            "sms",
            "Đăng ký SMS Banking",
            COLOR_WHITE,
            "Nguồn: File SMS\n"
            "• Lọc cột crtusr = Mã GDV\n"
            "• CHỈ tính dòng có uptdtm RỖNG\n"
            "  (uptdtm có giá trị = bản cập nhật SĐT, không tính)\n"
            "• Đếm số đăng ký mới theo ngày entydt",
            "0", None,
        ),
        (
            "emob",
            "Đăng ký E-Mobile Banking",
            COLOR_WHITE,
            "Nguồn: File E-Mobile Banking\n"
            "• Lọc cột crtusr = Mã GDV\n"
            "• CHỈ tính dòng có uptdtm RỖNG\n"
            "  (uptdtm có giá trị = bản cập nhật, không tính)\n"
            "• Đếm số đăng ký mới theo ngày entydt",
            "0", None,
        ),
        (
            "bill",
            "Hạch toán Billpayment TM, CK",
            COLOR_WHITE,
            "Nguồn: File Bảng kê chứng từ giao dịch chi tiết\n"
            "• Cột header STT nằm ở dòng B10 (Excel row 10)\n"
            "• Cột 'Ngày GD': định dạng yyyymmddHHMMSS\n"
            "  (vd: 20260327154258 = 27/03/2026 15:42:58)\n"
            "• Không lọc theo user — file đã được GDV lọc sẵn\n"
            "• Đếm số giao dịch trong ngày",
            "0", None,
        ),
        (
            "baono",
            "Sắp xếp Báo Nợ, Báo có, In sổ phụ  [GRATHIEU]",
            COLOR_GRATHIEU,
            "Chỉ áp dụng cho GDV GRATHIEU.\n"
            "Tổng tháng tính ngẫu nhiên:\n"
            "= random(350 – 370) × 2  +  random(78 – 85)\n"
            "Toàn bộ tổng tháng được gán vào cột ngày đầu tháng (ngày 1).",
            "0", None,
        ),
        (
            "sort_doc",
            "Sắp xếp chứng từ giao Hậu kiểm",
            COLOR_FIXED_BG,
            "Số liệu cố định — áp dụng tất cả GDV:\n"
            "• Ngày làm việc (T2–T6): = 1\n"
            "• Thứ 7, Chủ nhật: = 0",
            "0", None,
        ),
        (
            "report",
            "Lập báo cáo",
            COLOR_FIXED_BG,
            "Số liệu cố định — áp dụng tất cả GDV:\n"
            "• Ngày 01 (đầu tháng) : 4 điểm\n"
            "• Ngày 15 (giữa tháng): 1 điểm\n"
            "• Ngày cuối tháng     : 2 điểm\n"
            "• Các ngày còn lại    : 0\n"
            "(Tính cả ngày nghỉ nếu rơi vào T7/CN)",
            "0", None,
        ),
        (
            "cbkt",
            "CB KT Quản lý tin học  [GRATHIEU]",
            COLOR_GRATHIEU,
            "Chỉ áp dụng cho GDV GRATHIEU.\n"
            "Số liệu ước tính:\n"
            "• Ngày làm việc (T2–T6): random từ 0 đến 2\n"
            "• Thứ 7, Chủ nhật: = 0",
            "0", None,
        ),
        (
            "atm_manage",
            "Quản lý ATM  [GRATHIEU]",
            COLOR_GRATHIEU,
            "Chỉ áp dụng cho GDV GRATHIEU.\n"
            "Số liệu cố định:\n"
            "• Tất cả các ngày (kể cả T7, Chủ nhật): = 1.5",
            "0.##", None,
        ),
        (
            "atm_refill",
            "Tiếp quỹ ATM  [GRATHIEU]",
            COLOR_GRATHIEU,
            "Chỉ áp dụng cho GDV GRATHIEU.\n"
            "Số liệu cố định theo thứ trong tuần:\n"
            "• Thứ 2 (Monday) : 4.5 điểm\n"
            "• Thứ 3 (Tuesday): 0\n"
            "• Thứ 4 (Wednesday): 0\n"
            "• Thứ 5 (Thursday): 3 điểm\n"
            "• Thứ 6 (Friday) : 4.5 điểm\n"
            "• Thứ 7, Chủ nhật: 0",
            "0.##", None,
        ),
    ]

    # ── Ghi các dòng chỉ tiêu ────────────────────────────────────
    _row_totals_by_metric = []   # dùng để tính dòng Tổng cộng

    for row_offset, (key, label, row_bg, note_text, num_fmt, override_total) in enumerate(metrics):
        excel_row = 3 + row_offset

        # Xen kẽ màu cho dòng trắng thuần
        bg = row_bg
        if bg == COLOR_WHITE and row_offset % 2 == 1:
            bg = COLOR_ROW_ALT

        # STT
        _cell(ws, excel_row, 1, row_offset + 1, font_size=10,
              bg_color=bg, border=THIN_BORDER)

        # Chỉ tiêu + chú thích ?
        _cell(ws, excel_row, 2, label, font_size=10,
              bg_color=bg, border=THIN_BORDER, h_align="left")
        _add_note(ws, excel_row, 2, note_text)

        if override_total is not None:
            # Metric chỉ có tổng tháng, không có số liệu theo ngày
            for d in range(1, days + 1):
                _cell(ws, excel_row, 2 + d,
                      value=None, font_size=9,
                      bg_color=bg, border=THIN_BORDER)
            row_total = override_total
        else:
            # Metric bình thường: cộng dồn từng ngày
            row_total = 0
            for d in range(1, days + 1):
                val = rows_data[d - 1][key] if key else 0
                row_total += val
                _cell(ws, excel_row, 2 + d,
                      value=val if val else None,
                      font_size=9, bg_color=bg, border=THIN_BORDER,
                      number_format=num_fmt)

        # Tổng
        _cell(ws, excel_row, total_col,
              value=row_total if row_total else None,
              bold=True, font_size=10,
              bg_color=COLOR_TOTAL_BG, border=THIN_BORDER,
              number_format=num_fmt)

        ws.row_dimensions[excel_row].height = 20

        # Tích lũy vào danh sách row_totals để tính dòng Tổng cộng
        _row_totals_by_metric.append((row_total, num_fmt))

    # ── Dòng Tổng cộng ───────────────────────────────────────────
    grand_row = 3 + len(metrics)
    ws.row_dimensions[grand_row].height = 22

    _cell(ws, grand_row, 1, "",
          bold=True, font_size=10,
          bg_color=COLOR_TITLE_BG, font_color=COLOR_WHITE, border=THIN_BORDER)
    _cell(ws, grand_row, 2, "TỔNG CỘNG",
          bold=True, font_size=10,
          bg_color=COLOR_TITLE_BG, font_color=COLOR_WHITE,
          border=THIN_BORDER, h_align="center")

    grand_total = 0
    for d in range(1, days + 1):
        # Cộng tất cả chỉ tiêu cho ngày d
        day_sum = sum(rows_data[d - 1].get(m[0], 0)
                      for m in metrics if m[0] is not None)
        grand_total += day_sum
        _cell(ws, grand_row, 2 + d,
              value=day_sum if day_sum else None,
              bold=True, font_size=9,
              bg_color=COLOR_TITLE_BG, font_color=COLOR_WHITE,
              border=THIN_BORDER, number_format="0.##")

    _cell(ws, grand_row, total_col,
          value=grand_total if grand_total else None,
          bold=True, font_size=10,
          bg_color=COLOR_TOTAL_BG, border=THIN_BORDER,
          number_format="0.##")

    # ── Ghi chú cuối bảng ────────────────────────────────────────
    # grand_row = 3 + len(metrics), note_row ngay sau đó
    note_row = 3 + len(metrics) + 1
    ws.merge_cells(start_row=note_row, start_column=1,
                   end_row=note_row, end_column=total_col)
    _cell(ws, note_row, 1,
          value=(
              "(*) Hover chuột vào tên chỉ tiêu để xem hướng dẫn lấy số liệu.  "
              "Dòng nền vàng nhạt = chỉ tiêu riêng của GRATHIEU."
          ),
          font_size=9, italic=True, font_color="666666",
          h_align="left", bg_color=COLOR_WHITE)

    # ── Freeze panes ─────────────────────────────────────────────
    ws.freeze_panes = ws.cell(row=3, column=3)

    # ── Xuất BytesIO ─────────────────────────────────────────────
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
