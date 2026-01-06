"""
KPI Calculator Module
=====================
Module tính toán và cập nhật KPI vào file Excel mẫu
"""

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from io import BytesIO
import re


class KPICalculator:
    """
    Class tính toán và cập nhật KPI vào Excel template
    """

    def __init__(self, coefficients):
        """
        Khởi tạo KPICalculator

        Args:
            coefficients: Dict chứa các hệ số quy đổi
                {
                    'card': 3.0,
                    'signature': 3.0,
                    'sms': 4.0,
                    'archive': 0.5,
                    'cif': 3.0
                }
        """
        self.coefficients = coefficients
        self.wb = None
        self.ws = None
        self.header_row = None
        self.day_columns = {}  # Mapping: {day: column_letter}
        self.kpi_rows = {}  # Mapping: {kpi_name: row_number}

    def _find_header_row(self):
        """
        Tìm dòng header chứa các số ngày (1, 2, 3...31)

        Returns:
            int: Số dòng header (hoặc None nếu không tìm thấy)
        """
        for row_idx in range(1, min(20, self.ws.max_row + 1)):  # Tìm trong 20 dòng đầu
            row_values = []
            for col_idx in range(1, self.ws.max_column + 1):
                cell = self.ws.cell(row_idx, col_idx)
                row_values.append(cell.value)

            # Kiểm tra xem có chứa dãy số 1, 2, 3... không
            numbers_found = []
            for val in row_values:
                if isinstance(val, (int, float)) and 1 <= val <= 31:
                    numbers_found.append(int(val))

            # Nếu tìm thấy ít nhất 5 số liên tiếp
            if len(numbers_found) >= 5:
                # Kiểm tra có phải dãy tăng dần không
                is_sequential = True
                for i in range(len(numbers_found) - 1):
                    if numbers_found[i + 1] != numbers_found[i] + 1:
                        is_sequential = False
                        break

                if is_sequential and numbers_found[0] == 1:
                    return row_idx

        return None

    def _map_day_columns(self):
        """
        Tạo mapping giữa ngày (1-31) và cột Excel
        """
        if self.header_row is None:
            raise ValueError("Không tìm thấy dòng header chứa các ngày!")

        self.day_columns = {}

        for col_idx in range(1, self.ws.max_column + 1):
            cell_value = self.ws.cell(self.header_row, col_idx).value
            if isinstance(cell_value, (int, float)) and 1 <= cell_value <= 31:
                day = int(cell_value)
                self.day_columns[day] = get_column_letter(col_idx)

    def _find_kpi_rows(self):
        """
        Tìm các dòng KPI dựa trên STT hoặc từ khóa

        Tìm:
        - STT 1 (CIF): "Đăng ký TT KH cá nhân"
        - STT 3 (Signature): "Quét chữ ký"
        - STT 4 (Archive): "lưu trữ"
        - STT 9 (SMS): "SMS"
        - STT 12 (Card): "Phát hành thẻ"
        """
        self.kpi_rows = {}

        # Duyệt qua các dòng để tìm
        for row_idx in range(1, self.ws.max_row + 1):
            # Tìm cột STT và Nội dung
            stt_value = None
            content_value = None

            for col_idx in range(1, min(10, self.ws.max_column + 1)):  # Tìm trong 10 cột đầu
                cell = self.ws.cell(row_idx, col_idx)
                cell_value = str(cell.value).strip() if cell.value else ""

                # Kiểm tra cột STT (thường là số)
                if cell_value.isdigit():
                    stt_value = int(cell_value)

                # Kiểm tra cột nội dung (text dài hơn)
                if len(cell_value) > 10:
                    content_value = cell_value.lower()

            # Mapping dựa trên STT - ưu tiên STT trước, sau đó mới dùng từ khóa
            # STT 12: Phát hành thẻ (kiểm tra TRƯỚC để tránh conflict)
            if stt_value == 12:
                self.kpi_rows['card'] = row_idx

            # STT 1: CIF
            elif stt_value == 1:
                self.kpi_rows['cif'] = row_idx

            # STT 3: Quét chữ ký
            elif stt_value == 3:
                self.kpi_rows['signature'] = row_idx

            # STT 4: Lưu trữ hồ sơ
            elif stt_value == 4:
                self.kpi_rows['archive'] = row_idx

            # STT 9: Đăng ký SMS
            elif stt_value == 9:
                self.kpi_rows['sms'] = row_idx

            # Fallback: Tìm theo từ khóa nếu chưa tìm thấy qua STT
            elif content_value:
                # Tìm dòng Phát hành thẻ (nếu chưa có)
                if 'card' not in self.kpi_rows and 'phát hành thẻ' in content_value:
                    self.kpi_rows['card'] = row_idx

                # Tìm dòng CIF (nếu chưa có)
                elif 'cif' not in self.kpi_rows and 'đăng ký tt kh' in content_value:
                    self.kpi_rows['cif'] = row_idx

                # Tìm dòng Chữ ký (nếu chưa có)
                elif 'signature' not in self.kpi_rows and 'quét chữ ký' in content_value:
                    self.kpi_rows['signature'] = row_idx

                # Tìm dòng Lưu trữ (nếu chưa có)
                elif 'archive' not in self.kpi_rows and 'lưu trữ' in content_value and 'gdv' in content_value:
                    self.kpi_rows['archive'] = row_idx

                # Tìm dòng SMS (nếu chưa có) - CẨN THẬN với từ khóa SMS
                elif 'sms' not in self.kpi_rows and 'đăng ký sms' in content_value:
                    self.kpi_rows['sms'] = row_idx

        # Debug: In ra các dòng đã tìm thấy
        print("\n📋 Debug - Các dòng KPI đã tìm thấy:")
        for key, row in self.kpi_rows.items():
            cell_content = self.ws.cell(row, 3).value  # Cột C thường chứa nội dung
            print(f"  - {key.upper()}: Dòng {row} - '{cell_content}'")
        print()

    def _update_cell(self, row, day, value):
        """
        Cập nhật giá trị vào một ô cụ thể

        Args:
            row: Số dòng
            day: Ngày (1-31)
            value: Giá trị cần ghi
        """
        if day not in self.day_columns:
            return  # Bỏ qua nếu không có cột cho ngày này

        col_letter = self.day_columns[day]
        cell = self.ws[f"{col_letter}{row}"]
        cell.value = value if value > 0 else None

    def _update_row_totals(self, row, coefficient):
        """
        Cập nhật tổng hàng ngang và bút toán quy đổi

        Args:
            row: Số dòng cần cập nhật
            coefficient: Hệ số quy đổi
        """
        # Tính tổng từ các ô ngày
        total = 0
        for day in range(1, 32):
            if day in self.day_columns:
                col_letter = self.day_columns[day]
                cell_value = self.ws[f"{col_letter}{row}"].value
                if cell_value and isinstance(cell_value, (int, float)):
                    total += cell_value

        # Tìm cột "Cộng" (thường sau cột ngày 31)
        # Giả sử cột "Cộng" nằm ngay sau cột ngày cuối cùng
        last_day_col = max(self.day_columns.values(), key=lambda x: ord(x[0]) if len(x) == 1 else ord(x[0]) * 26 + ord(x[1]))

        # Tìm cột Cộng bằng cách duyệt các cột sau ngày cuối
        cong_col = None
        for offset in range(1, 5):
            col_idx = self._column_letter_to_index(last_day_col) + offset
            col_letter = get_column_letter(col_idx)
            cell = self.ws[f"{col_letter}{self.header_row}"]
            cell_value = str(cell.value).lower() if cell.value else ""
            if 'cộng' in cell_value or 'tổng' in cell_value:
                cong_col = col_letter
                break

        # Cập nhật cột Cộng
        if cong_col:
            self.ws[f"{cong_col}{row}"].value = total if total > 0 else None

            # Cập nhật cột "Bút toán quy đổi" (cột tiếp theo sau "Cộng")
            bt_col_idx = self._column_letter_to_index(cong_col) + 1
            bt_col = get_column_letter(bt_col_idx)
            converted_value = total * coefficient
            self.ws[f"{bt_col}{row}"].value = converted_value if converted_value > 0 else None

    def _column_letter_to_index(self, col_letter):
        """
        Chuyển đổi chữ cái cột sang index (A=1, B=2, ...)

        Args:
            col_letter: Chữ cái cột (VD: 'A', 'AA')

        Returns:
            int: Index của cột
        """
        index = 0
        for i, char in enumerate(reversed(col_letter.upper())):
            index += (ord(char) - ord('A') + 1) * (26 ** i)
        return index

    def calculate_and_update_template(self, template_file, card_data, sms_data, emobile_data, days_in_month):
        """
        Tính toán KPI và cập nhật vào file Excel template

        Args:
            template_file: File object của template Excel
            card_data: DataFrame dữ liệu thẻ
            sms_data: DataFrame dữ liệu SMS
            emobile_data: DataFrame dữ liệu E-Mobile
            days_in_month: Số ngày trong tháng

        Returns:
            BytesIO: Buffer chứa file Excel đã cập nhật
        """
        # Reset file pointer
        template_file.seek(0)

        # Load workbook
        self.wb = load_workbook(template_file)
        self.ws = self.wb.active

        # Tìm header row và mapping
        self.header_row = self._find_header_row()
        if self.header_row is None:
            raise ValueError("Không tìm thấy dòng header chứa các ngày trong file mẫu!")

        self._map_day_columns()
        self._find_kpi_rows()

        # Kiểm tra các dòng KPI có tìm thấy không
        if not self.kpi_rows:
            raise ValueError("Không tìm thấy các dòng KPI trong file mẫu!")

        # Duyệt qua từng ngày và cập nhật
        for day in range(1, days_in_month + 1):
            # Lấy dữ liệu cho ngày này
            card_count = self._get_count(card_data, day, 'count')
            new_issue_count = self._get_count(card_data, day, 'new_issue_count')
            sms_count = self._get_count(sms_data, day, 'count')
            emobile_count = self._get_count(emobile_data, day, 'count')

            # Cập nhật các dòng KPI
            # STT 12: Phát hành thẻ
            if 'card' in self.kpi_rows:
                self._update_cell(self.kpi_rows['card'], day, card_count)

            # STT 3: Quét chữ ký (= New Issue * 2)
            if 'signature' in self.kpi_rows:
                self._update_cell(self.kpi_rows['signature'], day, new_issue_count * 2)

            # STT 1: CIF mới (= New Issue)
            if 'cif' in self.kpi_rows:
                self._update_cell(self.kpi_rows['cif'], day, new_issue_count)

            # STT 9: Đăng ký SMS
            if 'sms' in self.kpi_rows:
                self._update_cell(self.kpi_rows['sms'], day, sms_count)

            # STT 4: Lưu trữ hồ sơ (= Card + SMS + E-Mobile)
            if 'archive' in self.kpi_rows:
                archive_count = card_count + sms_count + emobile_count
                self._update_cell(self.kpi_rows['archive'], day, archive_count)

        # Cập nhật tổng cho từng dòng
        if 'card' in self.kpi_rows:
            self._update_row_totals(self.kpi_rows['card'], self.coefficients['card'])

        if 'signature' in self.kpi_rows:
            self._update_row_totals(self.kpi_rows['signature'], self.coefficients['signature'])

        if 'cif' in self.kpi_rows:
            self._update_row_totals(self.kpi_rows['cif'], self.coefficients['cif'])

        if 'sms' in self.kpi_rows:
            self._update_row_totals(self.kpi_rows['sms'], self.coefficients['sms'])

        if 'archive' in self.kpi_rows:
            self._update_row_totals(self.kpi_rows['archive'], self.coefficients['archive'])

        # Lưu vào BytesIO
        output = BytesIO()
        self.wb.save(output)
        output.seek(0)

        return output

    def _get_count(self, data_df, day, column='count'):
        """
        Lấy giá trị count cho một ngày cụ thể

        Args:
            data_df: DataFrame
            day: Ngày cần lấy
            column: Tên cột cần lấy

        Returns:
            int: Giá trị (0 nếu không có)
        """
        if data_df is None or data_df.empty:
            return 0

        day_data = data_df[data_df['day'] == day]
        if day_data.empty:
            return 0

        if column not in day_data.columns:
            return 0

        return int(day_data[column].iloc[0])
