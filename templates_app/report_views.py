"""
Views cho chức năng Báo cáo
"""
from django.shortcuts import render, redirect
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse, JsonResponse
from django.views.decorators.http import require_http_methods
from django.contrib import messages
from .models import MailEnvelopeTracking, ReportConfiguration
import pandas as pd
import openpyxl
import io
import json
import traceback
from openpyxl.utils.dataframe import dataframe_to_rows
from copy import copy
from datetime import datetime


@login_required
def report_hub(request):
    """Trang hub cho các loại báo cáo"""
    return render(request, 'templates_app/reports/report_hub.html')


@login_required
def lai_ton_dong_report_view(request):
    """Giao diện báo cáo Lãi tồn đọng"""
    return render(request, 'templates_app/reports/lai_ton_dong.html')


@login_required
@require_http_methods(["POST"])
def process_lai_ton_dong_report(request):
    """Xử lý báo cáo Lãi tồn đọng"""
    try:
        # Lấy cấu hình từ database
        try:
            config_obj = ReportConfiguration.objects.get(report_type='lai_ton_dong', is_active=True)
            ltd_config = config_obj.config_data
        except (ReportConfiguration.DoesNotExist, Exception):
            # Cấu hình mặc định
            ltd_config = {
                'allowed_acctcd': [701001, 701002, 701003],
                'template_row_index': 11,
                'start_row': 11
            }

        current_file = request.FILES.get('current_file')
        comparison_file = request.FILES.get('comparison_file')

        if not all([current_file, comparison_file]):
            messages.error(request, "Vui lòng tải lên đủ 2 file.")
            return redirect('lai_ton_dong_report')

        def process_raw_file(file):
            # Đọc custseq như text để tránh bị format thành số thập phân (7.20201)
            df = pd.read_excel(file, dtype={'custseq': str})

            allowed_acctcd = ltd_config.get('allowed_acctcd', [])
            if 'acctcd' in df.columns:
                df['acctcd'] = pd.to_numeric(df['acctcd'], errors='coerce')
                df = df[df['acctcd'].isin(allowed_acctcd)]

            if 'acrbamt' in df.columns:
                df['acrbamt'] = df['acrbamt'].abs()
            return df

        df_current = process_raw_file(current_file)
        df_comparison = process_raw_file(comparison_file)

        cols_to_keep = ['custseq', 'custnm', 'refno', 'acrbamt', 'bceqa']

        df_current_sel = df_current[cols_to_keep].rename(columns={
            'custseq': 'Mã khách hàng',
            'custnm': 'Tên khách hàng',
            'refno': 'LDS',
            'acrbamt': 'Dư nợ gốc kỳ này',
            'bceqa': 'Lãi kỳ này'
        })
        df_comparison_sel = df_comparison[cols_to_keep].rename(columns={
            'custseq': 'Mã khách hàng SS',
            'custnm': 'Tên khách hàng SS',
            'refno': 'LDS',
            'acrbamt': 'Dư nợ gốc kỳ SS',
            'bceqa': 'Lãi kỳ SS'
        })

        # OUTER JOIN để lấy tất cả LDS từ cả 2 kỳ
        merged_df = pd.merge(
            df_current_sel,
            df_comparison_sel,
            on='LDS',
            how='outer'
        )

        # Điền thông tin khách hàng: ưu tiên kỳ hiện tại, nếu không có thì lấy từ kỳ so sánh
        merged_df['Mã khách hàng'] = merged_df['Mã khách hàng'].fillna(merged_df['Mã khách hàng SS'])
        merged_df['Tên khách hàng'] = merged_df['Tên khách hàng'].fillna(merged_df['Tên khách hàng SS'])

        # Xóa cột tạm
        merged_df.drop(['Mã khách hàng SS', 'Tên khách hàng SS'], axis=1, inplace=True)

        # Điền 0 cho các giá trị số còn thiếu
        merged_df.fillna(0, inplace=True)

        merged_df['Lãi Tăng'] = (merged_df['Lãi kỳ này'] - merged_df['Lãi kỳ SS']).clip(lower=0)
        merged_df['Lãi Giảm'] = (merged_df['Lãi kỳ SS'] - merged_df['Lãi kỳ này']).clip(lower=0)

        final_df = merged_df[
            (merged_df['Dư nợ gốc kỳ này'] != 0) |
            (merged_df['Lãi kỳ này'] != 0) |
            (merged_df['Dư nợ gốc kỳ SS'] != 0) |
            (merged_df['Lãi kỳ SS'] != 0)
        ].copy()

        output_cols = [
            'Mã khách hàng', 'Tên khách hàng', 'LDS',
            'Dư nợ gốc kỳ SS', 'Lãi kỳ SS',
            'Dư nợ gốc kỳ này', 'Lãi kỳ này',
            'Lãi Tăng', 'Lãi Giảm'
        ]
        final_df = final_df[output_cols]

        # Tạo file Excel từ template
        template_path = 'packages/mau_lai_tondong.xlsx'  # Cần copy template vào đây
        try:
            workbook = openpyxl.load_workbook(template_path)
            sheet = workbook.active
        except FileNotFoundError:
            return HttpResponse("Không tìm thấy file mẫu. Vui lòng liên hệ quản trị viên.", status=500)

        template_row_index = ltd_config.get('template_row_index', 11)
        start_row = ltd_config.get('start_row', 11)

        # Xóa merged cells cũ
        merged_ranges = list(sheet.merged_cells.ranges)
        for merged_cell_range in merged_ranges:
            if merged_cell_range.min_row >= start_row:
                sheet.unmerge_cells(str(merged_cell_range))

        # Xóa dữ liệu cũ
        if sheet.max_row >= start_row:
            for row in sheet.iter_rows(min_row=start_row, max_row=sheet.max_row + 10):
                for cell in row:
                    cell.value = None

        # Copy style từ template row
        template_cells = sheet[template_row_index]
        rows_to_write = dataframe_to_rows(final_df, index=False, header=False)

        for r_idx, row_data in enumerate(rows_to_write, start_row):
            for c_idx, value in enumerate(row_data, 1):
                new_cell = sheet.cell(row=r_idx, column=c_idx)
                if c_idx <= len(template_cells):
                    template_cell = template_cells[c_idx - 1]
                    if template_cell.has_style:
                        new_cell.font = copy(template_cell.font)
                        new_cell.border = copy(template_cell.border)
                        new_cell.fill = copy(template_cell.fill)
                        new_cell.number_format = template_cell.number_format
                        new_cell.protection = copy(template_cell.protection)
                        new_cell.alignment = copy(template_cell.alignment)
                new_cell.value = value

        # Thêm dòng tổng
        total_row_index = start_row + len(final_df)
        sheet.cell(row=total_row_index, column=1, value='TỔNG CỘNG')
        numeric_cols_to_sum = ['Dư nợ gốc kỳ này', 'Lãi kỳ này', 'Dư nợ gốc kỳ SS', 'Lãi kỳ SS', 'Lãi Tăng', 'Lãi Giảm']
        col_index_map = {name: i+1 for i, name in enumerate(output_cols)}

        for col_name in numeric_cols_to_sum:
            col_sum = final_df[col_name].sum()
            if col_name in col_index_map:
                total_cell = sheet.cell(row=total_row_index, column=col_index_map[col_name])
                template_cell = template_cells[col_index_map[col_name] - 1]
                if template_cell.has_style:
                    total_cell.font = copy(template_cell.font)
                    total_cell.border = copy(template_cell.border)
                    total_cell.fill = copy(template_cell.fill)
                    total_cell.number_format = template_cell.number_format
                    total_cell.protection = copy(template_cell.protection)
                    total_cell.alignment = copy(template_cell.alignment)
                total_cell.value = col_sum

        # Xuất file
        output = io.BytesIO()
        workbook.save(output)
        output.seek(0)

        response = HttpResponse(
            output.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="BaoCao_SoSanh_LaiTonDong_KetQua.xlsx"'
        return response

    except Exception as e:
        traceback.print_exc()
        messages.error(request, f"Đã xảy ra lỗi: {e}")
        return redirect('lai_ton_dong_report')


# ── Helpers đọc file dùng chung cho các báo cáo (report_*_views.py) ──────────

def _read_excel_safe(file_obj, **kwargs):
    """Đọc file Excel, tự bỏ qua workbook corruption cho file .xls cũ.
    Truyền usecols=... để chỉ đọc cột cần thiết, giảm bộ nhớ cho file lớn."""
    file_bytes = file_obj.read() if hasattr(file_obj, 'read') else file_obj
    try:
        return pd.read_excel(io.BytesIO(file_bytes), **kwargs)
    except Exception:
        import xlrd
        wb = xlrd.open_workbook(file_contents=file_bytes, ignore_workbook_corruption=True)
        return pd.read_excel(wb, **kwargs)


def _read_csv_robust(file_obj, usecols=None):
    """
    Đọc file .csv thử qua nhiều encoding (file xuất từ IPCAS thường là Windows-1258/ANSI
    tiếng Việt, không phải UTF-8) kết hợp dấu phẩy/tab làm dấu phân cách — ưu tiên kết quả
    đọc ra nhiều hơn 1 cột. Tự giải mã bytes → text trước để tránh lỗi nội bộ của pandas khi
    kết hợp tự dò encoding/delimiter trực tiếp trên file object nhị phân.
    """
    raw = file_obj.read() if hasattr(file_obj, 'read') else file_obj
    if isinstance(raw, str):
        raw = raw.encode('utf-8')

    best = None
    last_err = None
    for encoding in ('utf-8', 'utf-8-sig', 'cp1258', 'latin1'):
        try:
            text = raw.decode(encoding)
        except UnicodeDecodeError as e:
            last_err = e
            continue
        for sep in (',', '\t'):
            try:
                df = pd.read_csv(io.StringIO(text), sep=sep)
            except pd.errors.ParserError as e:
                last_err = e
                continue
            if df.shape[1] > 1:
                df.columns = [str(c).strip() for c in df.columns]
                return df[usecols] if usecols else df
            if best is None:
                best = df
    if best is not None:
        best.columns = [str(c).strip() for c in best.columns]
        return best[usecols] if usecols else best
    raise last_err


def _strip_text_marker(series):
    """Bỏ dấu ' ở đầu giá trị (marker buộc kiểu text khi export CSV từ IPCAS) và khoảng trắng dư."""
    return series.astype(str).str.strip().str.lstrip("'").str.strip()


@login_required
def mail_envelope_tracking_view(request):
    """Giao diện nhận bì thư"""
    # Lấy danh sách user từ cấu hình nếu có
    try:
        config_obj = ReportConfiguration.objects.get(report_type='phat_hanh_the', is_active=True)
        pgd_user_map = config_obj.config_data.get('pgd_user_map', {})
        all_users = sorted([user for users in pgd_user_map.values() for user in users])
    except Exception:
        all_users = []

    return render(request, 'templates_app/reports/mail_envelope_tracking.html', {'all_users': all_users})


@login_required
@require_http_methods(["POST"])
def save_mail_envelope(request):
    """Lưu thông tin nhận bì thư"""
    try:
        data = json.loads(request.body)

        MailEnvelopeTracking.objects.create(
            envelope_code=data['ma_bithu'],
            receive_date=data['ngay_nhan'],
            receiver=data['nguoi_nhan']
        )

        return JsonResponse({'status': 'success', 'message': 'Đã lưu thành công!'}, status=200)
    except Exception as e:
        return JsonResponse({'status': 'error', 'message': str(e)}, status=500)


@login_required
def mail_envelope_report_view(request):
    """Báo cáo bì thư"""
    ma_bithu_search = request.GET.get('ma_bithu', '').strip()
    start_date_str = request.GET.get('start_date')
    end_date_str = request.GET.get('end_date')

    records = []

    if ma_bithu_search:
        records = MailEnvelopeTracking.objects.filter(
            envelope_code__icontains=ma_bithu_search
        ).order_by('-receive_date', '-created_at')
    elif start_date_str and end_date_str:
        start_date = datetime.strptime(start_date_str, '%Y-%m-%d').date()
        end_date = datetime.strptime(end_date_str, '%Y-%m-%d').date()
        records = MailEnvelopeTracking.objects.filter(
            receive_date__range=[start_date, end_date]
        ).order_by('-receive_date', '-created_at')

    context = {
        'records': records,
        'start_date': start_date_str,
        'end_date': end_date_str,
        'ma_bithu_search': ma_bithu_search
    }

    return render(request, 'templates_app/reports/mail_envelope_report.html', context)


# Dùng chung cho báo cáo Đóng/Mở TK, Đối chiếu HĐV, đăng ký HĐV/QR
def _find_col(df_cols, candidates):
    """Tìm tên cột trong df theo danh sách ứng viên (case-insensitive, bỏ qua khoảng trắng thừa)."""
    lower_map = {c.strip().lower(): c for c in df_cols}
    for cand in candidates:
        key = cand.strip().lower()
        if key in lower_map:
            return lower_map[key]
    return None
