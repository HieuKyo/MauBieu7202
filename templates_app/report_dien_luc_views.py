"""
Báo cáo Điện lực.
"""
from django.shortcuts import render, redirect
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.views.decorators.http import require_http_methods
from django.contrib import messages
import pandas as pd
import openpyxl
import io
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from datetime import datetime
from .report_views import _read_excel_safe


@login_required
def dien_luc_report_view(request):
    """Giao diện báo cáo Điện lực - xử lý client-side"""
    return render(request, 'templates_app/reports/dien_luc.html')


@login_required
@require_http_methods(["POST"])
def process_dien_luc_report(request):
    """Xử lý báo cáo Thu hộ tiền điện - xuất Excel bảng kê"""
    try:
        uploaded_file = request.FILES.get('dien_luc_file')
        start_date_str = request.POST.get('start_date', '').strip()
        end_date_str = request.POST.get('end_date', '').strip()

        if not uploaded_file:
            messages.error(request, 'Vui lòng chọn file Excel.')
            return redirect('dien_luc_report')

        if not start_date_str or not end_date_str:
            messages.error(request, 'Vui lòng nhập đầy đủ khoảng thời gian.')
            return redirect('dien_luc_report')

        start_date = datetime.strptime(start_date_str, '%Y-%m-%d').date()
        end_date = datetime.strptime(end_date_str, '%Y-%m-%d').date()

        # Đọc file Excel — dùng helper tự xử lý corruption cho .xls cũ
        df = _read_excel_safe(uploaded_file)

        # Chuẩn hóa tên cột (strip whitespace)
        df.columns = [c.strip() for c in df.columns]

        required_cols = ['NGAY_NOP', 'TONG_NOP']
        for col in required_cols:
            if col not in df.columns:
                messages.error(request, f'File thiếu cột bắt buộc: {col}. Các cột hiện có: {", ".join(df.columns)}')
                return redirect('dien_luc_report')

        # Parse NGAY_NOP — thử nhiều format, ưu tiên DD/MM/YYYY
        raw = df['NGAY_NOP']
        if pd.api.types.is_datetime64_any_dtype(raw):
            df['ngay_parsed'] = raw
        else:
            s = raw.astype(str).str.strip().str[:10]
            # Thử lần lượt các format phổ biến
            for fmt in ('%d/%m/%Y', '%Y-%m-%d', '%d-%m-%Y', '%m/%d/%Y'):
                parsed = pd.to_datetime(s, format=fmt, errors='coerce')
                ok = parsed.notna().sum()
                if ok > len(df) * 0.5:
                    df['ngay_parsed'] = parsed
                    break
            else:
                df['ngay_parsed'] = pd.NaT

        df = df.dropna(subset=['ngay_parsed'])
        if df.empty:
            messages.error(request, 'Không thể parse cột NGAY_NOP. Kiểm tra lại định dạng ngày trong file.')
            return redirect('dien_luc_report')

        df['ngay_date'] = df['ngay_parsed'].dt.date

        # Debug: show sample parsed dates and range
        sample_dates = sorted(df['ngay_date'].unique())[:5]
        sample_str = ', '.join(d.strftime('%d/%m/%Y') for d in sample_dates)

        # Filter by date range
        df_filtered = df[(df['ngay_date'] >= start_date) & (df['ngay_date'] <= end_date)]

        if df_filtered.empty:
            messages.warning(request,
                f'Không có dữ liệu trong khoảng {start_date.strftime("%d/%m/%Y")} – {end_date.strftime("%d/%m/%Y")}. '
                f'File có {len(df)} dòng, ngày đầu tiên parse được: {sample_str}')
            return redirect('dien_luc_report')

        df = df_filtered

        # Parse TONG_NOP as numeric — strip spaces/commas nếu là string
        tong_nop_raw = df['TONG_NOP'].astype(str).str.strip().str.replace(',', '', regex=False)
        df['tong_nop_num'] = pd.to_numeric(tong_nop_raw, errors='coerce').fillna(0)

        # Đếm số hóa đơn bằng size (đếm tất cả các dòng, không bỏ NaN)
        size_by_date = df.groupby('ngay_date').size().rename('so_luong')
        sum_by_date = df.groupby('ngay_date')['tong_nop_num'].sum().rename('so_tien')
        grouped = (
            pd.concat([size_by_date, sum_by_date], axis=1)
            .reset_index()
            .sort_values('ngay_date')
        )

        # Build Excel with openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Bảng kê tiền điện'

        # Fit tất cả cột vào 1 trang khi in
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0  # 0 = không giới hạn số trang dọc
        ws.page_setup.orientation = 'portrait'

        # --- Styles ---
        bold_font = Font(name='Times New Roman', bold=True, size=12)
        normal_font = Font(name='Times New Roman', size=12)
        center = Alignment(horizontal='center', vertical='center', wrap_text=True)
        left = Alignment(horizontal='left', vertical='center', wrap_text=True)
        right = Alignment(horizontal='right', vertical='center')

        thin = Side(border_style='thin', color='000000')
        all_border = Border(left=thin, right=thin, top=thin, bottom=thin)

        header_fill = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')

        # Column widths: A(STT), B(Ngày), C(Số HĐ), D(Số tiền), E(Ghi chú)
        ws.column_dimensions['A'].width = 8
        ws.column_dimensions['B'].width = 18
        ws.column_dimensions['C'].width = 22
        ws.column_dimensions['D'].width = 22
        ws.column_dimensions['E'].width = 25

        row = 1

        # Row 1: bank name (left A:C) + CHXHCNVN (right D:E)
        ws.merge_cells(f'A{row}:C{row}')
        c = ws.cell(row=row, column=1,
                    value='NGÂN HÀNG NÔNG NGHIỆP\nVÀ PHÁT TRIỂN NÔNG THÔN VIỆT NAM')
        c.font = Font(name='Times New Roman', bold=True, size=11)
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.row_dimensions[row].height = 32

        ws.merge_cells(f'D{row}:E{row}')
        c2 = ws.cell(row=row, column=4,
                     value='CỘNG HOÀ XÃ HỘI CHỦ NGHĨA VIỆT NAM')
        c2.font = Font(name='Times New Roman', bold=True, size=11)
        c2.alignment = center
        row += 1

        # Row 2: Chi nhánh (left A:C) + Độc lập (right D:E)
        ws.merge_cells(f'A{row}:C{row}')
        c = ws.cell(row=row, column=1, value='CHI NHÁNH GIÁ RAI BẠC LIÊU')
        c.font = Font(name='Times New Roman', bold=True, underline='single', size=11)
        c.alignment = center

        ws.merge_cells(f'D{row}:E{row}')
        c2 = ws.cell(row=row, column=4, value='Độc lập - Tự do - Hạnh phúc')
        c2.font = Font(name='Times New Roman', bold=True, underline='single', size=11)
        c2.alignment = center
        ws.row_dimensions[row].height = 18
        row += 1

        # Empty row
        row += 1

        # Title row
        start_fmt = start_date.strftime('%d/%m/%Y')
        end_fmt = end_date.strftime('%d/%m/%Y')
        title = f'BẢNG KÊ THANH TOÁN HOÁ ĐƠN TIỀN ĐIỆN\nTỪ NGÀY {start_fmt} ĐẾN NGÀY {end_fmt}'
        ws.merge_cells(f'A{row}:E{row}')
        c = ws.cell(row=row, column=1, value=title)
        c.font = Font(name='Times New Roman', bold=True, size=14)
        c.alignment = center
        ws.row_dimensions[row].height = 46
        row += 1

        # Empty row
        row += 1

        # Table header
        headers = ['STT', 'Ngày', 'Số lượng Hoá Đơn', 'Số tiền', 'Ghi Chú']
        for col_idx, h in enumerate(headers, start=1):
            c = ws.cell(row=row, column=col_idx, value=h)
            c.font = bold_font
            c.alignment = center
            c.border = all_border
            c.fill = header_fill
        ws.row_dimensions[row].height = 22
        row += 1

        # Data rows
        total_so_luong = 0
        total_so_tien = 0

        for i, data_row in enumerate(grouped.itertuples(), start=1):
            ngay_str = data_row.ngay_date.strftime('%d/%m/%Y')
            so_luong = int(data_row.so_luong)
            so_tien = float(data_row.so_tien)
            total_so_luong += so_luong
            total_so_tien += so_tien

            values = [i, ngay_str, so_luong, so_tien, '']
            aligns = [center, center, center, right, left]
            for col_idx, (val, aln) in enumerate(zip(values, aligns), start=1):
                c = ws.cell(row=row, column=col_idx, value=val)
                c.font = normal_font
                c.alignment = aln
                c.border = all_border
                if col_idx == 4:
                    c.number_format = '#,##0'
            ws.row_dimensions[row].height = 18
            row += 1

        # Total row
        total_values = ['', 'TỔNG CỘNG', total_so_luong, total_so_tien, '']
        total_aligns = [center, center, center, right, left]
        for col_idx, (val, aln) in enumerate(zip(total_values, total_aligns), start=1):
            c = ws.cell(row=row, column=col_idx, value=val)
            c.font = bold_font
            c.alignment = aln
            c.border = all_border
            if col_idx == 4:
                c.number_format = '#,##0'
        ws.row_dimensions[row].height = 18
        row += 1

        # Empty row before signatures
        row += 1

        # Signature title row: LẬP BẢNG / KIỂM SOÁT
        ws.merge_cells(f'B{row}:C{row}')
        c = ws.cell(row=row, column=2, value='LẬP BẢNG')
        c.font = bold_font
        c.alignment = center

        ws.merge_cells(f'D{row}:E{row}')
        c2 = ws.cell(row=row, column=4, value='KIỂM SOÁT')
        c2.font = bold_font
        c2.alignment = center
        ws.row_dimensions[row].height = 18
        row += 1

        # Sub-label row: (Ký, ghi rõ họ tên) — directly below, no extra gap
        ws.merge_cells(f'B{row}:C{row}')
        c = ws.cell(row=row, column=2, value='(Ký, ghi rõ họ tên)')
        c.font = normal_font
        c.alignment = center

        ws.merge_cells(f'D{row}:E{row}')
        c2 = ws.cell(row=row, column=4, value='(Ký, ghi rõ họ tên)')
        c2.font = normal_font
        c2.alignment = center
        ws.row_dimensions[row].height = 18
        row += 1

        # Empty rows for actual signature space
        ws.row_dimensions[row].height = 50
        row += 1

        # Output
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        filename = f'bang_ke_tien_dien_{start_fmt.replace("/","")}-{end_fmt.replace("/","")}.xlsx'
        response = HttpResponse(
            output.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response

    except Exception as e:
        messages.error(request, f'Lỗi xử lý file: {str(e)}')
        return redirect('dien_luc_report')
