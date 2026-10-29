"""
Báo cáo Thu hộ học phí và điện lực.
"""
from django.shortcuts import render, redirect
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.views.decorators.http import require_http_methods
from django.contrib import messages
import pandas as pd
import openpyxl
import io
import traceback
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from .report_views import _read_excel_safe


# ── Báo cáo Thu hộ học phí và điện lực ────────────────────────────────────────

# Mã DV -> Tên trường (theo phụ lục). Mã DV "759" là Điện lực (xử lý riêng bên dưới).
THU_HO_HOC_PHI_MA_DV_TRUONG = {
    '13211': 'Trường THPT Giá Rai',
    '13212': 'Trường THPT Nguyễn Trung Trực',
    '13213': 'Trường THPT Tân Phong',
    '15005': 'Trường THPT Giá Rai',
    '15013': 'Trường THPT Nguyễn Trung Trực',
    '15014': 'Trường Tiểu Học Phong Thạnh Tây',
    '15142': 'Trường Tiểu Học Phong Tân',
    '15143': 'Trường Tiểu Học Phong Thạnh',
    '15144': 'Trường Tiểu Học Phong Thạnh Đông',
    '15145': 'Trường Tiểu Học Phong Thạnh Tây',
    '15148': 'Trường Tiểu Học Tân Hiệp',
    '15149': 'Trường Tiểu Học Tân Thạnh A',
    '15169': 'Trường THCS Giá Rai A',
    '15170': 'Trường THCS Giá Rai B',
    '15171': 'Trường THCS Hộ Phòng',
    '15187': 'Trường THCS Phong Phú',
    '15188': 'Trường THCS Phong Tân',
    '15189': 'Trường THCS Phong Thạnh',
    '15190': 'Trường THCS Phong Thạnh Đông',
    '15194': 'Trường THCS Tân Hiệp',
    '15195': 'Trường THCS Thạnh Bình',
    '15320': 'Trường Mẫu Giáo Phong Phú',
    '15324': 'Trường Mẫu Giáo Tân Hiệp',
    '15325': 'Trường Tiểu Học Thạnh Bình',
    '15326': 'Trường Mầm Non Tuổi Thơ',
    '14289': 'Trường Tiểu Học Giá Rai A',
    '14290': 'Trường Tiểu Học Giá Rai B',
    '14291': 'Trường Tiểu Học Hộ Phòng A',
    '14292': 'Trường Tiểu Học Hộ Phòng B',
    '14293': 'Trường Tiểu Học Phong Thạnh',
    '14294': 'Trường Tiểu Học Phong Phú A',
    '14295': 'Trường Tiểu Học Phong Phú B',
    '14296': 'Trường Tiểu Học Phong Thạnh Tây',
    '14297': 'Trường Tiểu Học Phong Tân',
    '14298': 'Trường Tiểu Học Thạnh Bình',
    '14300': 'Trường Tiểu Học Tân Thạnh A',
    '14301': 'Trường Tiểu Học Phong Thạnh Đông',
    '15070': 'Trường Mầm Non Họa Mi',
    '15078': 'Trường Mầm Non Hương Sen',
    '15086': 'Trường Mầm Non Sơn Ca 1',
    '15094': 'Trường Mầm Non Phong Thạnh A',
    '15095': 'Trường Mầm Non Phong Thạnh Đông',
    '15096': 'Trường Mầm Non Sơn Ca 2',
    '15097': 'Trường Mẫu Giáo Thạnh Bình',
    '15101': 'Trường Tiểu Học Và THCS Phong Thạnh A',
    '15108': 'Trường Tiểu Học Giá Rai A',
    '15109': 'Trường Tiểu Học Giá Rai B',
    '15110': 'Trường Tiểu Học Hộ Phòng A',
    '15111': 'Trường Tiểu Học Hộ Phòng B',
    '15140': 'Trường Tiểu Học Phong Phú A',
    '15141': 'Trường Tiểu Học Phong Phú B',
}
THU_HO_DIEN_LUC_MA_DV = '759'
THU_HO_HOC_PHI_HEADER_ROW = 9  # Excel row 10 (0-indexed) chứa header STT/Mã DV/Số tiền...


def _thu_ho_normalize_ma_dv(value):
    """Chuẩn hóa giá trị Mã DV về string, bỏ '.0' nếu Excel đọc thành số thực."""
    s = str(value).strip()
    if s.endswith('.0'):
        s = s[:-2]
    return s


@login_required
def thu_ho_hoc_phi_report_view(request):
    """Giao diện báo cáo Thu hộ học phí và điện lực"""
    if 'thu_ho_hoc_phi_data' in request.session:
        del request.session['thu_ho_hoc_phi_data']
    return render(request, 'templates_app/reports/thu_ho_hoc_phi.html')


@login_required
@require_http_methods(["POST"])
def process_thu_ho_hoc_phi_report(request):
    """Đọc file giao dịch, phân loại theo Mã DV (từng trường + điện lực), thống kê số món/số tiền."""
    try:
        uploaded_file = request.FILES.get('data_file')
        if not uploaded_file:
            messages.error(request, 'Vui lòng chọn file Excel.')
            return redirect('thu_ho_hoc_phi_report')

        df = _read_excel_safe(uploaded_file, header=THU_HO_HOC_PHI_HEADER_ROW)
        df.columns = [str(c).strip() for c in df.columns]

        required_cols = ['Mã DV', 'Số tiền']
        for col in required_cols:
            if col not in df.columns:
                messages.error(
                    request,
                    f"File thiếu cột bắt buộc: {col}. Các cột hiện có: {', '.join(df.columns)}"
                )
                return redirect('thu_ho_hoc_phi_report')

        df = df.dropna(subset=['Mã DV'])
        df['_ma_dv'] = df['Mã DV'].apply(_thu_ho_normalize_ma_dv)
        tien_raw = df['Số tiền'].astype(str).str.strip().str.replace(',', '', regex=False)
        df['_so_tien'] = pd.to_numeric(tien_raw, errors='coerce').fillna(0)

        def _classify(ma_dv):
            if ma_dv == THU_HO_DIEN_LUC_MA_DV:
                return 'Điện lực'
            return THU_HO_HOC_PHI_MA_DV_TRUONG.get(ma_dv, f'Không xác định (Mã DV {ma_dv})')

        df['_nhom'] = df['_ma_dv'].apply(_classify)

        grouped = (
            df.groupby('_nhom')
            .agg(so_mon=('_so_tien', 'size'), so_tien=('_so_tien', 'sum'))
            .reset_index()
            .rename(columns={'_nhom': 'nhom'})
        )

        dien_luc_row = grouped[grouped['nhom'] == 'Điện lực']
        truong_rows = grouped[
            (grouped['nhom'] != 'Điện lực') & (~grouped['nhom'].str.startswith('Không xác định'))
        ].sort_values('nhom')
        unknown_rows = grouped[grouped['nhom'].str.startswith('Không xác định')].sort_values('nhom')

        truong_list = truong_rows.to_dict('records')
        unknown_list = unknown_rows.to_dict('records')
        dien_luc = dien_luc_row.to_dict('records')[0] if not dien_luc_row.empty else {'so_mon': 0, 'so_tien': 0}

        tong_so_mon = int(grouped['so_mon'].sum())
        tong_so_tien = float(grouped['so_tien'].sum())

        result = {
            'truong_list': truong_list,
            'unknown_list': unknown_list,
            'dien_luc': dien_luc,
            'tong_so_mon': tong_so_mon,
            'tong_so_tien': tong_so_tien,
            'truong_count': len(truong_list),
            'unknown_count': len(unknown_list),
        }
        request.session['thu_ho_hoc_phi_data'] = result

        return render(request, 'templates_app/reports/thu_ho_hoc_phi.html', {'result': result})

    except Exception as e:
        traceback.print_exc()
        messages.error(request, f"Đã xảy ra lỗi: {e}")
        return redirect('thu_ho_hoc_phi_report')


@login_required
def download_thu_ho_hoc_phi_report(request):
    """Xuất Excel bảng thống kê từ dữ liệu đã xử lý (lưu trong session)."""
    result = request.session.get('thu_ho_hoc_phi_data')
    if not result:
        messages.error(request, 'Chưa có dữ liệu để tải. Vui lòng xử lý báo cáo trước.')
        return redirect('thu_ho_hoc_phi_report')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Thu ho hoc phi va dien luc'

    bold_font = Font(name='Times New Roman', bold=True, size=12)
    normal_font = Font(name='Times New Roman', size=12)
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left = Alignment(horizontal='left', vertical='center', wrap_text=True)
    right = Alignment(horizontal='right', vertical='center')
    thin = Side(border_style='thin', color='000000')
    all_border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')
    dien_luc_fill = PatternFill(start_color='FFF2CC', end_color='FFF2CC', fill_type='solid')

    ws.column_dimensions['A'].width = 6
    ws.column_dimensions['B'].width = 45
    ws.column_dimensions['C'].width = 14
    ws.column_dimensions['D'].width = 18

    ws.merge_cells('A1:D1')
    c = ws.cell(row=1, column=1, value='THỐNG KÊ THU HỘ HỌC PHÍ VÀ ĐIỆN LỰC')
    c.font = Font(name='Times New Roman', bold=True, size=14)
    c.alignment = center
    ws.row_dimensions[1].height = 24

    row = 3
    headers = ['STT', 'Đơn vị', 'Số món', 'Số tiền']
    for col_idx, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=col_idx, value=h)
        c.font = bold_font
        c.alignment = center
        c.border = all_border
        c.fill = header_fill
    row += 1

    stt = 1
    for item in result['truong_list']:
        values = [stt, item['nhom'], int(item['so_mon']), float(item['so_tien'])]
        aligns = [center, left, center, right]
        for col_idx, (val, aln) in enumerate(zip(values, aligns), start=1):
            c = ws.cell(row=row, column=col_idx, value=val)
            c.font = normal_font
            c.alignment = aln
            c.border = all_border
            if col_idx == 4:
                c.number_format = '#,##0'
        row += 1
        stt += 1

    for item in result['unknown_list']:
        values = [stt, item['nhom'], int(item['so_mon']), float(item['so_tien'])]
        aligns = [center, left, center, right]
        for col_idx, (val, aln) in enumerate(zip(values, aligns), start=1):
            c = ws.cell(row=row, column=col_idx, value=val)
            c.font = normal_font
            c.alignment = aln
            c.border = all_border
            if col_idx == 4:
                c.number_format = '#,##0'
        row += 1
        stt += 1

    dl = result['dien_luc']
    values = [stt, 'Điện lực', int(dl['so_mon']), float(dl['so_tien'])]
    aligns = [center, left, center, right]
    for col_idx, (val, aln) in enumerate(zip(values, aligns), start=1):
        c = ws.cell(row=row, column=col_idx, value=val)
        c.font = bold_font
        c.alignment = aln
        c.border = all_border
        c.fill = dien_luc_fill
        if col_idx == 4:
            c.number_format = '#,##0'
    row += 1

    total_values = ['', 'TỔNG CỘNG', result['tong_so_mon'], result['tong_so_tien']]
    aligns = [center, center, center, right]
    for col_idx, (val, aln) in enumerate(zip(total_values, aligns), start=1):
        c = ws.cell(row=row, column=col_idx, value=val)
        c.font = bold_font
        c.alignment = aln
        c.border = all_border
        if col_idx == 4:
            c.number_format = '#,##0'

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    response = HttpResponse(
        output.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="ThuHoHocPhiDienLuc.xlsx"'
    return response
