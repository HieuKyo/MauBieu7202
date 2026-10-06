"""
Test đối chiếu sao kê DPTB18 với file CSP / MSPH02.
Dữ liệu tự tạo theo đúng định dạng file thật (tên, số TK là giả).
"""
import os
import shutil
import tempfile
from io import BytesIO

import pandas as pd
from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase, TestCase
from django.urls import reverse

from .bank_statement_enricher import _valid_name, enrich
from .bank_statement_parser import BankStatementParser
from .models import AccountName, BankStatement


def _dptb18_rows():
    base = {'trcdnm': 'Tiền gửi không kỳ hạn', 'trcd': 'X101', 'aftrbal': 1000000, 'tomgntno': '', 'toacctno': ''}
    rows = [
        # 0: nhận liên NH — khớp MSPH02 chiều đến qua thrref
        {'trdt': '2026-01-02', 'acctccyamt': 500000, 'thrref': '5485955', 'rem': 'NGUYEN VAN A chuyen tien'},
        # 1: chuyển liên NH — khớp MSPH02 chiều đi qua thrref
        {'trdt': '2026-01-02', 'acctccyamt': -50000, 'thrref': '140294708',
         'rem': 'Vietcombank:0891000644791:TRAN B chuyen tien'},
        # 2: nhận qua NAPAS — khớp CSP chiều đến qua mã 6 số đầu rem
        {'trdt': '2026-01-03', 'acctccyamt': 200000, 'thrref': 'BANKNETIBFT',
         'rem': '1000A87202 - 037620-LE C chuyen tien;MCC;20260101103736;0948487626;970448'},
        # 3: chuyển qua NAPAS — khớp CSP chiều đi (dòng Debit) qua mã 6 số đầu rem
        {'trdt': '2026-01-03', 'acctccyamt': -150000, 'thrref': 'BANKNETIBFT',
         'rem': '941556-OCB;0947365707;TRAN B chuyen tien'},
        # 4: nội bộ Agribank — không có trong file đối soát
        {'trdt': '2026-01-04', 'acctccyamt': 100000, 'thrref': 'TR20260104123456',
         'rem': 'MB(123456)(PHAM D chuyen tien)'},
    ]
    return [{**base, **r} for r in rows]


_MSPH02_BASE = {'ten_nguoi_chuyen': '', 'ten_nguoi_nhan': '', 'tk_chuyen': '', 'tk_nhan': '',
                'noi_dung': '', 'ngay_giao_dich': '02/01/2026 10:00:00'}


def _aux_frames():
    msph02_den = pd.DataFrame([
        {**_MSPH02_BASE, 'trace': '005485955', 'sttlamt': '500,000', 'rspcd': 'Đã trả KH',
         'nh_nhan_gui_nhan': 'Hoi So NH NT Viet Nam', 'ten_nguoi_chuyen': 'NGUYEN VAN A', 'tk_chuyen': '0791000073168'},
        # Thất bại → bỏ qua, không tính là không khớp
        {**_MSPH02_BASE, 'trace': '005485956', 'sttlamt': '700,000', 'rspcd': 'HT Lỗi',
         'nh_nhan_gui_nhan': 'Hoi So NH NT Viet Nam'},
        # Thành công nhưng không có trong DPTB18 → không khớp
        {**_MSPH02_BASE, 'trace': '000000999', 'sttlamt': '300,000', 'rspcd': 'Đã trả KH',
         'nh_nhan_gui_nhan': 'NH Quân đội Hà Nội', 'ten_nguoi_chuyen': 'HO VAN F', 'tk_chuyen': '1234567890'},
    ])
    msph02_di = pd.DataFrame([
        {**_MSPH02_BASE, 'trace': '140294708', 'sttlamt': '50,000', 'rspcd': 'Hoàn thành',
         'nh_nhan_gui_nhan': 'NH Quân đội Hà Nội', 'ten_nguoi_nhan': 'TRAN THI B', 'tk_nhan': '0891000644791'},
    ])
    csp_base = {'TRANSACTION TIME': '2026-01-01 10:37:36', 'CONTENT': ''}
    csp_den = pd.DataFrame([
        # TRACE IPCAS bị mất số 0 đầu (037620 → 37620)
        {**csp_base, 'FROM ACCOUNT': '948487626', 'SENDER NAME': 'LE VAN C', 'AMOUNT': '200000',
         'STATUS': '00 - Approved', 'TRACE IPCAS': '37620', 'FROM BANK CODE': '970448-OCB'},
        # Không có trong DPTB18 → không khớp; nhưng dùng để đối chiếu tên cho CSP chiều đi
        {**csp_base, 'FROM ACCOUNT': '947365707', 'SENDER NAME': 'LU THI E', 'AMOUNT': '90000',
         'STATUS': '00 - Approved', 'TRACE IPCAS': '111111', 'FROM BANK CODE': '970448-OCB'},
    ])
    di_base = {'TRANSACTION TIME': '2026-01-03 09:24:10', 'AMOUNT': '150000',
               'REMARK': 'OCB;0947365707;TRAN B chuyen tien', 'TO_ACCOUNT': '', 'TO_BANK_CODE': ''}
    csp_di = pd.DataFrame([
        {**di_base, 'TRANCODE': '369-Transfer Debit', 'TRACE': '941556', 'RESPONSE': '1-Approved'},
        {**di_base, 'TRANCODE': '91-Transfer Credit', 'TRACE': '177233', 'RESPONSE': '1-Approved',
         'TO_ACCOUNT': '947365707', 'TO_BANK_CODE': '970448-OCB'},
        # Debit bị từ chối → bỏ qua
        {**di_base, 'TRANCODE': '369-Transfer Debit', 'TRACE': '941557', 'RESPONSE': '59-Insufficient funds'},
    ])
    return {'msph02_den': msph02_den, 'msph02_di': msph02_di, 'csp_den': csp_den, 'csp_di': csp_di}


def _xlsx_bytes(df):
    buf = BytesIO()
    df.to_excel(buf, index=False)
    return buf.getvalue()


class EnrichTests(SimpleTestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, self.tmp)
        self.aux_paths = {}
        for kind, df in _aux_frames().items():
            path = os.path.join(self.tmp, f'{kind}.xlsx')
            df.to_excel(path, index=False)
            self.aux_paths[kind] = path
        dptb_path = os.path.join(self.tmp, 'dptb18.xlsx')
        pd.DataFrame(_dptb18_rows()).to_excel(dptb_path, index=False)
        self.parser = BankStatementParser(dptb_path)
        self.assertEqual(self.parser.validate_file(), (True, ''))

    def _process(self, aux_paths):
        enrichment, unmatched = enrich(self.parser.df, aux_paths)
        return self.parser.process(enrichment=enrichment), unmatched

    def test_all_files(self):
        rows, unmatched = self._process(self.aux_paths)
        got = [(r['nguon'], r['ngan_hang'], r['so_tai_khoan'], r['ten_nguoi']) for r in rows[:4]]
        self.assertEqual(got, [
            ('MSPH02 chiều đến', 'Vietcombank', '0791000073168', 'NGUYEN VAN A'),
            ('MSPH02 chiều đi', 'MB Bank', '0891000644791', 'TRAN THI B'),
            # Số TK giữ bản có số 0 đầu từ rem
            ('CSP chiều đến', 'OCB', '0948487626', 'LE VAN C'),
            ('CSP chiều đi (tên từ CSP chiều đến)', 'OCB', '0947365707', 'LU THI E'),
        ])
        # Khớp file đối soát → loại liên ngân hàng theo chiều tiền; loại MCC giữ nguyên tên
        self.assertEqual([r['ghi_chu'] for r in rows[:4]], [
            'Nhận chuyển khoản liên ngân hàng',
            'Chuyển khoản đi khác ngân hàng',
            'Nhận thanh toán MCC',
            'Chuyển khoản đi khác ngân hàng',
        ])
        # Giao dịch nội bộ: không có trong file đối soát → giữ kết quả parse rem
        self.assertEqual((rows[4]['nguon'], rows[4]['ten_nguoi']), ('', 'PHAM D'))
        self.assertEqual(sorted((u['file'], u['trace']) for u in unmatched), [
            ('CSP chiều đến', '111111'),
            ('MSPH02 chiều đến', '000000999'),
        ])

    def test_without_aux_files_behaves_as_before(self):
        rows, unmatched = self._process({})
        self.assertEqual(unmatched, [])
        self.assertTrue(all(r['nguon'] == '' for r in rows))

    def test_cross_check_requires_same_bank(self):
        paths = dict(self.aux_paths)
        frames = _aux_frames()
        frames['csp_den'].loc[1, 'FROM BANK CODE'] = '970416-ACB'
        paths['csp_den'] = os.path.join(self.tmp, 'csp_den_acb.xlsx')
        frames['csp_den'].to_excel(paths['csp_den'], index=False)
        rows, _ = self._process(paths)
        self.assertEqual((rows[3]['nguon'], rows[3]['ten_nguoi']), ('CSP chiều đi', ''))

    def test_missing_column_raises(self):
        bad = os.path.join(self.tmp, 'bad.xlsx')
        _aux_frames()['csp_den'].drop(columns=['TRACE IPCAS']).to_excel(bad, index=False)
        with self.assertRaisesMessage(ValueError, 'File CSP chiều đến thiếu cột: TRACE IPCAS'):
            enrich(self.parser.df, {'csp_den': bad})

    def test_valid_name_rejects_bank_names(self):
        self.assertFalse(_valid_name('SACOMBANK'))
        self.assertFalse(_valid_name('MoMo'))
        self.assertFalse(_valid_name(''))
        self.assertTrue(_valid_name('Lam Chi Dinh'))


class ClassifyTests(SimpleTestCase):
    def test_sms_ott_fees_are_service_fees(self):
        parser = BankStatementParser('unused.xlsx')
        base = {'trcd': 'X201', 'trcdnm': 'Rút tiền (tiền gửi KKH)', 'husrid': '1000SMS10', 'acctccyamt': -8800}
        for rem in ['Phi Tin nhan OTT DV Agribank Plus Thang 02/2026',
                    'Phi DV SMS Banking T12/2025 SDT:0947095774 26 tin']:
            self.assertEqual(parser.classify_transaction({**base, 'rem': rem}), 'Phí dịch vụ', rem)

    def test_payment_hub_api_is_interbank(self):
        parser = BankStatementParser('unused.xlsx')
        base = {'trcd': 'X101', 'trcdnm': 'Tiền gửi không kỳ hạn', 'husrid': '7202API2', 'thrref': '7866082',
                'lclbrnm': 'Agribank CN Giá Rai Bạc Liêu', 'rem': 'NGUYEN THI NHU PHUONG chuyen tien'}
        self.assertEqual(parser.classify_transaction({**base, 'acctccyamt': 1000000}),
                         'Nhận chuyển khoản liên ngân hàng')
        self.assertEqual(parser.classify_transaction(
            {**base, 'trcd': 'X201', 'trcdnm': 'Rút tiền (tiền gửi KKH)', 'acctccyamt': -50000}),
            'Chuyển khoản đi khác ngân hàng')

    def test_osb_both_directions(self):
        parser = BankStatementParser('unused.xlsx')
        base = {'trcd': 'X101', 'trcdnm': 'Tiền gửi không kỳ hạn', 'husrid': '7202OSB',
                'thrref': 'BB20260102316607', 'tomgntno': '7202205203745',
                'lclbrnm': 'Agribank CN Giá Rai Bạc Liêu', 'rem': 'LAM CHI DINH chuyen tien'}
        self.assertEqual(parser.classify_transaction({**base, 'acctccyamt': 93000}), 'Giao dịch OSB')
        # Chiều đi có trcdnm "Rút tiền" trước đây bị nhầm thành "Rút tiền mặt"
        self.assertEqual(parser.classify_transaction(
            {**base, 'trcd': 'X201', 'trcdnm': 'Rút tiền (tiền gửi KKH)', 'acctccyamt': -500000}),
            'Giao dịch OSB')


class FillNamesByAccountTests(SimpleTestCase):
    def test_outgoing_internal_gets_name_from_incoming(self):
        tmp = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, tmp)
        base = {'trcd': 'X101', 'trcdnm': 'Tiền gửi không kỳ hạn', 'aftrbal': 1000000,
                'acctno': '7202205203745', 'lclbrnm': 'Agribank CN Giá Rai Bạc Liêu'}
        rows = [
            # Nhận từ TK 7207205188157 → biết tên chủ TK
            {**base, 'trdt': '2026-02-05', 'acctccyamt': 5000000, 'tomgntno': '7207205188157', 'toacctno': '',
             'rem': 'MB(366304)(HUYNH VAN CONG chuyen tien)'},
            # Chuyển đi TK 7207205188157 → rem chỉ có tên chủ TK sao kê
            {**base, 'trdt': '2026-02-06', 'acctccyamt': -5000000, 'tomgntno': '', 'toacctno': '7207205188157',
             'rem': 'MB(896901)(LAM CHI DINH chuyen tien)'},
        ]
        path = os.path.join(tmp, 'dptb18.xlsx')
        pd.DataFrame(rows).to_excel(path, index=False)
        parser = BankStatementParser(path)
        parser.validate_file()
        result = parser.process()
        self.assertEqual((result[0]['ten_nguoi'], result[0]['nguon']), ('HUYNH VAN CONG', ''))
        self.assertEqual(
            (result[1]['so_tai_khoan'], result[1]['ten_nguoi'], result[1]['nguon']),
            ('7207205188157', 'HUYNH VAN CONG', 'Đối chiếu số TK'))


class BankStatementViewTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user('tester', password='x')
        self.client.force_login(self.user)

    def test_upload_result_export(self):
        files = {'statement_file': SimpleUploadedFile('dptb18.xlsx', _xlsx_bytes(pd.DataFrame(_dptb18_rows())))}
        for kind, df in _aux_frames().items():
            files[f'{kind}_file'] = SimpleUploadedFile(f'{kind}.xlsx', _xlsx_bytes(df))

        response = self.client.post(reverse('bank_statement_upload'), files)
        statement = BankStatement.objects.get(uploaded_by=self.user)
        self.assertRedirects(response, reverse('bank_statement_result', args=[statement.id]))
        self.assertEqual(len(statement.unmatched_rows), 2)
        self.assertEqual(statement.transactions.exclude(source='').count(), 4)

        page = self.client.get(reverse('bank_statement_result', args=[statement.id]))
        self.assertContains(page, 'Dòng không khớp')
        self.assertContains(page, 'Cập nhật tên từ file Excel')
        self.assertContains(page, 'MSPH02 chiều đến')
        # Tên từ file đối soát được giữ trong bảng TK nhiều lần (không bị lọc bởi _is_proper_name)
        names = {a['account_number']: a['name'] for a in page.context['frequent_recipients']}
        self.assertEqual(names['0947365707'], 'LU THI E')

        export = self.client.get(reverse('bank_statement_export', args=[statement.id]))
        sheets = pd.read_excel(BytesIO(export.content), sheet_name=None)
        self.assertIn('Không khớp', sheets)
        self.assertEqual(len(sheets['Không khớp']), 2)
        self.assertIn('Nguồn', sheets['Chi tiết'].columns)
        stats = sheets['Thống kê'].iloc[:, 0].astype(str).tolist()
        self.assertIn('THỐNG KÊ THEO NGUỒN', stats)
        self.assertIn('MSPH02 chiều đến', stats)
        self.assertIn('Nội dung GD', stats)

    def test_export_name_formula_for_unknown_account(self):
        # GD chuyển đi liên NH, không có file đối soát → có Số TK nhưng không có tên
        rows = [{'trdt': '2026-01-02', 'acctccyamt': -100000, 'aftrbal': 900000, 'trcd': 'X201',
                 'trcdnm': 'Rút tiền (tiền gửi KKH)', 'thrref': '149045486', 'tomgntno': '', 'toacctno': '',
                 'rem': 'OCB;0123456789;chuyen tien'}]
        files = {'statement_file': SimpleUploadedFile('dptb18.xlsx', _xlsx_bytes(pd.DataFrame(rows)))}
        self.client.post(reverse('bank_statement_upload'), files)
        statement = BankStatement.objects.get(uploaded_by=self.user)
        export = self.client.get(reverse('bank_statement_export', args=[statement.id]))
        if os.environ.get('BANK_EXPORT_DUMP'):
            with open(os.environ['BANK_EXPORT_DUMP'], 'wb') as f:
                f.write(export.content)

        from openpyxl import load_workbook
        wb = load_workbook(BytesIO(export.content))
        detail, accounts = wb['Chi tiết'], wb['TK giao dịch nhiều lần']
        self.assertEqual((detail['F2'].value, detail['G2'].value), ('OCB', '0123456789'))
        self.assertEqual((accounts['B3'].value, accounts['C3'].value, accounts['D3'].value),
                         ('OCB', '0123456789', None))
        self.assertEqual(
            detail['H2'].value,
            "=IFERROR(LOOKUP(2,1/(('TK giao dịch nhiều lần'!$B$3:$B$3=F2)*('TK giao dịch nhiều lần'!$C$3:$C$3=G2)"
            "*('TK giao dịch nhiều lần'!$D$3:$D$3<>\"\")),'TK giao dịch nhiều lần'!$D$3:$D$3),\"\")")

    def test_import_names_saves_book_and_reuses(self):
        rows = [{'trdt': '2026-01-02', 'acctccyamt': -100000, 'aftrbal': 900000, 'trcd': 'X201',
                 'trcdnm': 'Rút tiền (tiền gửi KKH)', 'thrref': '149045486', 'tomgntno': '', 'toacctno': '',
                 'rem': 'OCB;0123456789;chuyen tien'}]

        def upload():
            files = {'statement_file': SimpleUploadedFile('dptb18.xlsx', _xlsx_bytes(pd.DataFrame(rows)))}
            self.client.post(reverse('bank_statement_upload'), files)
            return BankStatement.objects.filter(uploaded_by=self.user).latest('id')

        statement = upload()
        export = self.client.get(reverse('bank_statement_export', args=[statement.id]))

        # Người dùng điền tên vào ô vàng rồi lưu file
        from openpyxl import load_workbook
        wb = load_workbook(BytesIO(export.content))
        wb['TK giao dịch nhiều lần']['D3'] = 'NGUYEN VAN X'
        filled = BytesIO()
        wb.save(filled)

        url = reverse('bank_statement_import_names', args=[statement.id])
        self.client.post(url, {'names_file': SimpleUploadedFile('filled.xlsx', filled.getvalue())})
        book = AccountName.objects.get()
        self.assertEqual((book.account_key, book.bank_key, book.name, book.updated_by),
                         ('123456789', 'ocb', 'NGUYEN VAN X', self.user))
        trans = statement.transactions.get()
        self.assertEqual((trans.beneficiary_name, trans.source), ('NGUYEN VAN X', 'Danh bạ'))

        # Upload lại cùng file → không có tên mới
        response = self.client.post(url, {'names_file': SimpleUploadedFile('filled.xlsx', filled.getvalue())},
                                    follow=True)
        self.assertContains(response, 'Không tìm thấy tên mới nào trong file.')

        # Sao kê sau: tài khoản trùng tự có tên từ danh bạ
        trans2 = upload().transactions.get()
        self.assertEqual((trans2.beneficiary_name, trans2.source), ('NGUYEN VAN X', 'Danh bạ'))

    def test_import_names_rejects_wrong_file(self):
        statement = BankStatement.objects.create(file_name='x', total_transactions=0, uploaded_by=self.user)
        response = self.client.post(
            reverse('bank_statement_import_names', args=[statement.id]),
            {'names_file': SimpleUploadedFile('x.xlsx', _xlsx_bytes(pd.DataFrame({'a': [1]})))}, follow=True)
        self.assertContains(response, 'File không đúng')
        self.assertFalse(AccountName.objects.exists())
