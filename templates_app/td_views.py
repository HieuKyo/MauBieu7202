"""
Views cho module phân tích danh mục tín dụng MSIT80 — hoàn toàn offline LAN.
"""
import io
import json
import os
import tempfile
from datetime import date, timedelta

from django.contrib.auth.decorators import login_required
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.views.decorators.http import require_POST

from .models import FTPConfig, TDSnapshot
from .models import Template as WordTemplate
from .td_analysis import (
    analyze_portfolio,
    analysis_to_template_vars,
    compute_so_sanh,
    load_msit80,
)
from .utils import render_word_template


# ---------------------------------------------------------------------------
# 1. Upload & phân tích
# ---------------------------------------------------------------------------

@login_required
def td_upload_view(request):
    """GET: form upload. POST (AJAX): xử lý MSIT80 → lưu snapshot → session → JSON."""
    if request.method != 'POST':
        return render(request, 'templates_app/td_upload.html')

    uploaded = request.FILES.get('msit80_file')
    if not uploaded:
        return JsonResponse({'error': 'Chưa chọn file'}, status=400)

    suffix = '.xlsx' if uploaded.name.lower().endswith('.xlsx') else '.xls'
    tmp_path = None
    try:
        with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as tmp:
            for chunk in uploaded.chunks():
                tmp.write(chunk)
            tmp_path = tmp.name

        ftp_configs = list(FTPConfig.objects.values())
        df = load_msit80(tmp_path)
        analysis = analyze_portfolio(df, ftp_configs=ftp_configs)

        # Lưu records riêng (nặng) để không làm session quá lớn khi gọi API
        records = analysis.pop('records', [])
        request.session['td_analysis'] = analysis
        request.session['td_records'] = records
        request.session['td_file_name'] = uploaded.name

        # Xác định ngày file từ dsbsdt đầu tiên (fallback = hôm nay)
        file_date = date.today()
        if records and records[0].get('dsbsdt'):
            try:
                file_date = date.fromisoformat(str(records[0]['dsbsdt'])[:10])
            except Exception:
                pass

        # Lưu snapshot (mỗi file_date chỉ giữ lại bản mới nhất)
        kpi = analysis.get('kpi', {})
        TDSnapshot.objects.update_or_create(
            file_date=file_date,
            defaults={
                'uploaded_by': request.user,
                'file_name':   uploaded.name,
                'so_lds':      kpi.get('so_lds', 0),
                'so_kh':       kpi.get('so_kh', 0),
                'so_hd':       kpi.get('so_hd', 0),
                'tong_du_no':  int(kpi.get('tong_du_no', 0)),
                'qua_han':     kpi.get('qua_han', 0),
                'nhom2':       kpi.get('nhom2', 0),
                'nhom35':      kpi.get('nhom35', 0),
                'kh_dac_biet': kpi.get('kh_dac_biet', 0),
                'summary_json': {
                    'tong_du_no_ty': kpi.get('tong_du_no_ty', 0),
                    'so_lds':        kpi.get('so_lds', 0),
                    'so_kh':         kpi.get('so_kh', 0),
                    'so_hd':         kpi.get('so_hd', 0),
                    'qua_han':       kpi.get('qua_han', 0),
                    'nhom2':         kpi.get('nhom2', 0),
                    'nhom35':        kpi.get('nhom35', 0),
                    'kh_dac_biet':   kpi.get('kh_dac_biet', 0),
                },
            },
        )

        return JsonResponse({'ok': True, 'redirect': '/tin-dung/dashboard/'})

    except ValueError as e:
        return JsonResponse({'error': str(e)}, status=400)
    except Exception as e:
        return JsonResponse({'error': f'Lỗi phân tích: {str(e)}'}, status=500)
    finally:
        if tmp_path and os.path.exists(tmp_path):
            os.unlink(tmp_path)


# ---------------------------------------------------------------------------
# 2. Dashboard tổng quan
# ---------------------------------------------------------------------------

@login_required
def td_dashboard_view(request):
    """Hiển thị dashboard phân tích danh mục tín dụng từ session."""
    analysis = request.session.get('td_analysis')
    if not analysis:
        return redirect('td_analysis')

    records = request.session.get('td_records', [])
    file_name = request.session.get('td_file_name', '')

    cbtd_list = analysis.get('cbtd_list', [])
    goi_list  = analysis.get('goi_list', [])

    word_templates = WordTemplate.objects.filter(is_active=True).order_by('name')

    # Lấy danh sách snapshot để so sánh kỳ
    snapshots = list(
        TDSnapshot.objects.order_by('-file_date').values(
            'id', 'file_date', 'file_name', 'so_lds', 'so_kh', 'so_hd',
            'tong_du_no', 'qua_han', 'nhom2', 'nhom35', 'kh_dac_biet', 'summary_json'
        )[:60]
    )

    ctx = {
        'analysis':       analysis,
        'kpi':            analysis.get('kpi', {}),
        'hom_nay':        analysis.get('hom_nay', []),
        'can_bo_list':    analysis.get('can_bo_tin_dung', []),
        'canh_bao_kh':    analysis.get('canh_bao_khach_hang', []),
        'tong_ket':       analysis.get('tong_ket_van_ban', {}),
        'phan_loai_no':   analysis.get('phan_loai_no', {}),
        'rui_ro':         analysis.get('rui_ro', {}),
        'qua_han_detail': analysis.get('qua_han', {}),
        'hop_dong':       analysis.get('hop_dong', {}),
        'cbtd_list':      cbtd_list,
        'goi_list':       goi_list,
        'total_records':  len(records),
        'file_name':      file_name,
        'word_templates': word_templates,
        'snapshots_json': json.dumps(snapshots, default=str),
        'today':          date.today().strftime('%d/%m/%Y'),
    }
    return render(request, 'templates_app/td_dashboard.html', ctx)


# ---------------------------------------------------------------------------
# 3. Filter API (AJAX)
# ---------------------------------------------------------------------------

@login_required
def td_filter_api(request):
    """POST AJAX: lọc records theo 6 chiều + Bộ lọc Top."""
    if request.method != 'POST':
        return JsonResponse({'error': 'Method not allowed'}, status=405)

    records = request.session.get('td_records', [])
    if not records:
        return JsonResponse({'error': 'Chưa có dữ liệu'}, status=400)

    try:
        body = json.loads(request.body)
    except Exception:
        body = {}

    f_cbtd     = body.get('cbtd', '')
    f_nhomno   = body.get('nhomno', '')
    f_trangthai= body.get('trangthai', '')
    f_rui_ro   = body.get('rui_ro', '')
    f_goi      = body.get('goi', '')
    f_loai_kh  = body.get('loai_kh', '')
    f_top_field= body.get('top_field', '')
    f_top_n    = int(body.get('top_n', 20))

    filtered = records

    if f_cbtd:
        filtered = [r for r in filtered if str(r.get('ofcnm', '')).strip() == f_cbtd.strip()]
    if f_nhomno:
        grp_map = {'1': [1], '2': [2], '3': [3], '4': [4], '5': [5], '35': [3, 4, 5]}
        targets = grp_map.get(f_nhomno, [])
        if targets:
            filtered = [r for r in filtered if r.get('grpno') in targets]
    if f_trangthai:
        filtered = [r for r in filtered if r.get('_trangthai') == f_trangthai]
    if f_rui_ro:
        filtered = [r for r in filtered if r.get('_level') == f_rui_ro]
    if f_goi:
        filtered = [r for r in filtered if str(r.get('udpcd1', '')) == f_goi]
    if f_loai_kh:
        if f_loai_kh == 'Cá nhân':
            filtered = [r for r in filtered if 'cá nhân' in str(r.get('custtpnm', '')).lower()]
        elif f_loai_kh == 'Tổ chức':
            filtered = [r for r in filtered if 'cá nhân' not in str(r.get('custtpnm', '')).lower()]

    if f_top_field and filtered:
        reverse = f_top_field not in ('_ngay_mon',)
        try:
            filtered = sorted(
                [r for r in filtered if r.get(f_top_field) is not None],
                key=lambda r: r.get(f_top_field, 0),
                reverse=reverse
            )[:f_top_n]
        except Exception:
            pass

    total_bal = sum(float(r.get('dsbsbal', 0)) for r in filtered)

    return JsonResponse({
        'count':        len(filtered),
        'total_bal_ty': round(total_bal / 1e9, 3),
        'records':      filtered[:500],
    })


# ---------------------------------------------------------------------------
# 4. Xuất Excel (danh sách đang lọc)
# ---------------------------------------------------------------------------

@login_required
@require_POST
def td_export_excel_view(request):
    """POST: nhận danh sách record IDs (dsbsseq) từ client, xuất Excel."""
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        return JsonResponse({'error': 'openpyxl chưa được cài đặt'}, status=500)

    records = request.session.get('td_records', [])
    if not records:
        return JsonResponse({'error': 'Chưa có dữ liệu'}, status=400)

    try:
        body = json.loads(request.body)
        ids  = set(body.get('ids', []))
    except Exception:
        ids = set()

    if ids:
        rows = [r for r in records if str(r.get('dsbsseq', '')) in ids]
    else:
        rows = records

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Danh sách tín dụng'

    header_fill = PatternFill('solid', fgColor='1F6FD0')
    bold_white  = Font(bold=True, color='FFFFFF')

    headers = [
        'STT', 'Tên khách hàng', 'Cán bộ TD', 'Số hợp đồng', 'Dư nợ (đồng)',
        'Lãi suất (%)', 'Nhóm nợ', 'Ngày giải ngân', 'Ngày đáo hạn khoản',
        'Lịch trả lãi', 'Lịch trả gốc', 'Gói ưu đãi',
        'Điểm RR', 'Mức RR', 'Trạng thái', 'Cảnh báo', 'NIM (%)',
    ]
    ws.append(headers)
    for col, cell in enumerate(ws[1], 1):
        cell.fill = header_fill
        cell.font = bold_white
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    level_fills = {
        'Xanh':   PatternFill('solid', fgColor='D4EDDA'),
        'Vàng':   PatternFill('solid', fgColor='FFF3CD'),
        'Cam':    PatternFill('solid', fgColor='FFE4B5'),
        'Đỏ':     PatternFill('solid', fgColor='F8D7DA'),
        'Đỏ đậm': PatternFill('solid', fgColor='C0392B'),
    }

    for i, r in enumerate(rows, 1):
        nim_val = r.get('_nim')
        ws.append([
            i,
            str(r.get('custnm', '')),
            str(r.get('ofcnm', '')),
            str(r.get('apprseq', '')),
            float(r.get('dsbsbal', 0)),
            float(r.get('sprd', 0)),
            int(r.get('grpno', 1)),
            str(r.get('dsbsdt', '')),
            str(r.get('dsbsmatdt', '')),
            str(r.get('nxtintschddt', '')),
            str(r.get('nxtrpmtschddt', '')),
            str(r.get('udpcd1', '')),
            int(r.get('_score', 0)),
            str(r.get('_level', '')),
            str(r.get('_trangthai', '')),
            str(r.get('_canh_bao', '')),
            round(nim_val, 2) if nim_val is not None else '',
        ])
        lvl = str(r.get('_level', ''))
        if lvl in level_fills:
            for cell in ws[ws.max_row]:
                cell.fill = level_fills[lvl]

    col_widths = [5, 30, 20, 18, 18, 12, 8, 14, 14, 14, 14, 15, 8, 10, 20, 15, 10]
    for col, w in enumerate(col_widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    resp = HttpResponse(
        buf.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    resp['Content-Disposition'] = 'attachment; filename="danh_sach_tin_dung.xlsx"'
    return resp


# ---------------------------------------------------------------------------
# 5. Xuất Word
# ---------------------------------------------------------------------------

@login_required
@require_POST
def td_export_word_view(request):
    """POST: lấy phân tích từ session → render Word template → trả file .docx."""
    analysis = request.session.get('td_analysis')
    if not analysis:
        return JsonResponse({'error': 'Chưa có dữ liệu phân tích. Upload file trước.'}, status=400)

    template_id = request.POST.get('template_id')
    if not template_id:
        return JsonResponse({'error': 'Thiếu template_id'}, status=400)

    word_template = get_object_or_404(WordTemplate, pk=template_id, is_active=True)
    context = analysis_to_template_vars(analysis)

    try:
        docx_bytes = render_word_template(word_template.file.path, context)
    except Exception as e:
        return JsonResponse({'error': f'Lỗi xuất Word: {str(e)}'}, status=500)

    resp = HttpResponse(
        docx_bytes,
        content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
    )
    resp['Content-Disposition'] = 'attachment; filename="bao_cao_tin_dung.docx"'
    return resp


# ---------------------------------------------------------------------------
# 6. So sánh kỳ (API trả JSON)
# ---------------------------------------------------------------------------

@login_required
def td_so_sanh_api(request):
    """GET: trả delta KPI so với snapshot.
    Params:
      ky   = ngay|tuan|thang|quy|nam  (tương đối)
      date = YYYY-MM-DD               (ngày cụ thể, ưu tiên hơn ky)
    """
    analysis = request.session.get('td_analysis')
    if not analysis:
        return JsonResponse({'error': 'Chưa có dữ liệu'}, status=400)

    today = date.today()
    specific_date = request.GET.get('date', '')
    ky = request.GET.get('ky', 'ngay')

    if specific_date:
        try:
            target_date = date.fromisoformat(specific_date)
        except ValueError:
            return JsonResponse({'error': 'Ngày không hợp lệ (YYYY-MM-DD)'}, status=400)
        ky = 'custom'
    elif ky == 'ngay':
        target_date = today - timedelta(days=1)
    elif ky == 'tuan':
        target_date = today - timedelta(weeks=1)
    elif ky == 'thang':
        target_date = today.replace(day=1) - timedelta(days=1)
    elif ky == 'quy':
        m = ((today.month - 1) // 3) * 3 + 1
        target_date = today.replace(month=m, day=1) - timedelta(days=1)
    elif ky == 'nam':
        target_date = today.replace(month=1, day=1) - timedelta(days=1)
    else:
        target_date = today - timedelta(days=1)

    prev_snap = (
        TDSnapshot.objects
        .filter(file_date__lte=target_date)
        .order_by('-file_date')
        .values('id', 'file_date', 'summary_json', 'file_name',
                'so_lds', 'so_kh', 'tong_du_no', 'qua_han', 'nhom2', 'nhom35', 'kh_dac_biet')
        .first()
    )

    prev_kpi = prev_snap['summary_json'] if prev_snap else None
    delta    = compute_so_sanh(analysis.get('kpi', {}), prev_kpi)

    return JsonResponse({
        'ky':        ky,
        'prev_date': str(prev_snap['file_date']) if prev_snap else None,
        'delta':     delta,
    })


# ---------------------------------------------------------------------------
# 7. FTP Config CRUD
# ---------------------------------------------------------------------------

@login_required
def td_ftp_config_view(request):
    """GET: danh sách + form. POST: tạo mới hoặc sửa."""
    if request.method == 'POST':
        pk = request.POST.get('pk')
        data = {
            'ma_goi':           request.POST.get('ma_goi', '').strip(),
            'ten_goi':          request.POST.get('ten_goi', '').strip(),
            'ftp_pct':          float(request.POST.get('ftp_pct', 5.0)),
            'bien_do_pct':      float(request.POST.get('bien_do_pct', 0.0)),
            'cach_tinh':        request.POST.get('cach_tinh', 'SPRD-FTP+DIEU_CHINH'),
            'nim_co_dinh_pct':  float(request.POST.get('nim_co_dinh_pct', 0.0)),
            'loai_kh':          request.POST.get('loai_kh', 'Tất cả'),
            'is_default':       request.POST.get('is_default') == 'on',
        }
        if not data['ma_goi'] or not data['ten_goi']:
            return JsonResponse({'error': 'Mã gói và tên gói không được để trống'}, status=400)

        # Chỉ một bản ghi là default
        if data['is_default']:
            FTPConfig.objects.exclude(pk=pk or 0).update(is_default=False)

        if pk:
            cfg = get_object_or_404(FTPConfig, pk=pk)
            for k, v in data.items():
                setattr(cfg, k, v)
            cfg.save()
            msg = 'Đã cập nhật cấu hình FTP'
        else:
            FTPConfig.objects.create(**data)
            msg = 'Đã thêm cấu hình FTP mới'

        return JsonResponse({'ok': True, 'message': msg})

    configs = FTPConfig.objects.all().order_by('-is_default', 'ma_goi')
    return render(request, 'templates_app/td_ftp_config.html', {'configs': configs})


@login_required
@require_POST
def td_ftp_delete_view(request, pk):
    """POST: xóa cấu hình FTP."""
    cfg = get_object_or_404(FTPConfig, pk=pk)
    if cfg.is_default:
        return JsonResponse({'error': 'Không thể xóa cấu hình mặc định'}, status=400)
    cfg.delete()
    return JsonResponse({'ok': True})


# ---------------------------------------------------------------------------
# 8. FAQ / Trợ giúp
# ---------------------------------------------------------------------------

@login_required
def td_faq_view(request):
    """Trang trợ giúp & FAQ — nội dung tĩnh."""
    return render(request, 'templates_app/td_faq.html')


# ---------------------------------------------------------------------------
# Debug: xem cột thực tế trong file MSIT80 (chỉ dùng khi cần kiểm tra)
# ---------------------------------------------------------------------------

@login_required
def td_debug_cols_view(request):
    """POST: upload file → trả thông tin debug cột nhóm nợ."""
    if request.method != 'POST':
        return JsonResponse({'error': 'POST only'}, status=405)
    uploaded = request.FILES.get('msit80_file')
    if not uploaded:
        return JsonResponse({'error': 'No file'}, status=400)
    suffix = '.xlsx' if uploaded.name.lower().endswith('.xlsx') else '.xls'
    tmp_path = None
    try:
        import pandas as pd
        import math
        with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as tmp:
            for chunk in uploaded.chunks():
                tmp.write(chunk)
            tmp_path = tmp.name
        try:
            df = pd.read_excel(tmp_path, header=0, engine='xlrd')
        except Exception:
            df = pd.read_excel(tmp_path, header=0, engine='openpyxl')
        cols = [str(c).strip().lower() for c in df.columns]
        df.columns = cols

        def clean(v):
            if isinstance(v, float) and math.isnan(v): return None
            return str(v) if not isinstance(v, (int, float, str, bool, type(None))) else v

        def col_info(col):
            if col not in df.columns:
                return {'exists': False}
            uniq = [clean(v) for v in df[col].dropna().unique()[:20]]
            return {'exists': True, 'unique_values': uniq, 'dtype': str(df[col].dtype)}

        # Tìm các cột liên quan nhóm nợ (chứa 'grp', 'nhom', 'aqcc', 'clas', 'group')
        nhom_cols = [c for c in cols if any(k in c for k in ('grp','nhom','aqcc','clas','group','fin'))]

        return JsonResponse({
            'all_columns': cols,
            'grpno':     col_info('grpno'),
            'aqccdfin':  col_info('aqccdfin'),
            'custseq':   col_info('custseq'),
            'brcd':      col_info('brcd'),
            'nhom_related_cols': nhom_cols,
            'sample_3_rows': [{k: clean(v) for k, v in row.items()}
                              for row in df.head(3).to_dict(orient='records')],
        })
    except Exception as e:
        return JsonResponse({'error': str(e)}, status=500)
    finally:
        if tmp_path and os.path.exists(tmp_path):
            os.unlink(tmp_path)


# ---------------------------------------------------------------------------
# 9. So sánh 2 kỳ (trang riêng)
# ---------------------------------------------------------------------------

@login_required
def td_so_sanh_view(request):
    """GET: render trang so sánh. POST (AJAX multipart): nhận 2 file → phân tích cả 2 → trả JSON."""
    if request.method != 'POST':
        return render(request, 'templates_app/td_so_sanh.html', {})

    file_cur = request.FILES.get('file_cur')
    file_cmp = request.FILES.get('file_cmp')
    if not file_cur or not file_cmp:
        return JsonResponse({'error': 'Vui lòng chọn đủ 2 file MSIT80'}, status=400)

    ftp_configs = list(FTPConfig.objects.values())
    tmp_paths = []

    def _analyse(uploaded):
        suffix = '.xlsx' if uploaded.name.lower().endswith('.xlsx') else '.xls'
        with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as tmp:
            for chunk in uploaded.chunks():
                tmp.write(chunk)
            path = tmp.name
        tmp_paths.append(path)
        df = load_msit80(path)
        analysis = analyze_portfolio(df, ftp_configs=ftp_configs)
        records = analysis.pop('records', [])
        return analysis, records

    try:
        analysis_cur, records_cur = _analyse(file_cur)
        analysis_cmp, records_cmp = _analyse(file_cmp)

        request.session['td_analysis_cur']  = analysis_cur
        request.session['td_records']        = records_cur   # dùng chung key với dashboard filter
        request.session['td_analysis_cmp']  = analysis_cmp
        request.session['td_records_cmp']   = records_cmp

        return JsonResponse({
            'ok':       True,
            'kpi_cur':  analysis_cur.get('kpi', {}),
            'kpi_cmp':  analysis_cmp.get('kpi', {}),
            'file_cur': file_cur.name,
            'file_cmp': file_cmp.name,
        })

    except ValueError as e:
        return JsonResponse({'error': str(e)}, status=400)
    except Exception as e:
        return JsonResponse({'error': f'Lỗi phân tích: {str(e)}'}, status=500)
    finally:
        for p in tmp_paths:
            if os.path.exists(p):
                os.unlink(p)


@login_required
def td_so_sanh_records_api(request):
    """POST AJAX: trả danh sách records đã diff giữa 2 kỳ, có bộ lọc."""
    if request.method != 'POST':
        return JsonResponse({'error': 'Method not allowed'}, status=405)

    records_cur = request.session.get('td_records', [])
    records_cmp = request.session.get('td_records_cmp', [])

    if not records_cur:
        return JsonResponse({'error': 'Chưa có dữ liệu kỳ hiện tại'}, status=400)
    if not records_cmp:
        return JsonResponse({'error': 'Chưa upload file kỳ so sánh'}, status=400)

    try:
        body = json.loads(request.body)
    except Exception:
        body = {}

    f_cbtd      = body.get('cbtd', '')
    f_nhomno    = body.get('nhomno', '')
    f_trangthai = body.get('trangthai', '')
    f_trang_thai_diff = body.get('trang_thai_diff', '')  # Mới/Tất toán/Hiện hữu

    # Index kỳ so sánh theo dsbsseq để tra nhanh
    cmp_idx = {str(r.get('dsbsseq', '')): r for r in records_cmp}
    cur_idx = {str(r.get('dsbsseq', '')): r for r in records_cur}

    rows = []

    # Khoản hiện hữu ở kỳ hiện tại
    for r in records_cur:
        seq = str(r.get('dsbsseq', ''))
        prev = cmp_idx.get(seq)
        if prev:
            trang_thai_diff = 'Hiện hữu'
            delta_bal   = float(r.get('dsbsbal', 0)) - float(prev.get('dsbsbal', 0))
            delta_grp   = int(r.get('grpno', 1)) - int(prev.get('grpno', 1))
            delta_score = int(r.get('_score', 0)) - int(prev.get('_score', 0))
        else:
            trang_thai_diff = 'Mới phát sinh'
            delta_bal = delta_grp = delta_score = None

        rows.append({**r, '_diff': trang_thai_diff,
                     '_delta_bal': delta_bal, '_delta_grp': delta_grp,
                     '_delta_score': delta_score})

    # Khoản đã tất toán (có ở kỳ cmp nhưng không có ở kỳ cur)
    for seq, r in cmp_idx.items():
        if seq not in cur_idx:
            rows.append({**r, '_diff': 'Tất toán',
                         '_delta_bal': None, '_delta_grp': None, '_delta_score': None})

    # Áp bộ lọc
    if f_cbtd:
        rows = [r for r in rows if str(r.get('ofcnm', '')).strip() == f_cbtd.strip()]
    if f_nhomno:
        grp_map = {'1': [1], '2': [2], '3': [3], '4': [4], '5': [5], '35': [3, 4, 5]}
        targets = grp_map.get(f_nhomno, [])
        if targets:
            rows = [r for r in rows if r.get('grpno') in targets]
    if f_trangthai:
        rows = [r for r in rows if r.get('_trangthai') == f_trangthai]
    if f_trang_thai_diff:
        rows = [r for r in rows if r.get('_diff') == f_trang_thai_diff]

    # Thống kê tổng hợp
    moi    = sum(1 for r in rows if r['_diff'] == 'Mới phát sinh')
    tat    = sum(1 for r in rows if r['_diff'] == 'Tất toán')
    hh     = sum(1 for r in rows if r['_diff'] == 'Hiện hữu')
    nhom_tang = sum(1 for r in rows if (r.get('_delta_grp') or 0) > 0)

    return JsonResponse({
        'count':       len(rows),
        'moi':         moi,
        'tat_toan':    tat,
        'hien_huu':    hh,
        'nhom_tang':   nhom_tang,
        'records':     rows[:600],
    })


# ---------------------------------------------------------------------------
# Danh sách snapshots (cho dropdown so sánh)
# ---------------------------------------------------------------------------

@login_required
def td_snapshots_api(request):
    """GET: trả danh sách tất cả snapshot đã lưu."""
    snaps = list(
        TDSnapshot.objects
        .order_by('-file_date')
        .values('id', 'file_date', 'file_name', 'so_lds', 'tong_du_no',
                'qua_han', 'nhom2', 'nhom35', 'uploaded_at')
    )
    for s in snaps:
        s['file_date'] = str(s['file_date'])
        s['uploaded_at'] = str(s['uploaded_at'])[:16]
        s['tong_du_no_ty'] = round(s['tong_du_no'] / 1e9, 1)
    return JsonResponse({'snapshots': snaps})


# ---------------------------------------------------------------------------
# Giữ backward compat: route cũ /tin-dung/phan-tich/ vẫn dùng upload view
# ---------------------------------------------------------------------------
td_analysis_view = td_upload_view
td_export_view   = td_export_word_view
