"""
Helper functions for importing Variables and Templates
"""
import csv
import openpyxl
from io import BytesIO
from django.core.files.uploadedfile import InMemoryUploadedFile
from .models import Variable, Template, Category, TemplateVariable


def import_variables_from_csv(file):
    """
    Import variables from CSV file

    CSV Format:
    name,label,field_type,required,default_value,help_text
    ho_ten,Họ và tên,text,True,,Nhập họ tên đầy đủ
    ngay_sinh,Ngày sinh,date,True,,Nhập ngày sinh

    Returns:
        tuple: (success_count, error_list)
    """
    success_count = 0
    errors = []

    try:
        # Decode if bytes
        if isinstance(file.read(1), bytes):
            file.seek(0)
            content = file.read().decode('utf-8')
            file = BytesIO(content.encode('utf-8'))
        else:
            file.seek(0)

        # Read CSV
        csv_reader = csv.DictReader(file.read().decode('utf-8').splitlines())

        for row_num, row in enumerate(csv_reader, start=2):
            try:
                # Validate required fields
                if not row.get('name') or not row.get('label'):
                    errors.append(f"Dòng {row_num}: Thiếu tên biến hoặc nhãn")
                    continue

                # Validate field_type
                valid_types = ['text', 'textarea', 'date', 'number']
                field_type = row.get('field_type', 'text').lower()
                if field_type not in valid_types:
                    errors.append(f"Dòng {row_num}: Kiểu dữ liệu không hợp lệ '{field_type}'")
                    continue

                # Parse required field
                required_str = row.get('required', 'True').strip()
                required = required_str.lower() in ['true', '1', 'yes', 'có']

                # Create or update variable
                variable, created = Variable.objects.update_or_create(
                    name=row['name'].strip(),
                    defaults={
                        'label': row['label'].strip(),
                        'field_type': field_type,
                        'required': required,
                        'default_value': row.get('default_value', '').strip(),
                        'help_text': row.get('help_text', '').strip(),
                    }
                )
                success_count += 1

            except Exception as e:
                errors.append(f"Dòng {row_num}: {str(e)}")

        return success_count, errors

    except Exception as e:
        return 0, [f"Lỗi đọc file CSV: {str(e)}"]


def import_variables_from_excel(file):
    """
    Import variables from Excel file (.xlsx)

    Excel Format (First row is header):
    | name    | label      | field_type | required | default_value | help_text           |
    |---------|------------|------------|----------|---------------|---------------------|
    | ho_ten  | Họ và tên  | text       | True     |               | Nhập họ tên đầy đủ |
    | ngay_sinh | Ngày sinh | date      | True     |               | Nhập ngày sinh     |

    Returns:
        tuple: (success_count, error_list)
    """
    success_count = 0
    errors = []

    try:
        # Load workbook
        workbook = openpyxl.load_workbook(file, read_only=True)
        sheet = workbook.active

        # Read header row
        headers = []
        for cell in sheet[1]:
            headers.append(cell.value.strip() if cell.value else '')

        # Validate headers
        required_headers = ['name', 'label']
        for header in required_headers:
            if header not in headers:
                return 0, [f"Thiếu cột bắt buộc: {header}"]

        # Process rows
        for row_num, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            try:
                # Create dict from row
                row_data = {}
                for i, value in enumerate(row):
                    if i < len(headers) and headers[i]:
                        row_data[headers[i]] = value if value is not None else ''

                # Skip empty rows
                if not row_data.get('name'):
                    continue

                # Validate required fields
                if not row_data.get('label'):
                    errors.append(f"Dòng {row_num}: Thiếu nhãn")
                    continue

                # Validate field_type
                valid_types = ['text', 'textarea', 'date', 'number']
                field_type = str(row_data.get('field_type', 'text')).lower().strip()
                if field_type not in valid_types:
                    errors.append(f"Dòng {row_num}: Kiểu dữ liệu không hợp lệ '{field_type}'")
                    continue

                # Parse required field
                required_val = row_data.get('required', True)
                if isinstance(required_val, bool):
                    required = required_val
                else:
                    required_str = str(required_val).strip().lower()
                    required = required_str in ['true', '1', 'yes', 'có']

                # Create or update variable
                variable, created = Variable.objects.update_or_create(
                    name=str(row_data['name']).strip(),
                    defaults={
                        'label': str(row_data['label']).strip(),
                        'field_type': field_type,
                        'required': required,
                        'default_value': str(row_data.get('default_value', '')).strip(),
                        'help_text': str(row_data.get('help_text', '')).strip(),
                    }
                )
                success_count += 1

            except Exception as e:
                errors.append(f"Dòng {row_num}: {str(e)}")

        workbook.close()
        return success_count, errors

    except Exception as e:
        return 0, [f"Lỗi đọc file Excel: {str(e)}"]


def import_templates_bulk(files, category_id, variable_ids=None, allowed_group_ids=None):
    """
    Import multiple template files at once

    Args:
        files: List of uploaded .docx files
        category_id: Category ID to assign templates to
        variable_ids: List of variable IDs to assign to all templates (optional)
        allowed_group_ids: List of group IDs for permissions (optional)

    Returns:
        tuple: (success_count, error_list)
    """
    success_count = 0
    errors = []

    try:
        category = Category.objects.get(id=category_id)
    except Category.DoesNotExist:
        return 0, ["Danh mục không tồn tại"]

    for file in files:
        try:
            # Validate file extension
            if not file.name.endswith('.docx'):
                errors.append(f"{file.name}: Không phải file .docx")
                continue

            # Generate template name from filename
            template_name = file.name.replace('.docx', '').replace('_', ' ').strip()

            # Create template
            template = Template.objects.create(
                category=category,
                name=template_name,
                file=file,
                is_active=True,
                order=0
            )

            # Assign variables if provided
            if variable_ids:
                for order, var_id in enumerate(variable_ids, start=1):
                    try:
                        variable = Variable.objects.get(id=var_id)
                        TemplateVariable.objects.create(
                            template=template,
                            variable=variable,
                            order=order
                        )
                    except Variable.DoesNotExist:
                        errors.append(f"{file.name}: Biến ID {var_id} không tồn tại")

            # Assign groups if provided
            if allowed_group_ids:
                template.allowed_groups.set(allowed_group_ids)

            success_count += 1

        except Exception as e:
            errors.append(f"{file.name}: {str(e)}")

    return success_count, errors


def export_variables_to_csv():
    """
    Export all variables to CSV format

    Returns:
        str: CSV content
    """
    import io

    output = io.StringIO()
    writer = csv.writer(output)

    # Write header
    writer.writerow(['name', 'label', 'field_type', 'required', 'default_value', 'help_text'])

    # Write data
    for var in Variable.objects.all().order_by('name'):
        writer.writerow([
            var.name,
            var.label,
            var.field_type,
            var.required,
            var.default_value,
            var.help_text
        ])

    return output.getvalue()


def export_variables_to_excel():
    """
    Export all variables to Excel format

    Returns:
        BytesIO: Excel file content
    """
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Variables"

    # Write header
    headers = ['name', 'label', 'field_type', 'required', 'default_value', 'help_text']
    for col, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=col)
        cell.value = header
        cell.font = openpyxl.styles.Font(bold=True)

    # Write data
    for row, var in enumerate(Variable.objects.all().order_by('name'), start=2):
        sheet.cell(row=row, column=1).value = var.name
        sheet.cell(row=row, column=2).value = var.label
        sheet.cell(row=row, column=3).value = var.field_type
        sheet.cell(row=row, column=4).value = var.required
        sheet.cell(row=row, column=5).value = var.default_value
        sheet.cell(row=row, column=6).value = var.help_text

    # Auto-size columns
    for column in sheet.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(cell.value)
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        sheet.column_dimensions[column_letter].width = adjusted_width

    # Save to BytesIO
    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    return output
