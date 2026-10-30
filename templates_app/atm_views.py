"""
Views quản lý ATM: tiếp quỹ, thừa thiếu quỹ, lịch trực xe, công tác phí.
"""
# Standard library imports
import calendar
from datetime import date, datetime, timedelta

# Django imports
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.db import models
from django.db.models import Q
from django.http import HttpResponse, Http404
from django.shortcuts import render, get_object_or_404, redirect
from django.urls import reverse

# Local application imports
from .forms import ATMReplenishmentForm, ATMDiscrepancyForm
from .models import Template, BranchConfig, remove_vietnamese_diacritics, ATM, ATMManagementBoard, ATMReplenishment, ATMDiscrepancy
from .utils import render_word_template
from .atm_travel_claim import ATM_TRIP_CONFIG, MAX_TRIPS_PER_SHEET, build_payment_statement_rows, build_travel_log_groups, get_position_holder


# ============================================================================
# ATM Management Views
# ============================================================================

@login_required
def atm_dashboard(request):
    """Dashboard cho quản lý ATM (chỉ admin)"""
    if not request.user.is_superuser:
        messages.error(request, 'Bạn không có quyền truy cập trang này')
        return redirect('dashboard')

    # Thống kê
    total_atms = ATM.objects.filter(is_active=True).count()
    total_replenishments = ATMReplenishment.objects.count()
    total_discrepancies = ATMDiscrepancy.objects.count()
    pending_discrepancies = ATMDiscrepancy.objects.filter(status='pending').count()

    recent_replenishments = ATMReplenishment.objects.select_related(
        'atm', 'vehicle', 'driver', 'guard', 'created_by'
    ).order_by('-replenishment_date', '-created_at')[:10]

    recent_discrepancies = ATMDiscrepancy.objects.select_related(
        'atm', 'created_by'
    ).order_by('-audit_cycle_end', '-created_at')[:5]

    # Lấy danh sách templates (chỉ templates có tên chứa "ATM")
    if request.user.is_superuser:
        templates = Template.objects.filter(is_active=True, name__icontains='ATM').select_related('category').order_by('category__order', 'order', 'name')
    else:
        # Lấy các groups của user
        user_groups = request.user.groups.all()
        templates = Template.objects.filter(
            is_active=True,
            name__icontains='ATM'
        ).filter(
            models.Q(allowed_groups__isnull=True) | models.Q(allowed_groups__in=user_groups)
        ).distinct().select_related('category').order_by('category__order', 'order', 'name')

    context = {
        'total_atms': total_atms,
        'total_replenishments': total_replenishments,
        'total_discrepancies': total_discrepancies,
        'pending_discrepancies': pending_discrepancies,
        'recent_replenishments': recent_replenishments,
        'recent_discrepancies': recent_discrepancies,
        'templates': templates,
    }
    return render(request, 'templates_app/atm/dashboard.html', context)


@login_required
def atm_guide_download(request):
    """Tạo và tải file Word hướng dẫn sử dụng chức năng quản lý ATM."""
    if not request.user.is_superuser:
        messages.error(request, 'Bạn không có quyền truy cập trang này')
        return redirect('dashboard')

    from docx import Document as DocxDocument
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = DocxDocument()

    def add_title(text):
        p = doc.add_heading(text, level=0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def add_h(text, level=1):
        doc.add_heading(text, level=level)

    def add_p(text, bold=False):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold = bold
        return p

    def add_bullets(items):
        for item in items:
            doc.add_paragraph(item, style='List Bullet')

    add_title('HƯỚNG DẪN SỬ DỤNG CHỨC NĂNG QUẢN LÝ ATM')
    add_p('Tài liệu này hướng dẫn các chức năng chính trong mục Quản lý ATM: Tiếp quỹ, '
          'Thừa/thiếu quỹ, Bảng kê thanh toán công tác phí và Ban quản lý ATM.')

    add_h('1. Tiếp quỹ ATM')
    add_p('Dùng để lập phiếu tiếp quỹ cho từng máy ATM theo ngày, kèm số lượng tờ tiền '
          'theo từng mệnh giá và thông tin lệnh điều chuyển (phương tiện, tài xế, bảo vệ).')
    add_bullets([
        'Vào Quản lý ATM → "Tạo phiếu tiếp quỹ", chọn máy ATM, ngày tiếp quỹ, nhập số tờ '
        'từng mệnh giá (50k/100k/200k/500k) — hệ thống tự tính tổng số tiền.',
        'Mỗi phiếu được tự động gán "Số văn bản giấy đi đường" tăng dần theo từng năm '
        '(sang năm mới tự reset về 1) — dùng để in lên giấy đi đường, không cần nhập tay.',
        'Sau khi lưu, chọn mẫu biểu (Tiếp quỹ ATM 04/06, Tiếp quỹ ATM 05...) để tạo file Word.',
    ])

    add_h('2. Thừa/thiếu quỹ ATM')
    add_p('Dùng để ghi nhận các giao dịch thừa hoặc thiếu quỹ phát hiện khi đối chiếu, và '
          'in giấy đề nghị hạch toán gửi bộ phận kế toán.')
    add_bullets([
        'Thêm từng giao dịch: Họ tên, Số tài khoản là bắt buộc; Số thẻ, ID giao dịch, '
        'Số trace có thể để trống nếu không có — khi in, các nhãn tương ứng (SO THE/ID/'
        'TRACE) sẽ tự động ẩn nếu để trống, không in nhãn kèm chỗ trống.',
        'Có thể nhập nhiều giao dịch cùng lúc (nhập hàng loạt) cho cùng 1 chu kỳ kiểm quỹ.',
        'Khi in theo cả chu kỳ (nhiều giao dịch của cùng 1 máy ATM, cùng khoảng ngày): '
        'Mẫu 03/THE/QTGD sẽ gộp tất cả giao dịch vào 1 trang, mỗi giao dịch 1 dòng, số tiền '
        'của từng người hiển thị đúng riêng từng giao dịch; riêng ô "Số tiền bằng số/chữ" '
        'ở đầu trang hiển thị TỔNG của cả chu kỳ. Mẫu 04/THE/QTGD in mỗi giao dịch 1 trang riêng.',
    ])

    add_h('3. Bảng kê thanh toán công tác phí (Giấy đi đường)')
    add_p('Xem trước và xuất file Word Bảng kê thanh toán + Giấy đi đường cho Ban quản lý '
          'ATM theo khoảng ngày bất kỳ.')
    add_bullets([
        'Hệ thống tự tính đúng ai là người trực/phụ trách để thanh toán cho TỪNG NGÀY cụ thể, '
        'kể cả khi trong cùng 1 khoảng ngày có nhiều người luân phiên (vd: người A trực vài '
        'ngày, chuyển sang người B, rồi A quay lại) — mỗi giai đoạn được tính và liệt kê riêng, '
        'không gộp nhầm hay bỏ sót.',
        'Điều kiện để tính đúng: mỗi lần đổi người phụ trách PHẢI tạo một bản ghi MỚI trong '
        '"Ban quản lý ATM" (xem mục 4) — không sửa lại bản ghi cũ của người đã từng giữ vị trí.',
    ])

    add_h('4. Ban quản lý ATM')
    add_p('Quản lý danh sách người giữ các vị trí trong Ban quản lý ATM (Trưởng ban, Trưởng '
          'phòng KTNQ, Cán bộ phụ trách ATM) theo từng giai đoạn, kèm số/ngày quyết định thành lập.')
    add_bullets([
        'Quản lý qua trang Admin (/admin/) → "Ban quản lý ATM". Mỗi lần một người được phân '
        'công vào 1 vị trí (kể cả người đã từng giữ vị trí đó quay lại), bấm "Thêm" để tạo '
        'bản ghi mới — hệ thống tự động kết thúc hiệu lực của người đang giữ vị trí trước đó.',
        'Một bản ghi đã kết thúc hiệu lực (đã có người khác tiếp nhận sau đó) sẽ tự động bị '
        'khóa, không cho sửa — để không làm sai lệch số quyết định và bảng kê thanh toán của '
        'giai đoạn đã qua. Muốn thay đổi, luôn tạo bản ghi mới.',
        'Số quyết định/Ngày quyết định của Trưởng ban được dùng làm biến chung cho các mẫu '
        'in liên quan (board_decision_number, board_decision_date, disc_board_decision_number, '
        'disc_board_decision_date) — xem chi tiết tại trang "Thư viện Biến".',
    ])

    add_h('5. Mẫu in và Thư viện Biến')
    add_p('Toàn bộ biến có thể chèn vào các mẫu Word (file .docx) cho ATM được liệt kê tại '
          'trang "Thư viện Biến" (menu chính), thẻ "Biến ATM" và "Biến Thừa/Thiếu quỹ ATM" — '
          'tìm theo tên biến hoặc mô tả, bấm nút copy để sao chép tên biến dạng {{tên_biến}}.')

    for section in doc.sections:
        footer = section.footer
        footer.paragraphs[0].text = 'Hướng dẫn sử dụng — Quản lý ATM'
        footer.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    from io import BytesIO
    output_stream = BytesIO()
    doc.save(output_stream)
    output_stream.seek(0)

    response = HttpResponse(
        output_stream.read(),
        content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    )
    response['Content-Disposition'] = 'attachment; filename="Huong_Dan_Su_Dung_ATM.docx"'
    return response


@login_required
def atm_replenishment_create(request):
    """Tạo phiếu tiếp quỹ ATM mới"""
    if not request.user.is_superuser:
        messages.error(request, 'Bạn không có quyền truy cập trang này')
        return redirect('dashboard')

    if request.method == 'POST':
        form = ATMReplenishmentForm(request.POST)
        if form.is_valid():
            replenishment = form.save(commit=False)
            replenishment.created_by = request.user
            replenishment.save()
            messages.success(request, 'Đã tạo phiếu tiếp quỹ thành công')
            return redirect('atm_dashboard')
    else:
        form = ATMReplenishmentForm()

    # Lấy thông tin ban quản lý ATM
    management_board = {}
    for position in ['team_leader', 'treasury_head', 'atm_officer']:
        member = ATMManagementBoard.objects.filter(
            position=position, is_active=True
        ).first()
        management_board[position] = member

    context = {
        'form': form,
        'management_board': management_board,
    }
    return render(request, 'templates_app/atm/replenishment_form.html', context)


@login_required
def atm_replenishment_edit(request, replenishment_id):
    """Sửa phiếu tiếp quỹ ATM"""
    if not request.user.is_superuser:
        messages.error(request, 'Bạn không có quyền truy cập trang này')
        return redirect('dashboard')

    replenishment = get_object_or_404(ATMReplenishment, pk=replenishment_id)

    if request.method == 'POST':
        form = ATMReplenishmentForm(request.POST, instance=replenishment)
        if form.is_valid():
            form.save()
            messages.success(request, 'Đã cập nhật phiếu tiếp quỹ thành công')
            return redirect('atm_replenishment_list')
    else:
        form = ATMReplenishmentForm(instance=replenishment)

    management_board = {}
    for position in ['team_leader', 'treasury_head', 'atm_officer']:
        member = ATMManagementBoard.objects.filter(
            position=position, is_active=True
        ).first()
        management_board[position] = member

    context = {
        'form': form,
        'replenishment': replenishment,
        'management_board': management_board,
        'is_edit': True,
    }
    return render(request, 'templates_app/atm/replenishment_form.html', context)


@login_required
def atm_replenishment_delete(request, replenishment_id):
    """Xóa phiếu tiếp quỹ ATM"""
    if not request.user.is_superuser:
        messages.error(request, 'Bạn không có quyền truy cập trang này')
        return redirect('dashboard')

    replenishment = get_object_or_404(ATMReplenishment, pk=replenishment_id)

    if request.method == 'POST':
        replenishment.delete()
        messages.success(request, 'Đã xóa phiếu tiếp quỹ thành công')
        return redirect('atm_replenishment_list')

    context = {
        'replenishment': replenishment,
    }
    return render(request, 'templates_app/atm/replenishment_confirm_delete.html', context)


@login_required
def atm_load_replenishment_data(request, replenishment_id, template_id):
    """Tạo và tải file Word trực tiếp từ phiếu tiếp quỹ"""
    if not request.user.is_superuser:
        messages.error(request, 'Bạn không có quyền truy cập trang này')
        return redirect('dashboard')

    # Lấy replenishment
    replenishment = get_object_or_404(
        ATMReplenishment.objects.select_related(
            'atm', 'vehicle', 'driver', 'guard', 'created_by'
        ),
        pk=replenishment_id
    )

    # Lấy template và kiểm tra quyền
    template = get_object_or_404(Template, id=template_id, is_active=True)
    if not template.user_has_access(request.user):
        raise Http404("Bạn không có quyền truy cập mẫu biểu này")

    try:
        # Lấy dữ liệu từ replenishment
        data = replenishment.get_data_dict()

        # Thêm biến chung (chi nhánh + custom variables)
        branch_config = BranchConfig.get_for_user(request.user)
        data.update(branch_config.get_all_variables())

        # Thêm text_replenishment_date: "Địa danh, ngày dd tháng mm năm yyyy"
        if replenishment.replenishment_date:
            d = replenishment.replenishment_date
            dia_danh = data.get('dia_danh', '')
            if dia_danh:
                data['text_replenishment_date'] = f"{dia_danh}, ngày {d.day:02d} tháng {d.month:02d} năm {d.year}"
            else:
                data['text_replenishment_date'] = f"ngày {d.day:02d} tháng {d.month:02d} năm {d.year}"

        # Render template Word với dữ liệu
        template_path = template.file.path
        output_stream = render_word_template(template_path, data)

        # Trả về file Word để download
        response = HttpResponse(
            output_stream.read(),
            content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )
        filename = f"{template.name}_ATM_{replenishment.atm.machine_id}_{replenishment.replenishment_date.strftime('%Y%m%d')}.docx"
        response['Content-Disposition'] = f'attachment; filename="{filename}"'

        return response

    except Exception as e:
        messages.error(request, f"Lỗi khi tạo file Word: {str(e)}")
        return redirect('atm_dashboard')


@login_required
def atm_replenishment_list(request):
    """Danh sách phiếu tiếp quỹ ATM"""
    if not request.user.is_superuser:
        messages.error(request, 'Bạn không có quyền truy cập trang này')
        return redirect('dashboard')

    replenishments = ATMReplenishment.objects.select_related(
        'atm', 'vehicle', 'driver', 'guard', 'created_by'
    ).order_by('-replenishment_date', '-created_at')

    # Pagination
    from django.core.paginator import Paginator
    paginator = Paginator(replenishments, 20)  # 20 items per page
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)

    # Lấy danh sách templates (chỉ templates có tên chứa "ATM")
    if request.user.is_superuser:
        templates = Template.objects.filter(is_active=True, name__icontains='ATM').select_related('category').order_by('category__order', 'order', 'name')
    else:
        user_groups = request.user.groups.all()
        templates = Template.objects.filter(
            is_active=True,
            name__icontains='ATM'
        ).filter(
            Q(allowed_groups__isnull=True) | Q(allowed_groups__in=user_groups)
        ).distinct().select_related('category').order_by('category__order', 'order', 'name')

    context = {
        'page_obj': page_obj,
        'templates': templates,
    }
    return render(request, 'templates_app/atm/replenishment_list.html', context)


# ============================================================================
# ATM Discrepancy Views
# ============================================================================

@login_required
def atm_discrepancy_list(request):
    """Danh sách giao dịch thừa/thiếu quỹ ATM (gom nhóm theo ATM + chu kỳ)"""
    if not request.user.is_superuser:
        messages.error(request, 'Bạn không có quyền truy cập trang này')
        return redirect('dashboard')

    from django.db.models import Count, Min
    # Lấy các nhóm unique (atm, audit_cycle_start, audit_cycle_end), sắp xếp mới nhất trước
    cycles = (
        ATMDiscrepancy.objects
        .values('atm_id', 'audit_cycle_start', 'audit_cycle_end')
        .annotate(count=Count('id'), first_created=Min('created_at'))
        .order_by('-audit_cycle_end', '-first_created')
    )

    group_rows = []
    for cycle in cycles:
        items = list(
            ATMDiscrepancy.objects.filter(
                atm_id=cycle['atm_id'],
                audit_cycle_start=cycle['audit_cycle_start'],
                audit_cycle_end=cycle['audit_cycle_end'],
            ).select_related('atm', 'created_by').order_by('id')
        )
        if not items:
            continue
        rep = items[0]
        total_surplus = sum(int(d.amount) for d in items if d.discrepancy_type == 'surplus')
        total_deficit = sum(int(d.amount) for d in items if d.discrepancy_type == 'deficit')
        # Key dùng cho URL Word: atm_id + start_yyyymmdd + end_yyyymmdd
        cycle_key = f"{cycle['atm_id']}__{cycle['audit_cycle_start'].strftime('%Y%m%d')}__{cycle['audit_cycle_end'].strftime('%Y%m%d')}"
        group_rows.append({
            'cycle_key': cycle_key,
            'atm_id': cycle['atm_id'],
            'start': cycle['audit_cycle_start'],
            'end': cycle['audit_cycle_end'],
            'representative': rep,
            'items': items,
            'count': len(items),
            'total_surplus': total_surplus,
            'total_deficit': total_deficit,
        })

    # Pagination
    from django.core.paginator import Paginator
    paginator = Paginator(group_rows, 20)
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)

    # Chỉ lấy templates liên quan đến thừa/thiếu quỹ ATM (không lấy tiếp quỹ/pin)
    atm_templates = Template.objects.filter(
        is_active=True,
        category_id=6,
    ).filter(
        models.Q(name__icontains='hach toan') |
        models.Q(name__icontains='Hach toan') |
        models.Q(name__icontains='THE') |
        models.Q(name__icontains='hoan tra')
    ).order_by('order', 'name')

    context = {
        'page_obj': page_obj,
        'atm_templates': atm_templates,
    }
    return render(request, 'templates_app/atm/discrepancy_list.html', context)


@login_required
def atm_discrepancy_create(request):
    """Tạo giao dịch thừa/thiếu quỹ ATM mới (hỗ trợ nhiều giao dịch trong một chu kỳ)"""
    if not request.user.is_superuser:
        messages.error(request, 'Bạn không có quyền truy cập trang này')
        return redirect('dashboard')

    if request.method == 'POST':
        # Lấy thông tin chung của chu kỳ
        atm_id = request.POST.get('atm')
        audit_cycle_start = request.POST.get('audit_cycle_start')
        audit_cycle_end = request.POST.get('audit_cycle_end')
        status = request.POST.get('status', 'pending')
        notes = request.POST.get('notes', '')

        # Lấy danh sách giao dịch (dạng mảng)
        full_names = request.POST.getlist('full_name[]')
        account_numbers = request.POST.getlist('account_number[]')
        card_numbers = request.POST.getlist('card_number[]')
        trace_numbers = request.POST.getlist('trace_number[]')
        transaction_ids = request.POST.getlist('transaction_id[]')
        discrepancy_types = request.POST.getlist('discrepancy_type[]')
        amounts = request.POST.getlist('amount[]')

        errors = []

        # Validate dữ liệu chung
        if not atm_id:
            errors.append('Vui lòng chọn máy ATM.')
        if not audit_cycle_start:
            errors.append('Vui lòng nhập ngày bắt đầu chu kỳ.')
        if not audit_cycle_end:
            errors.append('Vui lòng nhập ngày kết thúc chu kỳ.')
        if audit_cycle_start and audit_cycle_end and audit_cycle_end < audit_cycle_start:
            errors.append('Ngày kết thúc chu kỳ phải sau ngày bắt đầu.')
        if not full_names or not any(n.strip() for n in full_names):
            errors.append('Vui lòng nhập ít nhất một giao dịch.')

        if not errors:
            try:
                import uuid as uuid_module
                from django.db import transaction as db_transaction
                from .models import ATM as ATMModel
                atm_obj = ATMModel.objects.get(pk=atm_id)
                batch_group_id = uuid_module.uuid4()
                saved_count = 0
                with db_transaction.atomic():
                    for i in range(len(full_names)):
                        full_name = full_names[i].strip() if i < len(full_names) else ''
                        if not full_name:
                            continue
                        ATMDiscrepancy.objects.create(
                            atm=atm_obj,
                            full_name=full_name,
                            account_number=account_numbers[i].strip() if i < len(account_numbers) else '',
                            card_number=card_numbers[i].strip() if i < len(card_numbers) else '',
                            trace_number=trace_numbers[i].strip() if i < len(trace_numbers) else '',
                            transaction_id=transaction_ids[i].strip() if i < len(transaction_ids) else '',
                            discrepancy_type=discrepancy_types[i] if i < len(discrepancy_types) else 'deficit',
                            amount=amounts[i] if i < len(amounts) else 0,
                            audit_cycle_start=audit_cycle_start,
                            audit_cycle_end=audit_cycle_end,
                            status=status,
                            notes=notes,
                            created_by=request.user,
                            group_id=batch_group_id,
                        )
                        saved_count += 1
                messages.success(request, f'Đã tạo {saved_count} giao dịch thừa/thiếu quỹ thành công')
                return redirect('atm_discrepancy_list')
            except Exception as e:
                errors.append(f'Lỗi khi lưu dữ liệu: {str(e)}')

        import json
        post_rows = []
        for i in range(len(full_names)):
            post_rows.append({
                'full_name': full_names[i] if i < len(full_names) else '',
                'account_number': account_numbers[i] if i < len(account_numbers) else '',
                'card_number': card_numbers[i] if i < len(card_numbers) else '',
                'trace_number': trace_numbers[i] if i < len(trace_numbers) else '',
                'transaction_id': transaction_ids[i] if i < len(transaction_ids) else '',
                'discrepancy_type': discrepancy_types[i] if i < len(discrepancy_types) else 'deficit',
                'amount': amounts[i] if i < len(amounts) else '',
            })
        from .models import ATM as ATMModel
        context = {
            'atm_list': ATMModel.objects.filter(is_active=True),
            'status_choices': ATMDiscrepancy.STATUS_CHOICES,
            'discrepancy_types': ATMDiscrepancy.DISCREPANCY_TYPES,
            'errors': errors,
            'post_data': request.POST,
            'post_rows_json': json.dumps(post_rows),
        }
        return render(request, 'templates_app/atm/discrepancy_create_multi.html', context)
    else:
        from .models import ATM as ATMModel
        context = {
            'atm_list': ATMModel.objects.filter(is_active=True),
            'status_choices': ATMDiscrepancy.STATUS_CHOICES,
            'discrepancy_types': ATMDiscrepancy.DISCREPANCY_TYPES,
        }
    return render(request, 'templates_app/atm/discrepancy_create_multi.html', context)


@login_required
def atm_discrepancy_edit(request, discrepancy_id):
    """Sửa giao dịch thừa/thiếu quỹ ATM"""
    if not request.user.is_superuser:
        messages.error(request, 'Bạn không có quyền truy cập trang này')
        return redirect('dashboard')

    discrepancy = get_object_or_404(ATMDiscrepancy, pk=discrepancy_id)

    if request.method == 'POST':
        form = ATMDiscrepancyForm(request.POST, instance=discrepancy)
        if form.is_valid():
            form.save()
            messages.success(request, 'Đã cập nhật giao dịch thành công')
            return redirect('atm_discrepancy_list')
    else:
        form = ATMDiscrepancyForm(instance=discrepancy)

    context = {
        'form': form,
        'discrepancy': discrepancy,
        'is_edit': True,
    }
    return render(request, 'templates_app/atm/discrepancy_form.html', context)


@login_required
def atm_discrepancy_delete(request, discrepancy_id):
    """Xóa giao dịch thừa/thiếu quỹ ATM"""
    if not request.user.is_superuser:
        messages.error(request, 'Bạn không có quyền truy cập trang này')
        return redirect('dashboard')

    discrepancy = get_object_or_404(ATMDiscrepancy, pk=discrepancy_id)

    if request.method == 'POST':
        discrepancy.delete()
        messages.success(request, 'Đã xóa giao dịch thành công')
        return redirect('atm_discrepancy_list')

    context = {
        'discrepancy': discrepancy,
    }
    return render(request, 'templates_app/atm/discrepancy_confirm_delete.html', context)


@login_required
def atm_discrepancy_group_word(request, atm_id, start_date, end_date):
    """Tự động tạo file Word cho tất cả giao dịch trong cùng chu kỳ (ATM + start + end)"""
    if not request.user.is_superuser:
        messages.error(request, 'Bạn không có quyền truy cập trang này')
        return redirect('dashboard')

    from datetime import datetime as dt
    try:
        start_dt = dt.strptime(start_date, '%Y%m%d').date()
        end_dt = dt.strptime(end_date, '%Y%m%d').date()
    except ValueError:
        messages.error(request, 'Ngày không hợp lệ')
        return redirect('atm_discrepancy_list')

    # atm_id ở đây là machine_id (string PK của ATM)
    discrepancies = ATMDiscrepancy.objects.filter(
        atm__machine_id=atm_id,
        audit_cycle_start=start_dt,
        audit_cycle_end=end_dt,
    ).select_related('atm', 'created_by').order_by('id')

    if not discrepancies.exists():
        messages.error(request, 'Không tìm thấy nhóm giao dịch này')
        return redirect('atm_discrepancy_list')

    try:
        from docx import Document
        from docx.shared import Pt, Cm
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from io import BytesIO

        first = discrepancies.first()
        atm = first.atm

        # Lấy thông tin ban quản lý ATM
        team_leader = ATMManagementBoard.objects.filter(position='team_leader', is_active=True).first()
        treasury_head = ATMManagementBoard.objects.filter(position='treasury_head', is_active=True).first()
        atm_officer = ATMManagementBoard.objects.filter(position='atm_officer', is_active=True).first()

        # Lấy thông tin chi nhánh
        branch_config = BranchConfig.get_for_user(request.user)
        branch_vars = branch_config.get_all_variables()
        dia_danh = branch_vars.get('dia_danh', '')
        ten_don_vi = branch_vars.get('ten_don_vi', '')

        doc = Document()

        # Thiết lập margin
        for section in doc.sections:
            section.top_margin = Cm(2)
            section.bottom_margin = Cm(2)
            section.left_margin = Cm(2.5)
            section.right_margin = Cm(2)

        def set_font(run, size=11, bold=False):
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.name = 'Times New Roman'

        def add_paragraph(text='', bold=False, size=11, align=WD_ALIGN_PARAGRAPH.LEFT):
            p = doc.add_paragraph()
            p.alignment = align
            run = p.add_run(text)
            set_font(run, size, bold)
            return p

        # Tiêu đề
        if ten_don_vi:
            p = add_paragraph(ten_don_vi.upper(), bold=True, size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
        add_paragraph('', size=10)

        disc_type_display = 'THỪA/THIẾU QUỸ' if discrepancies.filter(discrepancy_type='surplus').exists() and discrepancies.filter(discrepancy_type='deficit').exists() \
            else ('THỪA QUỸ' if discrepancies.first().discrepancy_type == 'surplus' else 'THIẾU QUỸ')
        p = add_paragraph(f'BÁO CÁO GIAO DỊCH {disc_type_display} ATM', bold=True, size=14, align=WD_ALIGN_PARAGRAPH.CENTER)

        # Thông tin chung
        add_paragraph('')
        p = doc.add_paragraph()
        p.add_run('Máy ATM: ').bold = True
        run = p.add_run(f'{atm.machine_id} - {atm.address}')
        set_font(run)

        p = doc.add_paragraph()
        r = p.add_run('Chu kỳ kiểm quỹ: ')
        r.bold = True
        set_font(r, bold=True)
        run = p.add_run(f'{first.audit_cycle_start.strftime("%d/%m/%Y")} đến {first.audit_cycle_end.strftime("%d/%m/%Y")}')
        set_font(run)

        p = doc.add_paragraph()
        r = p.add_run('Trạng thái: ')
        r.bold = True
        set_font(r, bold=True)
        run = p.add_run(first.get_status_display())
        set_font(run)

        if first.notes:
            p = doc.add_paragraph()
            r = p.add_run('Ghi chú: ')
            r.bold = True
            set_font(r, bold=True)
            run = p.add_run(first.notes)
            set_font(run)

        add_paragraph('')

        # Bảng danh sách giao dịch
        headers = ['STT', 'Họ tên', 'Số TK', 'Số thẻ', 'Số trace', 'ID giao dịch', 'Loại', 'Số tiền (đ)']
        table = doc.add_table(rows=1, cols=len(headers))
        table.style = 'Table Grid'

        # Header row
        hdr_cells = table.rows[0].cells
        for i, h in enumerate(headers):
            hdr_cells[i].text = h
            hdr_cells[i].paragraphs[0].runs[0].bold = True
            hdr_cells[i].paragraphs[0].runs[0].font.size = Pt(10)
            hdr_cells[i].paragraphs[0].runs[0].font.name = 'Times New Roman'
            hdr_cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

        total_surplus = 0
        total_deficit = 0
        for idx, disc in enumerate(discrepancies, 1):
            row_cells = table.add_row().cells
            type_str = 'Thừa' if disc.discrepancy_type == 'surplus' else 'Thiếu'
            amt = int(disc.amount)
            if disc.discrepancy_type == 'surplus':
                total_surplus += amt
            else:
                total_deficit += amt
            values = [
                str(idx),
                disc.full_name,
                disc.account_number,
                disc.card_number,
                disc.trace_number,
                disc.transaction_id,
                type_str,
                f"{amt:,}",
            ]
            for i, val in enumerate(values):
                row_cells[i].text = val
                row_cells[i].paragraphs[0].runs[0].font.size = Pt(10)
                row_cells[i].paragraphs[0].runs[0].font.name = 'Times New Roman'
                if i in (0, 6, 7):
                    row_cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Dòng tổng cộng
        total_row = table.add_row().cells
        total_row[0].merge(total_row[5])
        total_row[0].text = 'Tổng cộng'
        total_row[0].paragraphs[0].runs[0].bold = True
        total_row[0].paragraphs[0].runs[0].font.size = Pt(10)
        total_row[0].paragraphs[0].runs[0].font.name = 'Times New Roman'
        total_row[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        total_row[6].text = f"Thừa: {total_surplus:,}\nThiếu: {total_deficit:,}"
        total_row[6].paragraphs[0].runs[0].font.size = Pt(10)
        total_row[6].paragraphs[0].runs[0].font.name = 'Times New Roman'

        add_paragraph('')

        # Chữ ký
        sig_table = doc.add_table(rows=2, cols=3)
        sig_positions = [
            ('TRƯỞNG BAN', team_leader),
            ('TRƯỞNG PHÒNG KTNQ', treasury_head),
            ('CÁN BỘ PHỤ TRÁCH ATM', atm_officer),
        ]
        for col_idx, (title, member) in enumerate(sig_positions):
            cell = sig_table.rows[0].cells[col_idx]
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(title)
            r.bold = True
            r.font.size = Pt(11)
            r.font.name = 'Times New Roman'

            cell2 = sig_table.rows[1].cells[col_idx]
            p2 = cell2.paragraphs[0]
            p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
            name_text = member.full_name if member else '...........................'
            run2 = p2.add_run(name_text)
            run2.font.size = Pt(11)
            run2.font.name = 'Times New Roman'

        # Ngày ký
        add_paragraph('')
        if dia_danh:
            date_str = f"{dia_danh}, ngày {first.audit_cycle_end.day:02d} tháng {first.audit_cycle_end.month:02d} năm {first.audit_cycle_end.year}"
        else:
            date_str = f"Ngày {first.audit_cycle_end.day:02d} tháng {first.audit_cycle_end.month:02d} năm {first.audit_cycle_end.year}"
        add_paragraph(date_str, align=WD_ALIGN_PARAGRAPH.RIGHT)

        output = BytesIO()
        doc.save(output)
        output.seek(0)

        response = HttpResponse(
            output.read(),
            content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )
        disc_type_fn = 'ThuaThieu' if disc_type_display == 'THỪA/THIẾU QUỸ' else ('Thua' if 'THỪA' in disc_type_display else 'Thieu')
        filename = f"BaoCao_{disc_type_fn}_ATM_{atm.machine_id}_{first.audit_cycle_end.strftime('%Y%m%d')}.docx"
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response

    except Exception as e:
        messages.error(request, f"Lỗi khi tạo file Word: {str(e)}")
        return redirect('atm_discrepancy_list')


def _parse_travel_claim_dates(request):
    """Lấy start_date/end_date từ GET, mặc định là tháng hiện tại"""
    today = date.today()
    default_start = today.replace(day=1)
    default_end = today.replace(day=calendar.monthrange(today.year, today.month)[1])

    start_str = request.GET.get('start_date', '')
    end_str = request.GET.get('end_date', '')
    try:
        start_date = datetime.strptime(start_str, '%Y-%m-%d').date() if start_str else default_start
    except ValueError:
        start_date = default_start
    try:
        end_date = datetime.strptime(end_str, '%Y-%m-%d').date() if end_str else default_end
    except ValueError:
        end_date = default_end

    return start_date, end_date


@login_required
def vehicle_duty_schedule(request):
    """Trang chọn/lưu lịch trực xe theo tháng (mỗi ngày 1 tài xế trực)"""
    from .models import Person, VehicleDutySchedule

    if not request.user.is_superuser:
        messages.error(request, 'Bạn không có quyền truy cập trang này')
        return redirect('dashboard')

    today = date.today()
    try:
        year = int(request.GET.get('year', today.year))
    except (ValueError, TypeError):
        year = today.year
    try:
        month = int(request.GET.get('month', today.month))
    except (ValueError, TypeError):
        month = today.month
    try:
        date(year, month, 1)  # validate kết hợp year/month hợp lệ
    except ValueError:
        year, month = today.year, today.month

    drivers = list(Person.objects.filter(person_type='driver', is_active=True).order_by('id'))

    if request.method == 'POST':
        try:
            year = int(request.POST.get('year', year))
            month = int(request.POST.get('month', month))
        except (ValueError, TypeError):
            pass

        num_days = calendar.monthrange(year, month)[1]
        driver_by_id = {d.id: d for d in drivers}
        for day in range(1, num_days + 1):
            d = date(year, month, day)
            driver_id = request.POST.get(f'driver_{d.isoformat()}', '')
            if driver_id and driver_id.isdigit() and int(driver_id) in driver_by_id:
                VehicleDutySchedule.objects.update_or_create(
                    date=d, defaults={'driver_id': int(driver_id)}
                )
            else:
                VehicleDutySchedule.objects.filter(date=d).delete()

        messages.success(request, f'Đã lưu lịch trực xe tháng {month}/{year}.')
        return redirect(f"{reverse('vehicle_duty_schedule')}?year={year}&month={month}")

    num_days = calendar.monthrange(year, month)[1]
    existing = {
        s.date: s.driver_id
        for s in VehicleDutySchedule.objects.filter(date__year=year, date__month=month)
    }

    weekday_names = ['Th 2', 'Th 3', 'Th 4', 'Th 5', 'Th 6', 'Th 7', 'CN']
    days = []
    for day in range(1, num_days + 1):
        d = date(year, month, day)
        if d in existing:
            selected_driver_id = existing[d]
        elif drivers:
            # Mặc định xen kẽ theo ngày trong tháng, chỉ mang tính gợi ý ban đầu
            selected_driver_id = drivers[(day - 1) % len(drivers)].id
        else:
            selected_driver_id = None
        days.append({
            'date': d,
            'weekday': weekday_names[d.weekday()],
            'selected_driver_id': selected_driver_id,
        })

    prev_month_date = date(year, month, 1) - timedelta(days=1)
    next_month_date = date(year, month, num_days) + timedelta(days=1)

    context = {
        'year': year,
        'month': month,
        'days': days,
        'drivers': drivers,
        'prev_year': prev_month_date.year,
        'prev_month': prev_month_date.month,
        'next_year': next_month_date.year,
        'next_month': next_month_date.month,
    }
    return render(request, 'templates_app/atm/vehicle_duty_schedule.html', context)


@login_required
def atm_travel_claim(request):
    """Trang xem trước Bảng kê thanh toán + Giấy đi đường Ban quản lý ATM theo khoảng ngày"""
    if not request.user.is_superuser:
        messages.error(request, 'Bạn không có quyền truy cập trang này')
        return redirect('dashboard')

    start_date, end_date = _parse_travel_claim_dates(request)

    rows, totals = build_payment_statement_rows(start_date, end_date)
    sheets = build_travel_log_groups(start_date, end_date)

    sheets_summary = {}
    for s in sheets:
        key = (s['machine_id'], s['driver_name'])
        sheets_summary[key] = sheets_summary.get(key, 0) + 1

    context = {
        'start_date': start_date,
        'end_date': end_date,
        'rows': rows,
        'totals': totals,
        'sheets': sheets,
        'sheets_summary': sheets_summary,
        'sheet_count': len(sheets),
        'atm_trip_config': ATM_TRIP_CONFIG,
    }
    return render(request, 'templates_app/atm/travel_claim.html', context)


@login_required
def atm_payment_statement_word(request):
    """Xuất file Word Bảng kê thanh toán công tác phí Ban quản lý ATM"""
    if not request.user.is_superuser:
        messages.error(request, 'Bạn không có quyền truy cập trang này')
        return redirect('dashboard')

    start_date, end_date = _parse_travel_claim_dates(request)

    try:
        from datetime import date as date_cls
        from docx import Document
        from docx.shared import Pt, Cm
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from io import BytesIO
        from .models import num_to_vietnamese_words

        rows, totals = build_payment_statement_rows(start_date, end_date)

        branch_config = BranchConfig.get_for_user(request.user)
        branch_vars = branch_config.get_all_variables()
        dia_danh = branch_vars.get('dia_danh', '') or ''
        ten_chi_nhanh = branch_vars.get('ten_chi_nhanh', '') or ''
        ten_chi_nhanh_hoa = branch_vars.get('ten_chi_nhanh_hoa', '') or ''
        HEADER_TABLE_WIDTH = Cm(18.5)
        LEFT_COL_WIDTH = Cm(8.5)
        RIGHT_COL_WIDTH = Cm(10)

        declarant = get_position_holder('atm_officer', end_date)
        declarant_name = declarant.full_name if declarant else ''

        doc = Document()
        for section in doc.sections:
            section.top_margin = Cm(2)
            section.bottom_margin = Cm(2)
            section.left_margin = Cm(2.5)
            section.right_margin = Cm(2)

        def set_font(run, size=11, bold=False, italic=False, underline=False):
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.italic = italic
            run.font.underline = underline
            run.font.name = 'Times New Roman'

        def add_paragraph(text='', bold=False, size=11, align=WD_ALIGN_PARAGRAPH.LEFT, italic=False,
                           space_before=0, space_after=0):
            p = doc.add_paragraph()
            p.alignment = align
            p.paragraph_format.space_before = Pt(space_before)
            p.paragraph_format.space_after = Pt(space_after)
            run = p.add_run(text)
            set_font(run, size, bold, italic)
            return p

        def cell_paragraph(cell, text, bold=False, size=11, underline=False, italic=False,
                            align=WD_ALIGN_PARAGRAPH.CENTER, first=False):
            p = cell.paragraphs[0] if first else cell.add_paragraph()
            p.alignment = align
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(text)
            set_font(run, size, bold, italic=italic, underline=underline)
            return p

        # Letterhead 2 cột: tên ngân hàng (trái) / quốc hiệu (phải), rộng 18.5cm
        header_table = doc.add_table(rows=1, cols=2)
        header_table.autofit = False
        header_table.width = HEADER_TABLE_WIDTH
        left_cell, right_cell = header_table.rows[0].cells
        left_cell.width = LEFT_COL_WIDTH
        right_cell.width = RIGHT_COL_WIDTH

        cell_paragraph(left_cell, 'NGÂN HÀNG NÔNG NGHIỆP', size=11, first=True)
        cell_paragraph(left_cell, 'VÀ PHÁT TRIỂN NÔNG THÔN VIỆT NAM', size=11)
        if ten_chi_nhanh_hoa:
            cell_paragraph(left_cell, f'CHI NHÁNH {ten_chi_nhanh_hoa}', bold=True, size=11)

        cell_paragraph(right_cell, 'CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM', bold=True, size=12, first=True)
        cell_paragraph(right_cell, 'Độc lập - Tự do - Hạnh phúc', bold=True, size=12, underline=True)

        add_paragraph('BẢNG KÊ THANH TOÁN', bold=True, size=14, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=6, space_after=6)

        atm04_label = ATM_TRIP_CONFIG['7202ATM04']['label']
        atm06_label = ATM_TRIP_CONFIG['7202ATM06']['label']
        period_text = (
            f"từ ngày {start_date.day:02d} tháng {start_date.month:02d} năm {start_date.year} "
            f"đến ngày {end_date.day:02d} tháng {end_date.month:02d} năm {end_date.year}"
        )
        intro = doc.add_paragraph()
        intro.paragraph_format.first_line_indent = Cm(1)
        intro.paragraph_format.space_after = Pt(6)
        r = intro.add_run('Tôi tên: ')
        set_font(r)
        r = intro.add_run(declarant_name)
        set_font(r, bold=True)
        r = intro.add_run(
            f", là cán bộ thuộc phòng Kế toán – Ngân quỹ của Agribank Chi nhánh {ten_chi_nhanh}. "
            f"Nay tôi đề nghị thanh toán tiền phụ cấp trách nhiệm của Ban tiếp quỹ ATM04( {atm04_label} ), "
            f"ATM06 ( {atm06_label} ) {period_text}. Cụ thể như sau ( Kèm GĐĐ ):"
        )
        set_font(r)

        headers = ['TT', 'Họ và tên', 'Chức vụ',
                   'Số\nchuyến', f'ATM\n( Đặt tại\n{atm06_label} )',
                   'Số\nchuyến', f'ATM\n( Đặt tại\n{atm04_label} )',
                   'Số tiền\nTT', 'Tài khoản']
        table = doc.add_table(rows=1, cols=len(headers))
        table.style = 'Table Grid'
        hdr_cells = table.rows[0].cells
        for i, h in enumerate(headers):
            hdr_cells[i].text = h
            hdr_cells[i].paragraphs[0].runs[0].bold = True
            hdr_cells[i].paragraphs[0].runs[0].font.size = Pt(10)
            hdr_cells[i].paragraphs[0].runs[0].font.name = 'Times New Roman'
            hdr_cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

        for idx, row in enumerate(rows, 1):
            cells = table.add_row().cells
            unit_atm06 = ATM_TRIP_CONFIG['7202ATM06']['unit_price']
            unit_atm04 = ATM_TRIP_CONFIG['7202ATM04']['unit_price']
            values = [
                str(idx),
                row['full_name'],
                row['position_display'],
                str(row['trips_atm06']) if row['trips_atm06'] else '',
                f"{unit_atm06:,}" if row['trips_atm06'] else '',
                str(row['trips_atm04']) if row['trips_atm04'] else '',
                f"{unit_atm04:,}" if row['trips_atm04'] else '',
                f"{row['total']:,}",
                row['account_number'],
            ]
            for i, val in enumerate(values):
                cells[i].text = val
                cells[i].paragraphs[0].runs[0].font.size = Pt(10)
                cells[i].paragraphs[0].runs[0].font.name = 'Times New Roman'
                if i in (0, 3, 4, 5, 6, 7):
                    cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

        total_row = table.add_row().cells
        total_row[0].merge(total_row[6])
        total_row[0].text = 'Tổng cộng'
        total_row[0].paragraphs[0].runs[0].bold = True
        total_row[0].paragraphs[0].runs[0].font.size = Pt(10)
        total_row[0].paragraphs[0].runs[0].font.name = 'Times New Roman'
        total_row[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        total_row[7].text = f"{totals['total']:,}"
        total_row[7].paragraphs[0].runs[0].bold = True
        total_row[7].paragraphs[0].runs[0].font.size = Pt(10)
        total_row[7].paragraphs[0].runs[0].font.name = 'Times New Roman'
        total_row[7].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        r = p.add_run('Số tiền bằng chữ: ')
        set_font(r)
        r = p.add_run(num_to_vietnamese_words(totals['total']) + '.')
        set_font(r, bold=True, italic=True)

        add_paragraph(
            f"Đề nghị BGĐ và phòng KTNQ Agribank Chi nhánh {ten_chi_nhanh} thanh toán các khoản tiền nói trên."
        )

        today = date_cls.today()
        date_str = f"{dia_danh}, ngày {today.day:02d} tháng {today.month:02d} năm {today.year}" if dia_danh \
            else f"Ngày {today.day:02d} tháng {today.month:02d} năm {today.year}"
        add_paragraph(date_str, align=WD_ALIGN_PARAGRAPH.RIGHT, space_before=6)

        sig_table = doc.add_table(rows=2, cols=3)
        sig_titles = ['DUYỆT CỦA GIÁM ĐỐC', 'TP.KTNQ', 'NGƯỜI THANH TOÁN']
        for col_idx, title in enumerate(sig_titles):
            cell_paragraph(sig_table.rows[0].cells[col_idx], title, bold=True, first=True)
        sig_table.rows[1].height = Cm(2)

        output = BytesIO()
        doc.save(output)
        output.seek(0)

        response = HttpResponse(
            output.read(),
            content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )
        filename = f"BangKeThanhToan_ATM_{start_date.strftime('%Y%m%d')}_{end_date.strftime('%Y%m%d')}.docx"
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response

    except Exception as e:
        messages.error(request, f"Lỗi khi tạo file Word: {str(e)}")
        return redirect('atm_travel_claim')


@login_required
def atm_travel_log_word(request):
    """Xuất file .zip chứa các tờ Giấy đi đường (mỗi tờ 1 file .docx riêng)"""
    if not request.user.is_superuser:
        messages.error(request, 'Bạn không có quyền truy cập trang này')
        return redirect('dashboard')

    start_date, end_date = _parse_travel_claim_dates(request)

    try:
        from docx import Document
        from docx.shared import Pt, Cm
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.table import WD_ROW_HEIGHT_RULE
        from io import BytesIO
        import zipfile

        sheets = build_travel_log_groups(start_date, end_date)

        if not sheets:
            messages.error(request, 'Không có chuyến tiếp quỹ nào trong khoảng ngày đã chọn')
            return redirect('atm_travel_claim')

        branch_config = BranchConfig.get_for_user(request.user)
        branch_vars = branch_config.get_all_variables()
        dia_danh = branch_vars.get('dia_danh', '') or ''
        ten_chi_nhanh = branch_vars.get('ten_chi_nhanh', '') or ''
        ten_chi_nhanh_hoa = branch_vars.get('ten_chi_nhanh_hoa', '') or ''
        dots = '.' * 24
        HEADER_TABLE_WIDTH = Cm(18.5)
        LEFT_COL_WIDTH = Cm(8.5)
        RIGHT_COL_WIDTH = Cm(10)
        ROW_HEIGHT = Cm(4.52)

        def set_font(run, size=11, bold=False, italic=False, underline=False):
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.italic = italic
            run.font.underline = underline
            run.font.name = 'Times New Roman'

        def build_sheet_doc(sheet):
            doc = Document()
            for section in doc.sections:
                section.top_margin = Cm(2)
                section.bottom_margin = Cm(2)
                section.left_margin = Cm(2.5)
                section.right_margin = Cm(2)

            def add_paragraph(text='', bold=False, size=11, align=WD_ALIGN_PARAGRAPH.LEFT, italic=False,
                               space_before=0, space_after=0):
                p = doc.add_paragraph()
                p.alignment = align
                p.paragraph_format.space_before = Pt(space_before)
                p.paragraph_format.space_after = Pt(space_after)
                run = p.add_run(text)
                set_font(run, size, bold, italic)
                return p

            def cell_paragraph(cell, text, bold=False, size=11, underline=False, italic=False,
                               align=WD_ALIGN_PARAGRAPH.CENTER, first=False):
                p = cell.paragraphs[0] if first else cell.add_paragraph()
                p.alignment = align
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                run = p.add_run(text)
                set_font(run, size, bold, italic=italic, underline=underline)
                return p

            # Letterhead 2 cột: tên ngân hàng (trái) / quốc hiệu (phải), rộng 18.5cm
            header_table = doc.add_table(rows=1, cols=2)
            header_table.autofit = False
            header_table.width = HEADER_TABLE_WIDTH
            left_cell, right_cell = header_table.rows[0].cells
            left_cell.width = LEFT_COL_WIDTH
            right_cell.width = RIGHT_COL_WIDTH

            cell_paragraph(left_cell, 'NGÂN HÀNG NÔNG NGHIỆP', size=11, first=True)
            cell_paragraph(left_cell, 'VÀ PHÁT TRIỂN NÔNG THÔN VIỆT NAM', size=11)
            if ten_chi_nhanh_hoa:
                cell_paragraph(left_cell, f'CHI NHÁNH {ten_chi_nhanh_hoa}', bold=True, size=11)

            cell_paragraph(right_cell, 'CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM', bold=True, size=12, first=True)
            cell_paragraph(right_cell, 'Độc lập - Tự do - Hạnh phúc', bold=True, size=12, underline=True)

            add_paragraph('Số: …./GĐĐ-NHNoGR')
            add_paragraph('GIẤY ĐI ĐƯỜNG', bold=True, size=14, align=WD_ALIGN_PARAGRAPH.CENTER)

            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run('Cấp cho Ông (Bà): ')
            set_font(r)
            for idx, recipient in enumerate(sheet['recipients']):
                if idx > 0:
                    r_sep = p.add_run('; ')
                    set_font(r_sep)
                r_name = p.add_run(recipient['name'])
                set_font(r_name, bold=True)
                r_role = p.add_run(f" ({recipient['role']})")
                set_font(r_role)

            add_paragraph(f'Chức vụ, đơn vị công tác: Ban quản lý ATM Agribank chi nhánh {ten_chi_nhanh}')
            add_paragraph('Theo văn bản cử đi công tác số:')

            first_trip = sheet['trips'][0]
            last_trip = sheet['trips'][-1]
            add_paragraph(
                f"Từ ngày {first_trip.replenishment_date.day:02d} tháng {first_trip.replenishment_date.month:02d} năm {first_trip.replenishment_date.year} "
                f"đến ngày {last_trip.replenishment_date.day:02d} tháng {last_trip.replenishment_date.month:02d} năm {last_trip.replenishment_date.year}",
            )

            # Khối "Địa danh, ngày... / GIÁM ĐỐC" canh giữa với nhau, nằm bên phải trang
            date_str = f"{dia_danh}, Ngày ..... tháng ..... năm {last_trip.replenishment_date.year}" if dia_danh \
                else f"Ngày ..... tháng ..... năm {last_trip.replenishment_date.year}"
            sign_date_table = doc.add_table(rows=1, cols=2)
            sign_date_table.autofit = False
            sign_date_table.width = HEADER_TABLE_WIDTH
            blank_cell, sign_cell = sign_date_table.rows[0].cells
            blank_cell.width = LEFT_COL_WIDTH
            sign_cell.width = RIGHT_COL_WIDTH
            cell_paragraph(sign_cell, date_str, italic=True, size=11, first=True)
            cell_paragraph(sign_cell, 'GIÁM ĐỐC', bold=True, size=11)

            add_paragraph('Tiền ứng trước:')
            add_paragraph(f'Lương:{dots}đ')
            add_paragraph(f'Công tác phí:{dots}đ')
            add_paragraph(f'Cộng:{dots}đ')

            headers = ['Nơi đi, nơi đến', 'Ngày', 'Phương tiện', 'Số ngày\ncông tác', 'Lý do lưu trú', 'Chứng nhận của cơ quan\nnơi đến (ký tên đóng dấu)']
            table = doc.add_table(rows=1, cols=len(headers))
            table.style = 'Table Grid'
            hdr_cells = table.rows[0].cells
            for i, h in enumerate(headers):
                hdr_cells[i].text = h
                hdr_cells[i].paragraphs[0].runs[0].bold = True
                hdr_cells[i].paragraphs[0].runs[0].font.size = Pt(10)
                hdr_cells[i].paragraphs[0].runs[0].font.name = 'Times New Roman'
                hdr_cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

            noi_di = f"Agribank {ten_chi_nhanh}" if ten_chi_nhanh else 'Agribank'
            place = ATM_TRIP_CONFIG[sheet['machine_id']]['destination']
            for row_idx in range(MAX_TRIPS_PER_SHEET):
                row = table.add_row()
                row.height = ROW_HEIGHT
                row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
                cells = row.cells
                trip = sheet['trips'][row_idx] if row_idx < len(sheet['trips']) else None
                if trip:
                    noi_di_noi_den = f"Nơi đi:{noi_di}.\n\nNơi đến : {place}"
                    values = [noi_di_noi_den, trip.replenishment_date.strftime('%d/%m/%Y'), '', '', 'Tiếp quỹ ATM', '']
                else:
                    values = [f"Nơi đi:{noi_di}.\n\nNơi đến : ", '', '', '', '', '']
                for i, val in enumerate(values):
                    cells[i].text = val
                    cells[i].paragraphs[0].runs[0].font.size = Pt(10)
                    cells[i].paragraphs[0].runs[0].font.name = 'Times New Roman'

            add_paragraph('Phần thanh toán ( lập bảng kê chi tiết kèm theo nếu cần thiết )', italic=True)
            add_paragraph(f'1.Tiền chi phí đi lại:{dots}đ')
            add_paragraph(f'2.Tiền phòng ở:{dots}đ')
            add_paragraph(f'3.Phụ cấp lưu trú:{dots}đ')
            add_paragraph(f'4.Phụ cấp trách nhiệm:{dots}đ')

            # Khối "Ngày.../ Duyệt / Số tiền được thanh toán" canh giữa với nhau, nằm bên phải trang
            payout_table = doc.add_table(rows=1, cols=2)
            payout_table.autofit = False
            payout_table.width = HEADER_TABLE_WIDTH
            blank_cell2, payout_cell = payout_table.rows[0].cells
            blank_cell2.width = LEFT_COL_WIDTH
            payout_cell.width = RIGHT_COL_WIDTH
            cell_paragraph(payout_cell, f"Ngày ..... tháng ..... năm {last_trip.replenishment_date.year}", first=True)
            cell_paragraph(payout_cell, 'Duyệt')
            cell_paragraph(payout_cell, 'Số tiền được thanh toán là:' + '.' * 20 + 'đ')

            sig_table = doc.add_table(rows=1, cols=3)
            sig_titles = ['Người đi công tác', 'Phụ trách bộ phận', 'Kế toán trưởng']
            for col_idx, title in enumerate(sig_titles):
                cell_paragraph(sig_table.rows[0].cells[col_idx], title, bold=True, first=True)

            stream = BytesIO()
            doc.save(stream)
            stream.seek(0)
            return stream

        zip_buffer = BytesIO()
        with zipfile.ZipFile(zip_buffer, 'w', zipfile.ZIP_DEFLATED) as zf:
            for sheet in sheets:
                stream = build_sheet_doc(sheet)
                safe_name = remove_vietnamese_diacritics(sheet['driver_name']).replace(' ', '')
                machine_short = sheet['machine_id'].replace('7202', '')
                filename = f"GDD_{machine_short}_{safe_name}_to{sheet['sheet_index']}.docx"
                zf.writestr(filename, stream.read())
        zip_buffer.seek(0)

        response = HttpResponse(zip_buffer.read(), content_type='application/zip')
        zip_filename = f"GiayDiDuong_ATM_{start_date.strftime('%Y%m%d')}_{end_date.strftime('%Y%m%d')}.zip"
        response['Content-Disposition'] = f'attachment; filename="{zip_filename}"'
        return response

    except Exception as e:
        messages.error(request, f"Lỗi khi tạo file Giấy đi đường: {str(e)}")
        return redirect('atm_travel_claim')


@login_required
def atm_load_discrepancy_data(request, discrepancy_id, template_id):
    """Tạo và tải file Word trực tiếp từ giao dịch thừa/thiếu quỹ"""
    if not request.user.is_superuser:
        messages.error(request, 'Bạn không có quyền truy cập trang này')
        return redirect('dashboard')

    # Lấy discrepancy
    discrepancy = get_object_or_404(
        ATMDiscrepancy.objects.select_related('atm', 'created_by'),
        pk=discrepancy_id
    )

    # Lấy template và kiểm tra quyền
    template = get_object_or_404(Template, id=template_id, is_active=True)
    if not template.user_has_access(request.user):
        raise Http404("Bạn không có quyền truy cập mẫu biểu này")

    try:
        # Lấy dữ liệu từ discrepancy
        data = discrepancy.get_data_dict()

        # Thêm biến chung (chi nhánh + custom variables)
        branch_config = BranchConfig.get_for_user(request.user)
        data.update(branch_config.get_all_variables())

        # Render template Word với dữ liệu
        template_path = template.file.path
        output_stream = render_word_template(template_path, data)

        # Trả về file Word để download
        response = HttpResponse(
            output_stream.read(),
            content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )
        disc_type = "Thua" if discrepancy.discrepancy_type == 'surplus' else "Thieu"
        filename = f"{template.name}_ATM_{discrepancy.atm.machine_id}_{disc_type}_{discrepancy.audit_cycle_end.strftime('%Y%m%d')}.docx"
        response['Content-Disposition'] = f'attachment; filename="{filename}"'

        return response

    except Exception as e:
        messages.error(request, f"Lỗi khi tạo file Word: {str(e)}")
        return redirect('atm_discrepancy_list')


@login_required
def atm_load_discrepancy_cycle_template(request, atm_id, start_date, end_date, template_id):
    """Render Word template cho toàn bộ chu kỳ. Mẫu 04: mỗi GD một trang. Mẫu 03: tất cả KH trên 1 trang."""
    if not request.user.is_superuser:
        messages.error(request, 'Bạn không có quyền truy cập trang này')
        return redirect('dashboard')

    from datetime import datetime as dt
    from io import BytesIO
    from copy import deepcopy
    from docx import Document as DocxDocument
    from docx.oxml.ns import qn as oxml_qn
    from docx.oxml import OxmlElement as OxmlEl

    try:
        start_dt = dt.strptime(start_date, '%Y%m%d').date()
        end_dt = dt.strptime(end_date, '%Y%m%d').date()
    except ValueError:
        messages.error(request, 'Ngày không hợp lệ')
        return redirect('atm_discrepancy_list')

    discrepancies = list(ATMDiscrepancy.objects.filter(
        atm__machine_id=atm_id,
        audit_cycle_start=start_dt,
        audit_cycle_end=end_dt,
    ).select_related('atm', 'created_by').order_by('id'))

    if not discrepancies:
        messages.error(request, 'Không tìm thấy nhóm giao dịch này')
        return redirect('atm_discrepancy_list')

    template = get_object_or_404(Template, id=template_id, is_active=True)
    if not template.user_has_access(request.user):
        raise Http404("Bạn không có quyền truy cập mẫu biểu này")

    branch_config = BranchConfig.get_for_user(request.user)

    # Mẫu 04 dùng multi-page (1 trang/GD), Mẫu 03 dùng single-page (thêm dòng KH)
    is_multi_page = 'Mau_04' in template.file.name or '04_THE' in template.file.name

    try:
        if len(discrepancies) == 1:
            data = discrepancies[0].get_data_dict()
            data.update(branch_config.get_all_variables())
            output_stream = render_word_template(template.file.path, data)

        elif is_multi_page:
            # Render từng GD, ghép thành 1 file với page break
            rendered_docs = []
            for disc in discrepancies:
                data = disc.get_data_dict()
                data.update(branch_config.get_all_variables())
                stream = render_word_template(template.file.path, data)
                rendered_docs.append(DocxDocument(stream))

            combined = rendered_docs[0]
            for next_doc in rendered_docs[1:]:
                pg_p = OxmlEl('w:p')
                pg_r = OxmlEl('w:r')
                pg_br = OxmlEl('w:br')
                pg_br.set(oxml_qn('w:type'), 'page')
                pg_r.append(pg_br)
                pg_p.append(pg_r)

                body = combined.element.body
                sect_pr = body.find(oxml_qn('w:sectPr'))
                if sect_pr is not None:
                    sect_pr.addprevious(pg_p)
                else:
                    body.append(pg_p)

                for elem in next_doc.element.body:
                    if elem.tag.endswith('}sectPr'):
                        continue
                    new_elem = deepcopy(elem)
                    if sect_pr is not None:
                        sect_pr.addprevious(new_elem)
                    else:
                        body.append(new_elem)

            output_stream = BytesIO()
            combined.save(output_stream)
            output_stream.seek(0)

        else:
            # Mẫu 03: render với GD đầu tiên, chèn thêm dòng KH cho các GD sau
            data = discrepancies[0].get_data_dict()
            data.update(branch_config.get_all_variables())
            # Nếu có nhiều GD: ghi đè disc_amount và disc_amount_words bằng TỔNG — chỉ ảnh
            # hưởng ô "Số tiền bằng số/chữ" ở bảng trên; dòng "KH: ..." của từng GD (kể cả
            # GD đầu tiên) luôn được ghi đè lại bên dưới bằng số tiền RIÊNG của GD đó, không
            # dùng chung placeholder disc_amount để tránh GD đầu tiên bị hiện nhầm thành tổng.
            if len(discrepancies) > 1:
                from .models import num_to_vietnamese_words
                total = sum(int(d.amount) for d in discrepancies)
                data['disc_amount'] = f"{total:,}"
                data['disc_amount_words'] = num_to_vietnamese_words(total)
            stream = render_word_template(template.file.path, data)

            doc = DocxDocument(stream)

            def _kh_line_text(disc):
                parts = [f"KH: {disc.full_name} STK: {disc.account_number}"]
                if disc.card_number:
                    parts.append(f"SO THE: {disc.card_number}")
                if disc.transaction_id:
                    parts.append(f"ID: {disc.transaction_id}")
                if disc.trace_number:
                    parts.append(f"TRACE: {disc.trace_number}")
                parts.append(f"Số tiền: {disc.amount:,}đ")
                return ", ".join(parts)

            def _set_p_text(p_elem, text):
                runs = p_elem.findall(oxml_qn('w:r'))
                if not runs:
                    return
                t_el = runs[0].find(oxml_qn('w:t'))
                if t_el is None:
                    t_el = OxmlEl('w:t')
                    runs[0].append(t_el)
                t_el.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
                t_el.text = text
                for r in runs[1:]:
                    p_elem.remove(r)

            kh_para = None
            for para in doc.paragraphs:
                if para.text.strip().startswith('KH:'):
                    kh_para = para
                    break

            if kh_para is not None:
                _set_p_text(kh_para._p, _kh_line_text(discrepancies[0]))
                last_p = kh_para._p
                for disc in discrepancies[1:]:
                    new_p = deepcopy(kh_para._p)
                    _set_p_text(new_p, _kh_line_text(disc))
                    last_p.addnext(new_p)
                    last_p = new_p

            output_stream = BytesIO()
            doc.save(output_stream)
            output_stream.seek(0)

        response = HttpResponse(
            output_stream.read(),
            content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )
        filename = f"{template.name}_ATM_{atm_id}_{end_dt.strftime('%Y%m%d')}.docx"
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response

    except Exception as e:
        messages.error(request, f"Lỗi khi tạo file Word: {str(e)}")
        return redirect('atm_discrepancy_list')
