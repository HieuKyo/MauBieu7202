"""
Xuất danh sách khách hàng ra Excel (tab Khách hàng)
1. Excel thường: các thông tin chính của khách hàng
2. Theo mẫu "Mẫu mở CIF - dịch vụ đồng thời (Khách hàng mới)_V7.1-1": chỉ điền các ô nhập liệu,
   giữ nguyên STT, cột công thức (mã tra từ danh mục), danh sách chọn và các sheet danh mục của file gốc.
"""
import io
import re
import unicodedata
from datetime import date, datetime
from pathlib import Path

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.db.models import Q
from django.http import HttpResponse
from django.shortcuts import redirect

from .models import Customer

CIF_TEMPLATE = Path(__file__).resolve().parent / 'export_templates' / 'mau_mo_cif_V7.1-1.xlsx'
CIF_SHEET = 'DANG_KY_MOI'
CIF_FIRST_ROW = 4
CIF_LAST_ROW = 503          # file mẫu có sẵn công thức / danh sách chọn đến dòng 503 (500 khách)
CIF_LAST_COL = 51           # cột AY - cột cuối của mẫu

# Nghề nghiệp trên form → danh mục NGHE_NGHIEP của mẫu CIF ('' = không rõ, để GĐV chọn trong file)
NGHE_NGHIEP_CIF = {
    'Công chức viên chức': 'Công chức/Viên chức',
    'Nông dân': 'Nông dân',
    'Giáo viên': 'Giáo viên',
    'Bác sĩ': 'Bác sỹ',
    'Công nhân': 'Công nhân',
    'Kinh doanh tự do': 'Kinh doanh tự do',
    'Nội trợ': 'Nội trợ',
    'Kỹ sư': 'Kĩ sư',
    'Khác': 'Khác',
}

# Nơi cấp trên form → tên trong sheet NOI_CAP của mẫu (mã lấy từ chính file mẫu)
NOI_CAP_CIF = {
    'Cục CSQLHC về TTXH': 'Cục CSQLHC về TTXH',
    'Bộ Công An': 'Bộ Công An',                     # Căn cước mẫu mới (từ 01/07/2024) → mã 004
}


def _filtered_customers(request):
    """Khách hàng theo ô tìm kiếm hiện tại + khoảng ngày cập nhật (ngày in/lưu gần nhất)."""
    qs = Customer.objects.all()
    q = request.GET.get('q', '').strip()
    if q:
        qs = qs.filter(Q(ho_ten__icontains=q) | Q(so_cmnd__icontains=q) |
                       Q(so_dien_thoai__icontains=q) | Q(email__icontains=q) | Q(ghi_chu__icontains=q))
    for param, lookup in (('from', 'updated_at__date__gte'), ('to', 'updated_at__date__lte')):
        value = request.GET.get(param, '').strip()
        if value:
            try:
                qs = qs.filter(**{lookup: datetime.strptime(value, '%Y-%m-%d').date()})
            except ValueError:
                pass
    return qs.order_by('updated_at', 'id')


def _plain(text):
    """Bỏ dấu tiếng Việt (giữ chữ Đ → D)."""
    text = (text or '').replace('đ', 'd').replace('Đ', 'D')
    return ''.join(c for c in unicodedata.normalize('NFD', text) if unicodedata.category(c) != 'Mn')


def _norm_place(name):
    """So khớp tên địa danh: bỏ dấu, chữ thường, bỏ tiền tố Tỉnh/Thành phố/Phường/Xã..."""
    s = re.sub(r'\s+', ' ', _plain(name).lower()).strip()
    return re.sub(r'^(tinh|thanh pho|tp\.?|phuong|xa|dac khu|thi tran)\s+', '', s)


def _split_35(text):
    """Tách địa chỉ thành 2 ô tối đa 35 ký tự (cắt ở khoảng trắng)."""
    text = (text or '').strip()
    if len(text) <= 35:
        return text, ''
    cut = text.rfind(' ', 0, 36)
    cut = cut if cut > 0 else 35
    return text[:cut].rstrip(' ,'), text[cut:].strip(' ,')[:35]


def _ymd(d):
    return d.strftime('%Y%m%d') if d else ''


def _age(born):
    if not born:
        return None
    today = date.today()
    return today.year - born.year - ((today.month, today.day) < (born.month, born.day))


@login_required
def customer_export_view(request):
    kind = request.GET.get('kind', 'normal')
    customers = list(_filtered_customers(request))
    if not customers:
        messages.warning(request, 'Không có khách hàng nào phù hợp để xuất Excel.')
        return redirect('customer_list')
    if kind == 'cif':
        return _export_cif(request, customers)
    return _export_normal(customers)


def _xlsx_response(wb, filename):
    buf = io.BytesIO()
    wb.save(buf)
    response = HttpResponse(buf.getvalue(),
                            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


# ----------------------------------------------------------------------------- 1. Excel thường
def _class_name(note):
    """Lớp của học sinh trong ghi chú dạng 'Lớp: 10A1 - Trường: ...' (do Nhập DS học sinh ghi)."""
    m = re.search(r'^Lớp:\s*(.*?)(?:\s+-\s+Trường:.*)?$', note or '', re.M)
    return m.group(1).strip() if m else ''


def _export_normal(customers):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    yn = lambda v: 'Có' if v else ''
    dmy = lambda d: d.strftime('%d/%m/%Y') if d else ''
    columns = [
        ('STT', None), ('Mã khách hàng', lambda c: c.ma_khach_hang), ('Mã CIF', lambda c: c.cif),
        ('Họ và tên', lambda c: c.ho_ten), ('Lớp', lambda c: _class_name(c.ghi_chu)), ('Ngày sinh', lambda c: dmy(c.ngay_sinh)),
        ('Ghi chú tuổi', lambda c: (c.get_age_note() or ('', ''))[1]), ('Giới tính', lambda c: c.gioi_tinh),
        ('Số CCCD', lambda c: c.so_cmnd), ('Ngày cấp', lambda c: dmy(c.ngay_cap_cmnd)),
        ('Nơi cấp', lambda c: c.noi_cap_cmnd), ('Ngày hết hạn', lambda c: dmy(c.ngay_het_han_cmnd)),
        ('Địa chỉ', lambda c: c.dia_chi), ('Số điện thoại', lambda c: c.so_dien_thoai), ('Email', lambda c: c.email),
        ('Nghề nghiệp', lambda c: c.nghe_nghiep), ('Nơi làm việc', lambda c: c.noi_lam_viec),
        ('Số tài khoản', lambda c: c.so_tai_khoan), ('Loại tài khoản', lambda c: c.loai_tai_khoan),
        ('Loại tiền', lambda c: c.loai_tien_te), ('Loại thẻ', lambda c: c.loai_the), ('Hạng thẻ', lambda c: c.hang_the),
        ('Số thẻ', lambda c: c.so_the_atm), ('SMS Banking', lambda c: yn(c.dv_sms_banking)),
        ('Agribank Plus', lambda c: yn(c.dv_bankplus)), ('E-Mobile', lambda c: yn(c.dv_e_mobile)),
        ('Bảo hiểm ABIC', lambda c: yn(c.dv_abic)), ('Phân loại KH', lambda c: c.ket_qua_phan_loai_kh),
        ('Ngày in mẫu', lambda c: dmy(c.ngay_in)),
        ('Cập nhật lúc', lambda c: c.updated_at.strftime('%d/%m/%Y %H:%M') if c.updated_at else ''),
    ]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Khách hàng'
    head_fill = PatternFill('solid', fgColor='AE1C3F')
    for i, (title, _) in enumerate(columns, 1):
        cell = ws.cell(1, i, title)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    # Tô cả dòng theo tuổi tính đến ngày tải file: vàng = dưới 15 tuổi, xanh lá = lúc lưu dưới 15 nay đã đủ tuổi
    age_fill = {'under': PatternFill('solid', fgColor='FFF2CC'), 'reached': PatternFill('solid', fgColor='C6EFCE')}
    for r, c in enumerate(customers, 2):
        ws.cell(r, 1, r - 1)
        for i, (_, get) in enumerate(columns[1:], 2):
            ws.cell(r, i, get(c) or '').number_format = '@'   # giữ số 0 đầu (CCCD, SĐT, số TK)
        note = c.get_age_note()
        if note:
            for i in range(1, len(columns) + 1):
                ws.cell(r, i).fill = age_fill[note[0]]
    for i, (title, _) in enumerate(columns, 1):
        width = max([len(str(title))] + [len(str(ws.cell(r, i).value or '')) for r in range(2, min(ws.max_row, 200) + 1)])
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 8), 50)
    ws.freeze_panes = 'E2'
    ws.auto_filter.ref = ws.dimensions
    return _xlsx_response(wb, f'khach_hang_{datetime.now():%Y%m%d_%H%M}.xlsx')


# ----------------------------------------------------------------------------- 2. Mẫu mở CIF
def _export_cif(request, customers):
    import openpyxl

    if len(customers) > CIF_LAST_ROW - CIF_FIRST_ROW + 1:
        messages.error(request, f'Mẫu mở CIF chỉ chứa tối đa {CIF_LAST_ROW - CIF_FIRST_ROW + 1} khách hàng '
                                f'(đang chọn {len(customers)}). Vui lòng thu hẹp khoảng ngày hoặc tìm kiếm.')
        return redirect('customer_list')

    wb = openpyxl.load_workbook(CIF_TEMPLATE)
    ws = wb[CIF_SHEET]

    # Danh mục lấy từ chính file mẫu
    tinh = {_norm_place(r[1]): str(r[0]) for r in wb['TINH_TP'].iter_rows(min_row=2, values_only=True) if r[0] and r[1]}
    phuong = {}
    for code, name, city in wb['PHUONG_XA'].iter_rows(min_row=2, max_col=3, values_only=True):
        if code and name:
            phuong[(str(city), _norm_place(name))] = str(code)
    noi_cap = {str(r[1]).strip(): str(r[0]) for r in wb['NOI_CAP'].iter_rows(min_row=2, values_only=True) if r[0] and r[1]}
    loai_the = {r[0]: r[1] for r in wb['LOAI_THE'].iter_rows(min_row=2, values_only=True) if r[0]}

    # Xóa dữ liệu ví dụ: chỉ các ô nhập tay (không phải công thức) từ cột B đến AY; giữ STT (cột A)
    for row in ws.iter_rows(min_row=CIF_FIRST_ROW, max_row=CIF_LAST_ROW, min_col=2, max_col=CIF_LAST_COL):
        for cell in row:
            if cell.value is not None and not (isinstance(cell.value, str) and cell.value.startswith('=')):
                cell.value = None

    def put(r, col, value):
        cell = ws[f'{col}{r}']
        if isinstance(cell.value, str) and cell.value.startswith('='):
            return   # an toàn: không bao giờ ghi đè cột công thức của mẫu
        cell.value = value if value != '' else None

    for i, c in enumerate(customers):
        r = CIF_FIRST_ROW + i
        ho_ten = (c.ho_ten or '').strip()
        ten_anh = (c.ho_ten_tieng_anh or _plain(ho_ten)).upper().strip()

        # Địa chỉ mới dạng "..., Phường/Xã X, Tỉnh Y" → mã Tỉnh/Phường; phần còn lại vào Địa chỉ 1, 2
        parts = [p.strip() for p in (c.dia_chi or '').split(',') if p.strip()]
        ma_tinh = tinh.get(_norm_place(parts[-1])) if parts else None
        ma_phuong = phuong.get((ma_tinh, _norm_place(parts[-2]))) if ma_tinh and len(parts) >= 2 else None
        dia_chi_1 = ', '.join(parts[:-2]) if ma_phuong else ', '.join(parts)
        dc1, dc2 = _split_35(dia_chi_1)

        nghe = NGHE_NGHIEP_CIF.get(c.nghe_nghiep or '', '')
        if c.nghe_nghiep == 'Học sinh/Sinh viên':
            age = _age(c.ngay_sinh)
            nghe = 'Học sinh' if age is not None and age < 18 else 'Sinh viên'

        phat_hanh = bool(c.loai_the)
        the = ''
        if c.loai_the == 'Thẻ Ghi nợ nội địa':
            the = loai_the.get('PSuccess' if c.hang_the == 'Hạng vàng' else 'Success', '')
        dang_ky_tktt = bool(c.loai_tai_khoan or c.so_tai_khoan or c.dv_sms_banking or c.dv_bankplus or phat_hanh)

        values = {
            'B': ho_ten[:50], 'C': ho_ten[:35], 'D': ten_anh[:50], 'E': ten_anh[:35],
            'F': _ymd(c.ngay_sinh),
            'H': c.gioi_tinh if c.gioi_tinh in ('Nam', 'Nữ') else '',
            'I': re.sub(r'\D', '', c.so_dien_thoai or '')[:11],
            'K': 'VIETNAM', 'L': 'Y', 'N': 'Cá nhân', 'P': 'Tư nhân', 'R': nghe,
            'T': 'Chứng minh thư/Thẻ căn cước công dân', 'U': (c.so_cmnd or '').strip(),
            'W': _ymd(c.ngay_cap_cmnd), 'X': _ymd(c.ngay_het_han_cmnd), 'Y': 'N' if c.ngay_het_han_cmnd else '',
            'Z': noi_cap.get(NOI_CAP_CIF.get(c.noi_cap_cmnd or '', ''), ''),
            'AC': 'Nơi ở', 'AE': 'VIETNAM', 'AF': ma_tinh or '', 'AH': ma_phuong or '',
            'AJ': dc1, 'AK': dc2, 'AM': (c.email or '').strip(),
            'AN': 'Y' if dang_ky_tktt else 'N', 'AP': (c.ghi_chu or '')[:200],
            'AQ': 'Y' if c.dv_sms_banking else 'N', 'AR': 'Y' if c.dv_bankplus else 'N',
            'AS': 'Y' if phat_hanh else 'N',
        }
        if phat_hanh:
            values.update({'AU': the, 'AV': 'SITE', 'AW': 'NORMAL', 'AX': 1 if c.dv_abic else 0})
        for col, value in values.items():
            put(r, col, value)

    return _xlsx_response(wb, f'Mau_mo_CIF_{datetime.now():%Y%m%d_%H%M}.xlsx')
