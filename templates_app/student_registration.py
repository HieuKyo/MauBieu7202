"""
Học sinh tự đăng ký (trang web online → Google Sheet, dữ liệu MÃ HÓA) → nhập vào hệ thống LAN → in hồ sơ theo lô.

Mã hóa (trang web, WebCrypto):  AES-256-GCM cho dữ liệu, khóa AES mã hóa bằng RSA-OAEP-SHA256 (khóa công khai).
Mỗi hồ sơ là 1 chuỗi:  HS1:<khóa AES đã mã hóa>:<iv>:<dữ liệu đã mã hóa>   (base64)
Khóa bí mật chỉ nằm trên máy chủ LAN: <BASE_DIR>/private_media/keys/hoc_sinh_private.pem (không đưa lên git).
"""
import base64
import csv
import io
import json
import re
import unicodedata
import zipfile
from copy import deepcopy
from datetime import date, datetime
from pathlib import Path

from django.conf import settings
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse, JsonResponse
from django.shortcuts import redirect, render
from django.views.decorators.http import require_POST

from .cccd_utils import convert_address, parse_cccd_qr
from .models import BranchConfig, Customer, Template

TOKEN_RE = re.compile(r'HS1:[A-Za-z0-9+/=]+:[A-Za-z0-9+/=]+:[A-Za-z0-9+/=]+')
STUDENT_TEMPLATE_NAME = 'Phát hành thẻ Học sinh (Mẫu 1A)'
MAX_BATCH = 300


# ============================================================ Khóa
def _key_dir():
    return Path(getattr(settings, 'STUDENT_KEY_DIR', Path(settings.BASE_DIR) / 'private_media' / 'keys'))


def _private_key_path():
    return _key_dir() / 'hoc_sinh_private.pem'


def load_private_key():
    from cryptography.hazmat.primitives import serialization
    path = _private_key_path()
    if not path.exists():
        return None
    return serialization.load_pem_private_key(path.read_bytes(), password=None)


def public_key_b64(private_key):
    """Khóa công khai dạng SPKI (DER) base64 — dán vào trang web (WebCrypto importKey 'spki')."""
    from cryptography.hazmat.primitives import serialization
    der = private_key.public_key().public_bytes(serialization.Encoding.DER, serialization.PublicFormat.SubjectPublicKeyInfo)
    return base64.b64encode(der).decode()


@login_required
def student_keys_view(request):
    """Trang quản lý khóa (chỉ Superuser): tạo khóa 1 lần, xem khóa công khai để dán vào trang web."""
    if not request.user.is_superuser:
        messages.error(request, 'Chỉ quản trị viên được quản lý khóa đăng ký học sinh.')
        return redirect('customer_list')
    key = load_private_key()
    if request.method == 'POST' and request.POST.get('action') == 'create':
        if key is not None and request.POST.get('confirm') != 'TAO LAI':
            messages.error(request, 'Đã có khóa. Muốn tạo lại phải gõ "TAO LAI" — hồ sơ gửi bằng khóa cũ sẽ KHÔNG giải mã được nữa.')
        else:
            from cryptography.hazmat.primitives import serialization
            from cryptography.hazmat.primitives.asymmetric import rsa
            key = rsa.generate_private_key(public_exponent=65537, key_size=3072)
            _key_dir().mkdir(parents=True, exist_ok=True)
            _private_key_path().write_bytes(key.private_bytes(
                serialization.Encoding.PEM, serialization.PrivateFormat.PKCS8, serialization.NoEncryption()))
            messages.success(request, 'Đã tạo khóa mới. Dán khóa công khai bên dưới vào trang web đăng ký.')
        return redirect('student_keys')
    return render(request, 'templates_app/student_keys.html', {
        'public_key': public_key_b64(key) if key else '',
        'key_path': str(_private_key_path()),
    })


def decrypt_token(private_key, token):
    from cryptography.hazmat.primitives import hashes
    from cryptography.hazmat.primitives.asymmetric import padding
    from cryptography.hazmat.primitives.ciphers.aead import AESGCM
    _, k, iv, c = token.split(':')
    aes_key = private_key.decrypt(base64.b64decode(k), padding.OAEP(
        mgf=padding.MGF1(algorithm=hashes.SHA256()), algorithm=hashes.SHA256(), label=None))
    return json.loads(AESGCM(aes_key).decrypt(base64.b64decode(iv), base64.b64decode(c), None).decode('utf-8'))


# ============================================================ Nhập danh sách
def _tokens_from_file(upload):
    """Mọi chuỗi HS1:... trong file tải về: .zip / .json từ trang quản trị của host,
    hoặc Excel (.xlsx) / CSV từ Google Sheet."""
    def from_json(raw):
        # Đọc JSON thật (PHP có thể ghi "/" thành "\/"); file 1 hồ sơ {"data": ...} hoặc danh sách nhiều hồ sơ
        text = raw.decode('utf-8-sig', errors='ignore')
        try:
            obj = json.loads(text)
            items = obj if isinstance(obj, list) else [obj]
            return [m for it in items if isinstance(it, dict) for m in TOKEN_RE.findall(str(it.get('data', '')))]
        except ValueError:
            return TOKEN_RE.findall(text.replace('\\/', '/'))

    name = upload.name.lower()
    if name.endswith('.zip'):
        with zipfile.ZipFile(upload) as zf:
            return [m for n in zf.namelist() if n.lower().endswith('.json') for m in from_json(zf.read(n))]
    if name.endswith('.json'):
        return from_json(upload.read())
    if name.endswith(('.xlsx', '.xlsm')):
        import openpyxl
        wb = openpyxl.load_workbook(upload, read_only=True, data_only=True)
        cells = (str(v) for ws in wb.worksheets for row in ws.iter_rows(values_only=True) for v in row if v)
        tokens = [m for v in cells for m in TOKEN_RE.findall(v)]
        wb.close()
        return tokens
    text = upload.read().decode('utf-8-sig', errors='ignore')
    return [m for row in csv.reader(io.StringIO(text)) for v in row for m in TOKEN_RE.findall(v)]


def _student_preset():
    t = Template.objects.filter(name=STUDENT_TEMPLATE_NAME).first()
    return dict(t.preset_values) if t and t.preset_values else {}


def _set_class_note(note, lop, truong):
    """Ghi chú: thay dòng 'Lớp: ...' cũ (nếu có), giữ các ghi chú khác."""
    line = f'Lớp: {lop}' + (f' - Trường: {truong}' if truong else '')
    others = [x for x in (note or '').splitlines() if x.strip() and not x.startswith('Lớp:')]
    return '\n'.join([line] + others)


@login_required
@require_POST
def student_import_view(request):
    """Xem trước (apply=0) / nhập (apply=1) hồ sơ học sinh tự đăng ký (file .zip từ host, hoặc .xlsx/.csv từ Google Sheet).
    Hồ sơ có ảnh 2 mặt thẻ → lưu ảnh theo khách hàng; không đọc được QR → liệt kê riêng kèm ảnh để nhập tay."""
    from .customer_cccd_images import save_cccd_images, thumbnail_data_url
    upload = request.FILES.get('data_file')
    if not upload or not upload.name.lower().endswith(('.zip', '.json', '.xlsx', '.xlsm', '.csv')):
        return JsonResponse({'success': False, 'error': 'Vui lòng chọn file .zip tải từ trang quản trị (hoặc .xlsx / .csv từ Google Sheet)'}, status=400)
    key = load_private_key()
    if key is None:
        return JsonResponse({'success': False, 'error': 'Chưa có khóa giải mã trên máy chủ (Quản trị → Khóa đăng ký học sinh).'}, status=400)
    tokens = _tokens_from_file(upload)
    if not tokens:
        return JsonResponse({'success': False, 'error': 'Không tìm thấy hồ sơ mã hóa (chuỗi bắt đầu bằng HS1:) trong file.'}, status=400)

    apply = request.POST.get('apply') == '1'
    use_suggestion = request.POST.get('use_suggestion') == '1'

    # Giải mã + đọc QR; cùng 1 CCCD gửi nhiều lần → lấy lần gửi sau cùng
    records, no_qr, errors = {}, [], 0
    for tok in tokens:
        try:
            payload = decrypt_token(key, tok)
        except Exception:
            errors += 1
            continue
        extra = dict(sdt=re.sub(r'\D', '', str(payload.get('sdt', '')))[:15], lop=str(payload.get('lop', '')).strip()[:30],
                     truong=str(payload.get('truong', '')).strip()[:100], gui_luc=str(payload.get('t', '')),
                     co_anh=bool(payload.get('anh_truoc') or payload.get('anh_sau')), tok=tok)
        info = parse_cccd_qr(payload.get('qr', ''))
        if not info:
            if extra['co_anh']:
                # Có ảnh nhưng không đọc được QR → cán bộ xem ảnh và nhập tay
                no_qr.append({'ho_ten': str(payload.get('ho_ten', '')).strip()[:100], 'lop': extra['lop'], 'truong': extra['truong'],
                              'sdt': extra['sdt'], 'gui_luc': extra['gui_luc'],
                              'anh': [{'mat': label, 'thumb': thumbnail_data_url(payload.get(side, '')),
                                       'full': thumbnail_data_url(payload.get(side, ''), 2000)}
                                      for side, label in (('anh_truoc', 'Mặt trước'), ('anh_sau', 'Mặt sau')) if payload.get(side)]})
            else:
                errors += 1
            continue
        info.update(extra)
        old = records.get(info['so_cmnd'])
        if old is None or info['gui_luc'] >= old['gui_luc']:
            records[info['so_cmnd']] = info

    preset = _student_preset()
    existing = {c.so_cmnd: c for c in Customer.objects.filter(so_cmnd__in=list(records))}
    rows, saved_ids = [], []
    for cccd, r in sorted(records.items(), key=lambda kv: (kv[1]['truong'], kv[1]['lop'], kv[1]['ho_ten'])):
        conv = convert_address(r['dia_chi_tren_the'])
        if conv and conv[1]:
            dia_chi, addr_status, goi_y = conv[0], 'converted', ''
        elif conv:
            dia_chi, addr_status, goi_y = (conv[0], 'suggest_used', conv[0]) if use_suggestion else (r['dia_chi_tren_the'], 'suggest', conv[0])
        else:
            dia_chi, addr_status, goi_y = r['dia_chi_tren_the'], 'none', ''
        c = existing.get(cccd)
        row = {'cccd': cccd, 'ho_ten': r['ho_ten'], 'ngay_sinh': r['ngay_sinh'].strftime('%d/%m/%Y') if r['ngay_sinh'] else '',
               'tuoi': (lambda b: date.today().year - b.year - ((date.today().month, date.today().day) < (b.month, b.day)))(r['ngay_sinh']) if r['ngay_sinh'] else None,
               'sdt': r['sdt'], 'lop': r['lop'], 'truong': r['truong'], 'dia_chi': dia_chi,
               'dia_chi_tren_the': r['dia_chi_tren_the'], 'goi_y': goi_y, 'dia_chi_status': addr_status,
               'trang_thai': 'update' if c else 'new', 'loi': '', 'co_anh': r['co_anh']}
        if apply:
            try:
                if c is None:
                    c = Customer(created_by=request.user)
                    for field, value in preset.items():
                        if hasattr(c, field) and not isinstance(value, (list, dict)):
                            setattr(c, field, value)
                for field in ('so_cmnd', 'ho_ten', 'ngay_sinh', 'gioi_tinh', 'ngay_cap_cmnd', 'ngay_het_han_cmnd', 'noi_cap_cmnd'):
                    setattr(c, field, r[field])
                c.dia_chi = dia_chi
                if r['sdt']:
                    c.so_dien_thoai = r['sdt']
                if r['lop']:
                    c.ghi_chu = _set_class_note(c.ghi_chu, r['lop'], r['truong'])
                c.save()
                saved_ids.append(c.pk)
            except Exception as e:
                from .customer_views import _customer_error_text
                row['loi'] = _customer_error_text(e)
            else:
                if r['co_anh']:
                    try:
                        payload = decrypt_token(key, r['tok'])   # giải mã lại để lấy ảnh (không giữ ảnh của cả lô trong bộ nhớ)
                        save_cccd_images(c.pk, {'mat_truoc': payload.get('anh_truoc'), 'mat_sau': payload.get('anh_sau')})
                    except Exception:
                        row['loi'] = 'Đã lưu khách hàng nhưng không lưu được ảnh thẻ'
        rows.append(row)

    summary = {'new': sum(r['trang_thai'] == 'new' for r in rows), 'update': sum(r['trang_thai'] == 'update' for r in rows),
               'suggest': sum(r['dia_chi_status'] == 'suggest' for r in rows), 'loi_giai_ma': errors,
               'trung': len(tokens) - errors - len(records) - len(no_qr), 'loi_luu': sum(bool(r['loi']) for r in rows),
               'co_anh': sum(r['co_anh'] for r in rows), 'khong_qr': len(no_qr)}
    return JsonResponse({'success': True, 'applied': apply, 'rows': rows, 'summary': summary, 'saved_ids': saved_ids,
                         'no_qr': no_qr})


# ============================================================ In theo lô
def _merge_docx(streams):
    """Gộp nhiều file Word (cùng 1 mẫu) thành 1 file, mỗi hồ sơ bắt đầu ở TRANG LẺ (ngắt section oddPage):
    in 2 mặt thì hồ sơ sau luôn bắt đầu ở tờ mới, Word tự thêm trang trắng khi hồ sơ trước có số trang lẻ.
    Cùng mẫu nên ảnh / kiểu chữ / đánh số dùng chung id → chép nội dung thân bài là đủ."""
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    def start_on_odd_page(sect_pr):
        t = sect_pr.find(qn('w:type'))
        if t is None:
            t = OxmlElement('w:type')
            pg_sz = sect_pr.find(qn('w:pgSz'))      # w:type phải đứng trước w:pgSz theo thứ tự của Word
            pg_sz.addprevious(t) if pg_sz is not None else sect_pr.append(t)
        t.set(qn('w:val'), 'oddPage')

    master = Document(streams[0])
    body = master.element.body
    sect = body.find(qn('w:sectPr'))
    # Ảnh (wp:docPr), bookmark, ô nhập liệu (w:sdt) có mã id phải duy nhất — trùng là Word báo file hỏng
    doc_pr_id = max([int(d.get('id')) for d in body.iter(qn('wp:docPr'))] or [0])
    for i, s in enumerate(streams[1:], start=1):
        # Kết thúc hồ sơ trước bằng ngắt section (gắn vào đoạn cuối, không thêm đoạn trống dễ đẩy sang trang thừa)
        last = sect.getprevious()
        if last is None or last.tag != qn('w:p') or last.find(f"{qn('w:pPr')}/{qn('w:sectPr')}") is not None:
            last = OxmlElement('w:p')
            sect.addprevious(last)
        p_pr = last.find(qn('w:pPr'))
        if p_pr is None:
            p_pr = OxmlElement('w:pPr')
            last.insert(0, p_pr)
        p_pr.append(deepcopy(sect))
        first_sect = None
        for el in Document(s).element.body:
            if el.tag == qn('w:sectPr'):
                continue
            el = deepcopy(el)
            if first_sect is None:
                first_sect = next(el.iter(qn('w:sectPr')), None)
            for d in el.iter(qn('wp:docPr')):
                doc_pr_id += 1
                d.set('id', str(doc_pr_id))
            for b in el.iter(qn('w:bookmarkStart'), qn('w:bookmarkEnd')):
                b.set(qn('w:id'), str(int(b.get(qn('w:id'))) + 100000 * i))
            for sdt_id in el.iter(qn('w:id')):
                sdt_id.set(qn('w:val'), str((int(sdt_id.get(qn('w:val'))) + 7919 * i) % 2147483647))
            sect.addprevious(el)
        start_on_odd_page(first_sect if first_sect is not None else sect)
    out = io.BytesIO()
    master.save(out)
    return out.getvalue()


@login_required
def student_batch_print_view(request):
    """1 file Word gộp hồ sơ của nhiều khách (theo ids=..., hoặc theo tìm kiếm + khoảng ngày), dùng 1 mẫu biểu."""
    from .customer_cccd_images import cccd_image_sides
    from .customer_export import _class_name, _filtered_customers, _plain
    from .utils import render_word_template
    from .views import fill_empty_print_fields

    ids = [int(x) for x in request.GET.get('ids', '').split(',') if x.strip().isdigit()]
    keep_order = bool(ids) and request.GET.get('keep_order') == '1'
    customers = list(Customer.objects.filter(pk__in=ids).order_by('ghi_chu', 'ho_ten') if ids else _filtered_customers(request))
    if keep_order:
        # Danh sách đã xem trước và bỏ tick: in đúng thứ tự đã xem
        pos = {pk: i for i, pk in enumerate(ids)}
        customers.sort(key=lambda c: pos[c.pk])
    elif request.GET.get('sort') == 'lop':
        # Theo lớp (so số theo giá trị: 10A2 trước 10A10), trong lớp theo Tên rồi Họ như sổ lớp; không có lớp xếp cuối
        def vn_key(text):
            """Thứ tự chữ cái tiếng Việt: A Ă Â B C D Đ E Ê … O Ô Ơ … U Ư (bỏ qua dấu thanh)."""
            key = []
            for ch in text.upper():
                base, *marks = unicodedata.normalize('NFD', ch)
                mod = 1 if '̆' in marks else 2 if '̂' in marks else 3 if '̛' in marks else 0
                key.append(('D', 1) if ch == 'Đ' else (base, mod))
            return key

        def class_key(c):
            lop = _class_name(c.ghi_chu).upper()
            parts = [(0, int(p), '') if p.isdigit() else (1, 0, p) for p in re.findall(r'\d+|\D+', lop)]
            name = (c.ho_ten or '').split()
            return (not lop, parts, vn_key(name[-1] if name else ''), vn_key(' '.join(name)))
        customers.sort(key=class_key)
    if request.GET.get('preview') == '1':
        # Bước xem danh sách trước khi tạo file: cán bộ bỏ tick những em đã in
        return JsonResponse({'max': MAX_BATCH, 'rows': [
            {'id': c.pk, 'lop': _class_name(c.ghi_chu), 'ho_ten': c.ho_ten, 'cccd': c.so_cmnd or '',
             'ngay_sinh': c.ngay_sinh.strftime('%d/%m/%Y') if c.ngay_sinh else '',
             'anh': len(cccd_image_sides(c.pk)) == 2} for c in customers]})
    template = Template.objects.filter(pk=request.GET.get('template') or 0, is_active=True).first() \
        or Template.objects.filter(name=STUDENT_TEMPLATE_NAME, is_active=True).first()
    if template is None or not template.user_has_access(request.user):
        messages.error(request, 'Chưa chọn mẫu biểu hợp lệ để in theo lô.')
        return redirect('customer_list')
    if not customers:
        messages.warning(request, 'Không có khách hàng nào để in theo lô.')
        return redirect('customer_list')
    if len(customers) > MAX_BATCH:
        messages.error(request, f'Mỗi lô in tối đa {MAX_BATCH} hồ sơ (đang chọn {len(customers)}). Vui lòng thu hẹp danh sách.')
        return redirect('customer_list')
    if not Path(template.file.path).exists():
        messages.error(request, f'Không tìm thấy file mẫu "{template.name}" trên máy chủ.')
        return redirect('customer_list')

    try:
        ngay_in = datetime.strptime(request.GET.get('ngay_in', ''), '%Y-%m-%d').date()
    except ValueError:
        ngay_in = date.today()
    branch_vars = BranchConfig.get_for_user(request.user).get_all_variables()
    streams = []
    for c in customers:
        c.ngay_in = ngay_in                      # chỉ dùng để in, không lưu
        data = c.get_data_dict()
        data['ngay_lap'] = data.get('ngay_in', '')
        fill_empty_print_fields(data)
        data.update(branch_vars)
        streams.append(render_word_template(template.file.path, data))
    if request.GET.get('kieu', 'zip') == 'zip':
        # File nén: mỗi hồ sơ 1 file Word riêng (giống in đơn), số thứ tự đầu tên để giữ thứ tự in
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as zf:
            for i, (c, s) in enumerate(zip(customers, streams), 1):
                parts = [f'{i:03d}', _class_name(c.ghi_chu), c.ho_ten]
                name = re.sub(r'[\\/:*?"<>|\r\n]+', ' ', '_'.join(_plain(p).strip() for p in parts if p and p.strip()))
                zf.writestr(f'{name[:80].strip()}.docx', s.getvalue())
        response = HttpResponse(buf.getvalue(), content_type='application/zip')
        response['Content-Disposition'] = f'attachment; filename="In_theo_lo_{len(customers)}_ho_so_{datetime.now():%Y%m%d_%H%M}.zip"'
        return response
    response = HttpResponse(_merge_docx(streams),
                            content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
    response['Content-Disposition'] = f'attachment; filename="In_theo_lo_{len(customers)}_ho_so_{datetime.now():%Y%m%d_%H%M}.docx"'
    return response
