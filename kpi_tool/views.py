"""
Views cho module quy đổi bút toán
"""
import os
import tempfile
from datetime import datetime
from decimal import Decimal

from django.shortcuts import render, redirect, get_object_or_404
from django.contrib import messages
from django.http import HttpResponse, JsonResponse
from django.db.models import Sum, Count
from django.core.files.storage import default_storage
from django.views.decorators.http import require_http_methods

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

from .models import ConversionRule, TellerTransactionBatch, TellerTransactionDetail
from .services import DBFProcessor, RuleImporter


def index(request):
    """
    Trang chủ - Danh sách các lô đã xử lý
    """
    batches = TellerTransactionBatch.objects.all().order_by('-created_at')[:20]

    # Thống kê tổng quan
    stats = {
        'total_batches': TellerTransactionBatch.objects.filter(processing_status='completed').count(),
        'total_transactions': TellerTransactionDetail.objects.count(),
        'total_score': TellerTransactionDetail.objects.aggregate(Sum('score'))['score__sum'] or 0,
        'total_rules': ConversionRule.objects.filter(is_active=True).count(),
    }

    context = {
        'batches': batches,
        'stats': stats,
    }
    return render(request, 'kpi_tool/index.html', context)


def upload_view(request):
    """
    Trang upload file DBF
    """
    if request.method == 'POST':
        files = request.FILES.getlist('dbf_files')
        teller_name = request.POST.get('teller_name', '').strip()

        if not files:
            messages.error(request, 'Vui lòng chọn ít nhất một file để upload.')
            return redirect('kpi_tool:upload')

        # Lưu file tạm và xử lý
        processor = DBFProcessor()
        batches = []
        errors = []

        for uploaded_file in files:
            # Kiểm tra extension
            if not uploaded_file.name.upper().endswith('.DBF'):
                errors.append(f"{uploaded_file.name}: Không phải file DBF")
                continue

            # Lưu file tạm
            try:
                with tempfile.NamedTemporaryFile(delete=False, suffix='.dbf') as tmp_file:
                    for chunk in uploaded_file.chunks():
                        tmp_file.write(chunk)
                    tmp_file_path = tmp_file.name

                # Xử lý file - truyền tên file gốc để parse đúng
                batch = processor.process_dbf_file(
                    tmp_file_path,
                    teller_name,
                    original_filename=uploaded_file.name
                )
                batches.append(batch)

                # Xóa file tạm
                os.unlink(tmp_file_path)

            except Exception as e:
                errors.append(f"{uploaded_file.name}: {str(e)}")
                continue

        # Hiển thị kết quả
        if batches:
            messages.success(request, f'Đã xử lý thành công {len(batches)} file.')
        if errors:
            for error in errors:
                messages.error(request, error)

        # Redirect đến trang kết quả nếu chỉ có 1 batch
        if len(batches) == 1:
            return redirect('kpi_tool:batch_detail', batch_id=batches[0].id)
        elif batches:
            return redirect('kpi_tool:index')

    return render(request, 'kpi_tool/upload.html')


def batch_detail(request, batch_id):
    """
    Trang chi tiết lô giao dịch
    """
    batch = get_object_or_404(TellerTransactionBatch, id=batch_id)

    # Lấy chi tiết giao dịch
    transactions = batch.transactions.all().select_related('matched_rule')

    # Thống kê theo ngày
    daily_stats = batch.transactions.values('transaction_date').annotate(
        count=Count('id'),
        total_score=Sum('score')
    ).order_by('transaction_date')

    context = {
        'batch': batch,
        'transactions': transactions,
        'daily_stats': daily_stats,
    }
    return render(request, 'kpi_tool/batch_detail.html', context)


def batch_list(request):
    """
    Danh sách tất cả các lô
    """
    # Lọc theo teller, month, year
    teller_id = request.GET.get('teller_id')
    month = request.GET.get('month')
    year = request.GET.get('year')

    batches = TellerTransactionBatch.objects.all()

    if teller_id:
        batches = batches.filter(teller_id__icontains=teller_id)
    if month:
        batches = batches.filter(month=month)
    if year:
        batches = batches.filter(year=year)

    batches = batches.order_by('-created_at')

    # Danh sách teller để lọc
    tellers = TellerTransactionBatch.objects.values('teller_id', 'teller_name').distinct()

    context = {
        'batches': batches,
        'tellers': tellers,
    }
    return render(request, 'kpi_tool/batch_list.html', context)


def export_excel(request, batch_id):
    """
    Xuất báo cáo Excel theo tháng cho giao dịch viên
    """
    batch = get_object_or_404(TellerTransactionBatch, id=batch_id)
    transactions = batch.transactions.all().select_related('matched_rule')

    # Tạo workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Báo cáo tháng"

    # Header style
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=12)
    header_alignment = Alignment(horizontal="center", vertical="center")

    # Info style
    info_font = Font(bold=True, size=11)

    # Border style
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    # Tiêu đề báo cáo
    ws.merge_cells('A1:G1')
    title_cell = ws['A1']
    title_cell.value = f"BÁO CÁO QUY ĐỔI BÚT TOÁN THÁNG {batch.month}/{batch.year}"
    title_cell.font = Font(bold=True, size=14, color="FFFFFF")
    title_cell.alignment = header_alignment
    title_cell.fill = header_fill

    # Thông tin giao dịch viên
    current_row = 3
    ws[f'A{current_row}'] = "Mã giao dịch viên:"
    ws[f'A{current_row}'].font = info_font
    ws[f'B{current_row}'] = batch.teller_id
    ws.merge_cells(f'B{current_row}:C{current_row}')

    current_row += 1
    ws[f'A{current_row}'] = "Tên giao dịch viên:"
    ws[f'A{current_row}'].font = info_font
    ws[f'B{current_row}'] = batch.teller_name
    ws.merge_cells(f'B{current_row}:C{current_row}')

    current_row += 1
    ws[f'A{current_row}'] = "Tháng/Năm:"
    ws[f'A{current_row}'].font = info_font
    ws[f'B{current_row}'] = f"{batch.month}/{batch.year}"

    # Thống kê tổng hợp
    current_row += 2
    ws.merge_cells(f'A{current_row}:G{current_row}')
    ws[f'A{current_row}'] = "THỐNG KÊ TỔNG HỢP"
    ws[f'A{current_row}'].font = Font(bold=True, size=12)
    ws[f'A{current_row}'].alignment = header_alignment
    ws[f'A{current_row}'].fill = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")

    current_row += 1
    stats_data = [
        ["Chỉ tiêu", "Giá trị"],
        ["Tổng số giao dịch", batch.total_transactions],
        ["Giao dịch khớp quy tắc", batch.matched_transactions],
        ["Giao dịch không khớp", batch.unmatched_transactions],
        ["Tỷ lệ khớp", f"{(batch.matched_transactions/batch.total_transactions*100):.1f}%" if batch.total_transactions > 0 else "0%"],
        ["Tổng điểm quy đổi", float(batch.total_score)],
    ]

    for row_data in stats_data:
        ws[f'A{current_row}'] = row_data[0]
        ws[f'B{current_row}'] = row_data[1]
        ws[f'A{current_row}'].font = Font(bold=True) if row_data[0] == "Chỉ tiêu" else None
        ws[f'B{current_row}'].font = Font(bold=True) if row_data[0] == "Chỉ tiêu" else None

        if row_data[0] == "Chỉ tiêu":
            ws[f'A{current_row}'].fill = header_fill
            ws[f'B{current_row}'].fill = header_fill
            ws[f'A{current_row}'].font = header_font
            ws[f'B{current_row}'].font = header_font

        ws[f'A{current_row}'].border = thin_border
        ws[f'B{current_row}'].border = thin_border
        current_row += 1

    # Thống kê theo nghiệp vụ
    current_row += 1
    ws.merge_cells(f'A{current_row}:G{current_row}')
    ws[f'A{current_row}'] = "THỐNG KÊ THEO NGHIỆP VỤ"
    ws[f'A{current_row}'].font = Font(bold=True, size=12)
    ws[f'A{current_row}'].alignment = header_alignment
    ws[f'A{current_row}'].fill = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")

    current_row += 1

    # Header bảng nghiệp vụ
    headers = ['STT', 'Mã nghiệp vụ', 'Mô tả', 'Số lượng GD', 'Tổng điểm']
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=current_row, column=col)
        cell.value = header
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = thin_border

    # Tính toán thống kê theo nghiệp vụ
    from collections import defaultdict
    rule_stats = defaultdict(lambda: {'count': 0, 'total_score': 0, 'description': ''})

    for trans in transactions:
        if trans.matched_rule:
            key = trans.matched_rule.code
            rule_stats[key]['count'] += 1
            rule_stats[key]['total_score'] += float(trans.score)
            rule_stats[key]['description'] = trans.matched_rule.description

    # Giao dịch không khớp
    unmatched_count = batch.unmatched_transactions
    if unmatched_count > 0:
        rule_stats['N/A'] = {
            'count': unmatched_count,
            'total_score': 0,
            'description': 'Không khớp quy tắc'
        }

    # Dữ liệu nghiệp vụ
    current_row += 1
    for idx, (code, stats) in enumerate(sorted(rule_stats.items()), start=1):
        ws.cell(row=current_row, column=1, value=idx).border = thin_border
        ws.cell(row=current_row, column=2, value=code).border = thin_border
        ws.cell(row=current_row, column=3, value=stats['description']).border = thin_border
        ws.cell(row=current_row, column=4, value=stats['count']).border = thin_border
        ws.cell(row=current_row, column=5, value=stats['total_score']).border = thin_border
        current_row += 1

    # Điều chỉnh độ rộng cột
    ws.column_dimensions['A'].width = 8
    ws.column_dimensions['B'].width = 15
    ws.column_dimensions['C'].width = 35
    ws.column_dimensions['D'].width = 15
    ws.column_dimensions['E'].width = 15
    ws.column_dimensions['F'].width = 15
    ws.column_dimensions['G'].width = 15

    # Tạo response
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    filename = f"BaoCaoKPI_Thang{batch.month}_{batch.year}_{batch.teller_id}.xlsx"
    response['Content-Disposition'] = f'attachment; filename="{filename}"'

    wb.save(response)
    return response


def import_rules_view(request):
    """
    Trang import quy tắc từ file hesoquydoi.BAK
    """
    if request.method == 'POST':
        file = request.FILES.get('rule_file')

        if not file:
            messages.error(request, 'Vui lòng chọn file để import.')
            return redirect('kpi_tool:import_rules')

        # Kiểm tra extension
        if not (file.name.upper().endswith('.BAK') or file.name.upper().endswith('.DBF')):
            messages.error(request, 'File phải có định dạng .BAK hoặc .DBF')
            return redirect('kpi_tool:import_rules')

        # Lưu file tạm và xử lý
        try:
            with tempfile.NamedTemporaryFile(delete=False, suffix='.dbf') as tmp_file:
                for chunk in file.chunks():
                    tmp_file.write(chunk)
                tmp_file_path = tmp_file.name

            # Import rules
            importer = RuleImporter()
            count = importer.import_from_dbf(tmp_file_path)

            # Xóa file tạm
            os.unlink(tmp_file_path)

            messages.success(request, f'Đã import thành công {count} quy tắc.')
            return redirect('kpi_tool:rule_list')

        except Exception as e:
            messages.error(request, f'Lỗi khi import: {str(e)}')
            return redirect('kpi_tool:import_rules')

    # Thống kê quy tắc hiện có
    rules_count = ConversionRule.objects.count()
    active_rules_count = ConversionRule.objects.filter(is_active=True).count()

    context = {
        'rules_count': rules_count,
        'active_rules_count': active_rules_count,
    }
    return render(request, 'kpi_tool/import_rules.html', context)


def rule_list(request):
    """
    Danh sách quy tắc
    """
    rules = ConversionRule.objects.all().order_by('code')

    context = {
        'rules': rules,
    }
    return render(request, 'kpi_tool/rule_list.html', context)
