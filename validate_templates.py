#!/usr/bin/env python3
"""
Script validation để kiểm tra templates có bị corrupt không
Chạy script này để tìm và fix templates bị lỗi
"""
import os
import sys
from pathlib import Path


def validate_template(file_path):
    """
    Kiểm tra template file có hợp lệ không

    Args:
        file_path: Đường dẫn tới file .docx

    Returns:
        tuple: (is_valid, error_message)
    """
    try:
        # Try import python-docx
        try:
            from docx import Document
        except ImportError:
            return False, "ERROR: python-docx not installed. Run: pip install python-docx"

        # Check file exists
        if not Path(file_path).exists():
            return False, f"File not found: {file_path}"

        # Check file size
        file_size = Path(file_path).stat().st_size
        if file_size == 0:
            return False, "File is empty (0 bytes)"

        if file_size < 1000:  # Minimum valid .docx is ~1KB
            return False, f"File too small ({file_size} bytes) - likely corrupt"

        # Try to open with python-docx
        doc = Document(file_path)

        # Check basic structure
        para_count = len(doc.paragraphs)
        table_count = len(doc.tables)

        # Document has at least some content
        if para_count == 0 and table_count == 0:
            return False, "Document is empty (no paragraphs or tables)"

        return True, f"Valid ({para_count} paragraphs, {table_count} tables, {file_size:,} bytes)"

    except Exception as e:
        return False, f"Corrupt or invalid: {str(e)}"


def scan_directory(directory):
    """
    Scan toàn bộ thư mục tìm templates và validate

    Args:
        directory: Đường dẫn thư mục cần scan

    Returns:
        dict: Report kết quả
    """
    directory = Path(directory)

    if not directory.exists():
        print(f"❌ Directory not found: {directory}")
        return None

    # Find all .docx files
    docx_files = list(directory.rglob("*.docx"))

    if not docx_files:
        print(f"⚠ No .docx files found in {directory}")
        return None

    results = {
        'total': len(docx_files),
        'valid': 0,
        'invalid': 0,
        'files': []
    }

    print(f"\n🔍 Scanning {len(docx_files)} template files in {directory}\n")
    print("=" * 80)

    for file_path in sorted(docx_files):
        is_valid, message = validate_template(file_path)

        # Relative path for display
        rel_path = file_path.relative_to(directory)

        result = {
            'path': str(file_path),
            'rel_path': str(rel_path),
            'valid': is_valid,
            'message': message
        }
        results['files'].append(result)

        if is_valid:
            results['valid'] += 1
            print(f"✅ {rel_path}")
            print(f"   {message}")
        else:
            results['invalid'] += 1
            print(f"❌ {rel_path}")
            print(f"   {message}")

        print()

    # Summary
    print("=" * 80)
    print(f"\n📊 SUMMARY:")
    print(f"   Total files:   {results['total']}")
    print(f"   ✅ Valid:      {results['valid']}")
    print(f"   ❌ Invalid:    {results['invalid']}")

    if results['invalid'] > 0:
        print(f"\n⚠️  WARNING: {results['invalid']} corrupt templates found!")
        print("   These templates will cause 'We can't open...' errors")
        print("   → Replace them with valid .docx files")
    else:
        print(f"\n✨ All templates are valid!")

    return results


def main():
    """Main function"""
    import argparse

    parser = argparse.ArgumentParser(
        description="Validate Word template files (.docx)",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
Examples:
  # Validate single file
  python validate_templates.py /path/to/template.docx

  # Scan entire directory
  python validate_templates.py /home/user/maubieumoi

  # Scan media folder
  python validate_templates.py media/templates/docx
        """
    )
    parser.add_argument(
        'path',
        nargs='?',
        default=None,
        help='Path to template file or directory (default: auto-detect)'
    )
    parser.add_argument(
        '--fix',
        action='store_true',
        help='Auto-remove corrupt files (DANGEROUS!)'
    )

    args = parser.parse_args()

    # Auto-detect paths if not provided
    if args.path is None:
        # Try common paths
        possible_paths = [
            '/home/user/maubieumoi',
            'media/templates/docx',
            '/home/user/MauBieu7202/media/templates/docx',
        ]

        found_path = None
        for path in possible_paths:
            if Path(path).exists():
                found_path = path
                break

        if found_path:
            print(f"🔍 Auto-detected template folder: {found_path}\n")
            args.path = found_path
        else:
            print("❌ No template folder found. Please specify path:")
            print("   python validate_templates.py /path/to/templates")
            sys.exit(1)

    path = Path(args.path)

    # Single file
    if path.is_file():
        print(f"🔍 Validating single file: {path}\n")
        is_valid, message = validate_template(path)

        if is_valid:
            print(f"✅ Valid template")
            print(f"   {message}")
            sys.exit(0)
        else:
            print(f"❌ Invalid template")
            print(f"   {message}")
            sys.exit(1)

    # Directory
    elif path.is_dir():
        results = scan_directory(path)

        if results is None:
            sys.exit(1)

        # Fix mode
        if args.fix and results['invalid'] > 0:
            print(f"\n⚠️  FIX MODE ENABLED")
            print("=" * 80)

            for file_info in results['files']:
                if not file_info['valid']:
                    file_path = Path(file_info['path'])
                    print(f"🗑️  Removing corrupt file: {file_info['rel_path']}")

                    try:
                        file_path.unlink()
                        print("   ✓ Deleted")
                    except Exception as e:
                        print(f"   ✗ Failed: {e}")

            print("\n✨ Cleanup complete!")

        # Exit code
        sys.exit(0 if results['invalid'] == 0 else 1)

    else:
        print(f"❌ Path not found: {path}")
        sys.exit(1)


if __name__ == '__main__':
    main()
