#!/usr/bin/env python
"""
Script test để tìm và điền table trong Word
"""

import sys
from pathlib import Path
from docx import Document

def find_and_fill_card_name_table(template_path, output_path, card_name):
    """
    Tìm table có cấu trúc dãy ô vuông và điền tên thẻ vào

    Args:
        template_path: Đường dẫn template
        output_path: Đường dẫn output
        card_name: Tên trên thẻ (VD: "NGUYEN VAN A")
    """
    doc = Document(template_path)

    print(f"\n{'='*80}")
    print(f"ANALYZING TABLES")
    print(f"{'='*80}\n")

    print(f"Total tables in document: {len(doc.tables)}\n")

    # Analyze each table
    for idx, table in enumerate(doc.tables, 1):
        print(f"Table #{idx}:")
        print(f"  Rows: {len(table.rows)}")
        print(f"  Columns: {len(table.columns) if table.rows else 0}")

        # Check if this might be the card name table
        if table.rows and len(table.columns) >= 20:
            print(f"  → POTENTIAL CARD NAME TABLE (has {len(table.columns)} columns)")

            # Preview first row
            first_row = table.rows[0]
            cell_texts = [cell.text.strip() for cell in first_row.cells[:5]]
            print(f"  First 5 cells: {cell_texts}")

        print()

    # Now let's try to fill the first table with many columns
    filled = False
    for table in doc.tables:
        if not table.rows:
            continue

        # Look for table with 20+ columns (likely card name table)
        if len(table.columns) >= 20:
            print(f"\n{'='*80}")
            print(f"FILLING TABLE")
            print(f"{'='*80}\n")
            print(f"Card name: {card_name}")
            print(f"Table columns: {len(table.columns)}")
            print(f"Filling first row with characters...")

            # Fill first row (or you can specify which row)
            row = table.rows[0]  # Change to table.rows[1] if header row exists

            for i, char in enumerate(card_name):
                if i >= len(row.cells):
                    print(f"  Warning: Card name longer than table ({i+1} chars > {len(row.cells)} cells)")
                    break

                row.cells[i].text = char
                print(f"  Cell {i+1}: '{char}'")

            # Clear remaining cells
            for i in range(len(card_name), len(row.cells)):
                row.cells[i].text = ''

            filled = True
            break

    if filled:
        doc.save(output_path)
        print(f"\n✓ Saved to: {output_path}")
        print(f"\nOpen the file to verify the table is filled correctly!")
    else:
        print(f"\n✗ Could not find suitable table (need 20+ columns)")


if __name__ == '__main__':
    if len(sys.argv) < 4:
        print("Usage: python test_fill_table.py <template.docx> <output.docx> <card_name>")
        print("\nExample:")
        print('  python test_fill_table.py template.docx output.docx "NGUYEN VAN A"')
        sys.exit(1)

    template_path = sys.argv[1]
    output_path = sys.argv[2]
    card_name = sys.argv[3].upper()  # Uppercase for card name

    find_and_fill_card_name_table(template_path, output_path, card_name)
