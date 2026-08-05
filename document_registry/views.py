import io
from datetime import date
from django.shortcuts import render, get_object_or_404, redirect
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.core.paginator import Paginator
from django.db.models import Q
from django.http import HttpResponse

import openpyxl
from openpyxl.styles import (
    PatternFill, Font, Alignment, Border, Side, GradientFill
)
from openpyxl.utils import get_column_letter

from .models import OutgoingDocument, IncomingDocument
from .forms import OutgoingDocumentForm, IncomingDocumentForm


# ─── EXCEL HELPERS ──────────────────────────────────────────────────────────

def _make_border():
    thin = Side(style='thin', color='AAAAAA')
    return Border(left=thin, right=thin, top=thin, bottom=thin)


def _header_cell(ws, row, col, value, fill_hex, font_color='FFFFFF', font_size=10):
    cell = ws.cell(row=row, column=col, value=value)
    cell.fill = PatternFill('solid', fgColor=fill_hex)
    cell.font = Font(bold=True, color=font_color, size=font_size)
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    cell.border = _make_border()
    return cell


def _data_cell(ws, row, col, value, wrap=True, align='left'):
    cell = ws.cell(row=row, column=col, value=value)
    cell.alignment = Alignment(horizontal=align, vertical='top', wrap_text=wrap)
    cell.border = _make_border()
    return cell


def _set_col_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _fmt_date(d):
    return d.strftime('%d/%m/%Y') if d else ''


def _iso_to_dmy(iso_str):
    """Convert 'YYYY-MM-DD' string from HTML date input to 'dd/mm/YYYY'."""
    if not iso_str:
        return ''
    try:
        from datetime import datetime
        return datetime.strptime(iso_str, '%Y-%m-%d').strftime('%d/%m/%Y')
    except ValueError:
        return iso_str


@login_required
def dashboard(request):
    outgoing_count = OutgoingDocument.objects.count()
    incoming_count = IncomingDocument.objects.count()
    recent_outgoing = OutgoingDocument.objects.select_related('created_by')[:5]
    recent_incoming = IncomingDocument.objects.select_related('created_by')[:5]
    return render(request, 'document_registry/dashboard.html', {
        'outgoing_count': outgoing_count,
        'incoming_count': incoming_count,
        'recent_outgoing': recent_outgoing,
        'recent_incoming': recent_incoming,
    })


# ─── OUTGOING ───────────────────────────────────────────────────────────────

@login_required
def outgoing_list(request):
    qs = OutgoingDocument.objects.select_related('created_by')
    q = request.GET.get('q', '').strip()
    date_from = request.GET.get('date_from', '').strip()
    date_to = request.GET.get('date_to', '').strip()
    loai_chuyen_phat = request.GET.get('loai_chuyen_phat', '').strip()

    if q:
        qs = qs.filter(
            Q(so_ky_hieu__icontains=q) |
            Q(ten_loai_trich_yeu__icontains=q) |
            Q(nguoi_ky_khac__icontains=q) |
            Q(noi_nhan_khac__icontains=q)
        )
    if date_from:
        qs = qs.filter(ngay_van_ban__gte=date_from)
    if date_to:
        qs = qs.filter(ngay_van_ban__lte=date_to)
    if loai_chuyen_phat in ('noi_bo', 'buu_dien'):
        qs = qs.filter(loai_chuyen_phat=loai_chuyen_phat)

    paginator = Paginator(qs, 20)
    page = paginator.get_page(request.GET.get('page'))
    return render(request, 'document_registry/outgoing_list.html', {
        'page_obj': page,
        'q': q,
        'date_from': date_from,
        'date_to': date_to,
        'loai_chuyen_phat': loai_chuyen_phat,
    })


@login_required
def outgoing_create(request):
    today = date.today().isoformat()
    if request.method == 'POST':
        form = OutgoingDocumentForm(request.POST)
        if form.is_valid():
            doc = form.save(commit=False)
            doc.created_by = request.user
            doc.save()
            messages.success(request, f'Đã thêm văn bản đi: {doc.so_ky_hieu}')
            return redirect('document_registry:outgoing_list')
    else:
        form = OutgoingDocumentForm(initial={
            'ngay_van_ban': today,
            'ngay_chuyen': today,
            'so_luong_ban': 1,
        })
    return render(request, 'document_registry/outgoing_form.html', {
        'form': form,
        'title': 'Thêm văn bản đi',
        'action': 'create',
    })


@login_required
def outgoing_edit(request, pk):
    doc = get_object_or_404(OutgoingDocument, pk=pk)
    if request.method == 'POST':
        form = OutgoingDocumentForm(request.POST, instance=doc)
        if form.is_valid():
            form.save()
            messages.success(request, f'Đã cập nhật văn bản đi: {doc.so_ky_hieu}')
            return redirect('document_registry:outgoing_list')
    else:
        form = OutgoingDocumentForm(instance=doc)
    return render(request, 'document_registry/outgoing_form.html', {
        'form': form,
        'doc': doc,
        'title': f'Sửa văn bản đi: {doc.so_ky_hieu}',
        'action': 'edit',
    })


@login_required
def outgoing_delete(request, pk):
    doc = get_object_or_404(OutgoingDocument, pk=pk)
    if request.method == 'POST':
        label = str(doc)
        doc.delete()
        messages.success(request, f'Đã xóa văn bản đi: {label}')
        return redirect('document_registry:outgoing_list')
    return render(request, 'document_registry/confirm_delete.html', {
        'obj': doc,
        'title': 'Xóa văn bản đi',
        'cancel_url': 'document_registry:outgoing_list',
    })


# ─── INCOMING ───────────────────────────────────────────────────────────────

@login_required
def incoming_list(request):
    qs = IncomingDocument.objects.select_related('created_by')
    q = request.GET.get('q', '').strip()
    date_from = request.GET.get('date_from', '').strip()
    date_to = request.GET.get('date_to', '').strip()

    if q:
        qs = qs.filter(
            Q(so_den__icontains=q) |
            Q(so_ky_hieu__icontains=q) |
            Q(tac_gia__icontains=q) |
            Q(ten_loai_trich_yeu__icontains=q) |
            Q(don_vi_nguoi_nhan__icontains=q)
        )
    if date_from:
        qs = qs.filter(ngay_van_ban__gte=date_from)
    if date_to:
        qs = qs.filter(ngay_van_ban__lte=date_to)

    paginator = Paginator(qs, 20)
    page = paginator.get_page(request.GET.get('page'))
    return render(request, 'document_registry/incoming_list.html', {
        'page_obj': page,
        'q': q,
        'date_from': date_from,
        'date_to': date_to,
    })


@login_required
def incoming_create(request):
    today = date.today().isoformat()
    if request.method == 'POST':
        form = IncomingDocumentForm(request.POST)
        if form.is_valid():
            doc = form.save(commit=False)
            doc.created_by = request.user
            doc.save()
            messages.success(request, f'Đã thêm văn bản đến số: {doc.so_den}')
            return redirect('document_registry:incoming_list')
    else:
        form = IncomingDocumentForm(initial={
            'ngay_den': today,
            'ngay_van_ban': today,
            'ngay_chuyen': today,
        })
    return render(request, 'document_registry/incoming_form.html', {
        'form': form,
        'title': 'Thêm văn bản đến',
        'action': 'create',
    })


@login_required
def incoming_edit(request, pk):
    doc = get_object_or_404(IncomingDocument, pk=pk)
    if request.method == 'POST':
        form = IncomingDocumentForm(request.POST, instance=doc)
        if form.is_valid():
            form.save()
            messages.success(request, f'Đã cập nhật văn bản đến số: {doc.so_den}')
            return redirect('document_registry:incoming_list')
    else:
        form = IncomingDocumentForm(instance=doc)
    return render(request, 'document_registry/incoming_form.html', {
        'form': form,
        'doc': doc,
        'title': f'Sửa văn bản đến: {doc.so_den}',
        'action': 'edit',
    })


@login_required
def incoming_delete(request, pk):
    doc = get_object_or_404(IncomingDocument, pk=pk)
    if request.method == 'POST':
        label = str(doc)
        doc.delete()
        messages.success(request, f'Đã xóa văn bản đến: {label}')
        return redirect('document_registry:incoming_list')
    return render(request, 'document_registry/confirm_delete.html', {
        'obj': doc,
        'title': 'Xóa văn bản đến',
        'cancel_url': 'document_registry:incoming_list',
    })


# ─── EXPORT EXCEL ────────────────────────────────────────────────────────────

@login_required
def outgoing_export_excel(request):
    date_from = request.GET.get('date_from', '').strip()
    date_to = request.GET.get('date_to', '').strip()
    loai_chuyen_phat = request.GET.get('loai_chuyen_phat', '').strip()

    qs = OutgoingDocument.objects.all().order_by('ngay_van_ban', 'id')

    if date_from:
        qs = qs.filter(ngay_van_ban__gte=date_from)
    if date_to:
        qs = qs.filter(ngay_van_ban__lte=date_to)
    if loai_chuyen_phat in ('noi_bo', 'buu_dien'):
        qs = qs.filter(loai_chuyen_phat=loai_chuyen_phat)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Văn bản đi'
    ws.sheet_view.showGridLines = False

    # Title row
    title_label = 'SỔ ĐĂNG KÝ VĂN BẢN ĐI'
    if loai_chuyen_phat == 'noi_bo':
        title_label += ' – CHUYỂN PHÁT NỘI BỘ'
    elif loai_chuyen_phat == 'buu_dien':
        title_label += ' – CHUYỂN PHÁT BƯU ĐIỆN'
    if date_from or date_to:
        title_label += f'  (Từ {_iso_to_dmy(date_from) or "..."} đến {_iso_to_dmy(date_to) or "..."})'

    ws.merge_cells('A1:K1')
    title_cell = ws['A1']
    title_cell.value = title_label
    title_cell.font = Font(bold=True, size=13, color='FFFFFF')
    title_cell.fill = PatternFill('solid', fgColor='AE1C3F')
    title_cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 28

    # Header row
    HEADERS = [
        'STT',
        'Số, ký hiệu\nVăn bản',
        'Ngày\nVăn bản',
        'Tên loại và trích yếu\nnội dung Văn bản',
        'Người ký',
        'Nơi nhận\nVăn bản',
        'Đơn vị, người nhận\nbản lưu',
        'Số lượng\nbản',
        'Ngày\nchuyển',
        'Ký nhận',
        'Ghi chú',
    ]
    for col, h in enumerate(HEADERS, start=1):
        _header_cell(ws, 2, col, h, 'C0392B')
    ws.row_dimensions[2].height = 36

    # Data rows
    for i, doc in enumerate(qs, start=1):
        r = i + 2
        ky_nhan_val = doc.ky_nhan_buu_dien if doc.loai_chuyen_phat == 'buu_dien' else doc.ky_nhan
        _data_cell(ws, r, 1, i, align='center')
        _data_cell(ws, r, 2, doc.so_ky_hieu)
        _data_cell(ws, r, 3, _fmt_date(doc.ngay_van_ban), align='center')
        _data_cell(ws, r, 4, doc.ten_loai_trich_yeu)
        _data_cell(ws, r, 5, doc.get_nguoi_ky_display_name())
        _data_cell(ws, r, 6, doc.get_noi_nhan_display())
        _data_cell(ws, r, 7, doc.don_vi_nhan_ban_luu)
        _data_cell(ws, r, 8, doc.so_luong_ban, align='center')
        _data_cell(ws, r, 9, _fmt_date(doc.ngay_chuyen), align='center')
        _data_cell(ws, r, 10, ky_nhan_val)
        _data_cell(ws, r, 11, doc.ghi_chu)
        ws.row_dimensions[r].height = 18

    _set_col_widths(ws, [5, 18, 12, 42, 18, 30, 26, 9, 12, 16, 20])

    filename_parts = ['VanBanDi']
    if date_from:
        filename_parts.append(date_from.replace('-', ''))
    if date_to:
        filename_parts.append(date_to.replace('-', ''))
    if loai_chuyen_phat == 'noi_bo':
        filename_parts.append('NoiBo')
    elif loai_chuyen_phat == 'buu_dien':
        filename_parts.append('BuuDien')
    filename = '_'.join(filename_parts) + '.xlsx'

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    response = HttpResponse(
        buf.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


@login_required
def incoming_export_excel(request):
    date_from = request.GET.get('date_from', '').strip()
    date_to = request.GET.get('date_to', '').strip()

    qs = IncomingDocument.objects.all().order_by('ngay_van_ban', 'id')

    if date_from:
        qs = qs.filter(ngay_van_ban__gte=date_from)
    if date_to:
        qs = qs.filter(ngay_van_ban__lte=date_to)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Văn bản đến'
    ws.sheet_view.showGridLines = False

    title_label = 'SỔ ĐĂNG KÝ VĂN BẢN ĐẾN'
    if date_from or date_to:
        title_label += f'  (Từ {_iso_to_dmy(date_from) or "..."} đến {_iso_to_dmy(date_to) or "..."})'

    ws.merge_cells('A1:K1')
    title_cell = ws['A1']
    title_cell.value = title_label
    title_cell.font = Font(bold=True, size=13, color='FFFFFF')
    title_cell.fill = PatternFill('solid', fgColor='1A6E2C')
    title_cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 28

    HEADERS = [
        'STT',
        'Ngày\nđến',
        'Số\nđến',
        'Tác giả',
        'Số, ký hiệu\nVăn bản',
        'Ngày\nVăn bản',
        'Tên loại và trích yếu\nnội dung Văn bản',
        'Đơn vị hoặc\nngười nhận',
        'Ngày\nchuyển',
        'Ký nhận',
        'Ghi chú',
    ]
    for col, h in enumerate(HEADERS, start=1):
        _header_cell(ws, 2, col, h, '1A6E2C')
    ws.row_dimensions[2].height = 36

    for i, doc in enumerate(qs, start=1):
        r = i + 2
        _data_cell(ws, r, 1, i, align='center')
        _data_cell(ws, r, 2, _fmt_date(doc.ngay_den), align='center')
        _data_cell(ws, r, 3, doc.so_den, align='center')
        _data_cell(ws, r, 4, doc.tac_gia)
        _data_cell(ws, r, 5, doc.so_ky_hieu)
        _data_cell(ws, r, 6, _fmt_date(doc.ngay_van_ban), align='center')
        _data_cell(ws, r, 7, doc.ten_loai_trich_yeu)
        _data_cell(ws, r, 8, doc.don_vi_nguoi_nhan)
        _data_cell(ws, r, 9, _fmt_date(doc.ngay_chuyen), align='center')
        _data_cell(ws, r, 10, doc.ky_nhan)
        _data_cell(ws, r, 11, doc.ghi_chu)
        ws.row_dimensions[r].height = 18

    _set_col_widths(ws, [5, 12, 8, 24, 18, 12, 42, 26, 12, 16, 20])

    filename_parts = ['VanBanDen']
    if date_from:
        filename_parts.append(date_from.replace('-', ''))
    if date_to:
        filename_parts.append(date_to.replace('-', ''))
    filename = '_'.join(filename_parts) + '.xlsx'

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    response = HttpResponse(
        buf.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response
