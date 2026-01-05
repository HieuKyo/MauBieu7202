"""
Generate Sample Data for KPI Dashboard Testing
===============================================
Script tạo dữ liệu mẫu để test ứng dụng KPI Dashboard
"""

import pandas as pd
from datetime import datetime, timedelta
import random
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
import os


def generate_card_data(month=1, year=2024, user_id='GRATHIEU'):
    """
    Tạo dữ liệu mẫu cho file Thẻ
    """
    data = []
    start_date = datetime(year, month, 1)

    # Tạo dữ liệu ngẫu nhiên cho mỗi ngày trong tháng
    for day in range(1, 29):  # 28 ngày để đảm bảo tất cả tháng đều có
        current_date = start_date + timedelta(days=day - 1)

        # Tạo 0-5 bản ghi mỗi ngày
        num_records = random.randint(0, 5)
        for _ in range(num_records):
            issue_type = random.choice(['New Issue', 'Reissue', 'Reissue'])  # 1/3 là New Issue
            data.append({
                'dlvrydt': current_date.strftime('%d/%m/%Y'),
                'dlvryusrid': user_id,
                'isutycd': issue_type,
                'cardno': f'9704{random.randint(1000000000, 9999999999)}'
            })

    df = pd.DataFrame(data)
    return df


def generate_sms_data(month=1, year=2024, user_id='GRATHIEU'):
    """
    Tạo dữ liệu mẫu cho file SMS
    """
    data = []
    start_date = datetime(year, month, 1)

    for day in range(1, 29):
        current_date = start_date + timedelta(days=day - 1)

        # Tạo 0-3 bản ghi mỗi ngày
        num_records = random.randint(0, 3)
        for _ in range(num_records):
            data.append({
                'entydt': current_date.strftime('%Y-%m-%d'),
                'crtusr': user_id,
                'mobile': f'09{random.randint(10000000, 99999999)}',
                'service_type': 'SMS Banking'
            })

    df = pd.DataFrame(data)
    return df


def generate_emobile_data(month=1, year=2024, user_id='GRATHIEU'):
    """
    Tạo dữ liệu mẫu cho file E-Mobile Banking
    """
    data = []
    start_date = datetime(year, month, 1)

    for day in range(1, 29):
        current_date = start_date + timedelta(days=day - 1)

        # Tạo 0-4 bản ghi mỗi ngày
        num_records = random.randint(0, 4)
        for _ in range(num_records):
            data.append({
                'entydt': current_date.strftime('%Y-%m-%d'),
                'crtusr': user_id,
                'device_id': f'DEV{random.randint(10000, 99999)}',
                'app_version': random.choice(['2.5.1', '2.5.2', '2.6.0'])
            })

    df = pd.DataFrame(data)
    return df


def generate_kpi_template(month=1, year=2024):
    """
    Tạo file KPI Template mẫu
    """
    wb = Workbook()
    ws = wb.active
    ws.title = f"KPI_T{month}_{year}"

    # Styling
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    center_align = Alignment(horizontal="center", vertical="center")

    # Title
    ws.merge_cells('A1:AI1')
    ws['A1'] = f'BẢNG TỔNG HỢP KPI GIAO DỊCH VIÊN - THÁNG {month}/{year}'
    ws['A1'].font = Font(bold=True, size=14)
    ws['A1'].alignment = center_align

    # Header row 1
    ws['A3'] = 'STT'
    ws['B3'] = 'NỘI DUNG'
    ws['C3'] = 'ĐVT'
    ws['D3'] = 'TỔNG CHI NHÁNH'
    ws['E3'] = 'HỆ SỐ'

    # Days header (row 3, starting from column F)
    days_in_month = 31  # Tối đa 31 ngày
    for day in range(1, days_in_month + 1):
        col_idx = 5 + day  # F=6, G=7, ...
        cell = ws.cell(row=3, column=col_idx)
        cell.value = day
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align

    # "Cộng" và "Bút toán quy đổi"
    ws.cell(row=3, column=5 + days_in_month + 1).value = "Cộng"
    ws.cell(row=3, column=5 + days_in_month + 2).value = "Bút toán quy đổi"

    # Apply header style
    for col in ['A', 'B', 'C', 'D', 'E']:
        ws[f'{col}3'].fill = header_fill
        ws[f'{col}3'].font = header_font
        ws[f'{col}3'].alignment = center_align

    # KPI rows
    kpi_items = [
        (1, 'Đăng ký TT KH cá nhân', 'Hồ sơ', 3.0),
        (2, 'Rà soát, bổ sung thông tin KH', 'Hồ sơ', 1.0),
        (3, 'Quét chữ ký KH', 'Lượt', 3.0),
        (4, 'GDV lưu trữ HS', 'Hồ sơ', 0.5),
        (5, 'Thu hồi thẻ', 'Thẻ', 2.0),
        (6, 'Cấp lại thẻ', 'Thẻ', 2.0),
        (7, 'Đăng ký Internet Banking', 'KH', 4.0),
        (8, 'Đăng ký Agribank Plus', 'KH', 4.0),
        (9, 'Đăng ký SMS', 'KH', 4.0),
        (10, 'Đăng ký E-Statement', 'KH', 2.0),
        (11, 'Mở sổ tiết kiệm', 'Sổ', 1.0),
        (12, 'Phát hành thẻ', 'Thẻ', 3.0),
    ]

    start_row = 4
    for idx, (stt, content, unit, coefficient) in enumerate(kpi_items):
        row = start_row + idx
        ws.cell(row=row, column=1).value = stt
        ws.cell(row=row, column=2).value = content
        ws.cell(row=row, column=3).value = unit
        ws.cell(row=row, column=5).value = coefficient

        # Format
        ws.cell(row=row, column=1).alignment = center_align
        ws.cell(row=row, column=3).alignment = center_align
        ws.cell(row=row, column=5).alignment = center_align

    # Adjust column widths
    ws.column_dimensions['A'].width = 6
    ws.column_dimensions['B'].width = 35
    ws.column_dimensions['C'].width = 10
    ws.column_dimensions['D'].width = 15
    ws.column_dimensions['E'].width = 10

    for day in range(1, days_in_month + 1):
        col_letter = ws.cell(row=3, column=5 + day).column_letter
        ws.column_dimensions[col_letter].width = 4

    return wb


def main():
    """
    Tạo tất cả các file mẫu
    """
    print("🔄 Đang tạo dữ liệu mẫu...")

    # Tạo thư mục sample_data nếu chưa có
    output_dir = os.path.join(os.path.dirname(__file__), 'sample_data')
    os.makedirs(output_dir, exist_ok=True)

    month = 1
    year = 2024
    user_id = 'GRATHIEU'

    # 1. Tạo file Card data
    print("  📄 Tạo file dữ liệu Thẻ...")
    card_df = generate_card_data(month, year, user_id)
    card_path = os.path.join(output_dir, f'sample_card_data_{month:02d}_{year}.xlsx')
    card_df.to_excel(card_path, index=False)
    print(f"    ✅ Đã tạo: {card_path} ({len(card_df)} bản ghi)")

    # 2. Tạo file SMS data
    print("  📄 Tạo file dữ liệu SMS...")
    sms_df = generate_sms_data(month, year, user_id)
    sms_path = os.path.join(output_dir, f'sample_sms_data_{month:02d}_{year}.xlsx')
    sms_df.to_excel(sms_path, index=False)
    print(f"    ✅ Đã tạo: {sms_path} ({len(sms_df)} bản ghi)")

    # 3. Tạo file E-Mobile data
    print("  📄 Tạo file dữ liệu E-Mobile...")
    emobile_df = generate_emobile_data(month, year, user_id)
    emobile_path = os.path.join(output_dir, f'sample_emobile_data_{month:02d}_{year}.xlsx')
    emobile_df.to_excel(emobile_path, index=False)
    print(f"    ✅ Đã tạo: {emobile_path} ({len(emobile_df)} bản ghi)")

    # 4. Tạo KPI Template
    print("  📄 Tạo file KPI Template...")
    wb = generate_kpi_template(month, year)
    template_path = os.path.join(output_dir, f'sample_kpi_template_{month:02d}_{year}.xlsx')
    wb.save(template_path)
    print(f"    ✅ Đã tạo: {template_path}")

    print("\n✅ Hoàn thành! Tất cả file mẫu đã được tạo trong thư mục 'sample_data/'")
    print(f"\nThông tin:")
    print(f"  - Tháng: {month}/{year}")
    print(f"  - User ID: {user_id}")
    print(f"  - Số bản ghi Card: {len(card_df)}")
    print(f"  - Số bản ghi SMS: {len(sms_df)}")
    print(f"  - Số bản ghi E-Mobile: {len(emobile_df)}")


if __name__ == "__main__":
    main()
